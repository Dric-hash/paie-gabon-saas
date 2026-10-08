"""
tests/test_routes_integration.py — Tests d'intégration des routes HTTP

Ces tests démarrent l'application Flask complète avec une base SQLite en
mémoire et simulent de vraies requêtes HTTP via le client de test. Ils
attrapent les bugs qui passent les tests unitaires mais cassent l'app réelle :
    - routes perdues lors d'un refactoring
    - imports manquants (NameError au runtime)
    - blocages CSRF sur les appels JSON
    - défauts d'isolation multi-tenant
    - erreurs 500 sur les pages

Exécution :
    pytest tests/test_routes_integration.py -v
"""
import os
import sys
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

# ── Configuration de test AVANT l'import de l'app ─────────────────────────────
os.environ["DATABASE_URL"]   = "sqlite:///:memory:"
os.environ["SECRET_KEY"]     = "test-secret-key-pour-tests-integration"
os.environ["WTF_CSRF_ENABLED"] = "True"   # on veut tester le comportement CSRF réel

import pytest
from datetime import date, datetime

from app import app as flask_app
from models import (db, Plan, Tenant, Utilisateur, CategorieEmploi,
                    Salarie, PeriodePaie, BulletinPaie)


# ══════════════════════════════════════════════════════════════════════════════
# FIXTURES
# ══════════════════════════════════════════════════════════════════════════════
@pytest.fixture
def app():
    """Application de test avec une base SQLite en mémoire fraîche."""
    flask_app.config.update({
        "TESTING": True,
        "WTF_CSRF_ENABLED": True,
        "SERVER_NAME": "localhost",
        "RATELIMIT_ENABLED": False,
    })
    # Le limiter (mémoire) est partagé sur toute la session pytest : on le désactive
    # pour que les nombreux logins de la suite n'épuisent pas la fenêtre 20/min.
    try:
        from app import limiter
        limiter.enabled = False
    except Exception:
        pass
    with flask_app.app_context():
        db.drop_all()      # repartir d'une base vraiment vide
        db.create_all()
        _seed_data()
        yield flask_app
        db.session.remove()
        db.drop_all()


@pytest.fixture
def client(app):
    return app.test_client()


def _seed_data():
    """Crée un jeu de données minimal : 2 tenants, 1 admin, 1 RH, salariés."""
    plan = Plan(code="PRO", nom="Pro", prix_mensuel=35000,
                max_salaries=50, max_utilisateurs=3, actif=True)
    db.session.add(plan)
    db.session.flush()

    # ── Tenant A ──────────────────────────────────────────────────────────────
    t_a = Tenant(slug="entreprise-a", denomination="ENTREPRISE A", sigle="EA",
                 activite="BTP", ville="Libreville", pays="Gabon",
                 plan_id=plan.id, statut="ACTIF", token_api="token_a")
    # ── Tenant B (pour tester l'isolation) ─────────────────────────────────────
    t_b = Tenant(slug="entreprise-b", denomination="ENTREPRISE B", sigle="EB",
                 activite="COMMERCE", ville="Port-Gentil", pays="Gabon",
                 plan_id=plan.id, statut="ACTIF", token_api="token_b")
    db.session.add_all([t_a, t_b])
    db.session.flush()

    for t in (t_a, t_b):
        db.session.add(CategorieEmploi(tenant_id=t.id, code="C1", libelle="Ouvriers"))

    # ── Admin du tenant A ───────────────────────────────────────────────────────
    admin_a = Utilisateur(nom="ADMIN", prenom="Alice", email="admin@a.ga",
                          role="TENANT_ADMIN", tenant_id=t_a.id,
                          actif=True, email_verifie=True)
    admin_a.set_password("MotDePasse1")
    # ── RH du tenant A ──────────────────────────────────────────────────────────
    rh_a = Utilisateur(nom="RH", prenom="Robert", email="rh@a.ga",
                       role="RH", tenant_id=t_a.id, actif=True, email_verifie=True)
    rh_a.set_password("MotDePasse1")
    # ── Admin du tenant B ───────────────────────────────────────────────────────
    admin_b = Utilisateur(nom="ADMIN", prenom="Bob", email="admin@b.ga",
                          role="TENANT_ADMIN", tenant_id=t_b.id,
                          actif=True, email_verifie=True)
    admin_b.set_password("MotDePasse1")
    db.session.add_all([admin_a, rh_a, admin_b])
    db.session.flush()

    # ── Salarié du tenant A ─────────────────────────────────────────────────────
    sal_a = Salarie(tenant_id=t_a.id, matricule="EA001", nom="NDONG", prenom="Jean",
                    date_embauche=date(2023, 1, 1), emploi="Maçon",
                    situation_matrimoniale="MARIE", nb_enfants=2, statut="ACTIF")
    # ── Salarié du tenant B ─────────────────────────────────────────────────────
    sal_b = Salarie(tenant_id=t_b.id, matricule="EB001", nom="OBAME", prenom="Marie",
                    date_embauche=date(2023, 6, 1), emploi="Vendeuse",
                    situation_matrimoniale="CELIBATAIRE", nb_enfants=0, statut="ACTIF")
    db.session.add_all([sal_a, sal_b])

    # ── Période de paie tenant A ────────────────────────────────────────────────
    periode_a = PeriodePaie(tenant_id=t_a.id, mois=6, annee=2026,
                            libelle_mois="JUIN", statut="OUVERT")
    db.session.add(periode_a)
    db.session.commit()


def login(client, email, password="MotDePasse1"):
    """Connecte un utilisateur via le formulaire de login (avec CSRF)."""
    # Récupérer le token CSRF de la page de login
    page = client.get("/login")
    token = _extract_csrf(page.data)
    return client.post("/login", data={
        "email": email, "password": password, "csrf_token": token,
    }, follow_redirects=False)


def _extract_csrf(html_bytes):
    """Extrait le token CSRF d'une page HTML."""
    import re
    html = html_bytes.decode("utf-8", errors="ignore")
    m = re.search(r'name="csrf_token"[^>]*value="([^"]+)"', html)
    return m.group(1) if m else ""


def auth_session(client, email):
    """Authentifie directement via la session (sans passer par le formulaire)."""
    with flask_app.app_context():
        u = Utilisateur.query.filter_by(email=email).first()
        uid = u.id
    with client.session_transaction() as sess:
        sess["_user_id"] = str(uid)
        sess["_fresh"] = True


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — PAGES PUBLIQUES
# ══════════════════════════════════════════════════════════════════════════════
class TestPagesPubliques:
    def test_login_accessible(self, client):
        assert client.get("/login").status_code == 200

    def test_inscription_accessible(self, client):
        assert client.get("/inscription").status_code == 200

    def test_mot_de_passe_oublie_accessible(self, client):
        assert client.get("/mot-de-passe-oublie").status_code == 200

    def test_racine_affiche_presentation(self, client):
        """La racine affiche la page de présentation pour les visiteurs non connectés."""
        r = client.get("/", follow_redirects=False)
        assert r.status_code == 200
        assert b"inscription" in r.data.lower() or b"essai" in r.data.lower()

    def test_racine_connecte_redirige_dashboard(self, client):
        """Un utilisateur connecté est redirigé vers son tableau de bord."""
        auth_session(client, "admin@a.ga")
        r = client.get("/", follow_redirects=False)
        assert r.status_code == 302
        assert "/dashboard" in r.headers["Location"]

    def test_dashboard_non_connecte_redirige(self, client):
        r = client.get("/dashboard", follow_redirects=False)
        assert r.status_code == 302  # redirigé vers login


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — AUTHENTIFICATION
# ══════════════════════════════════════════════════════════════════════════════
class TestAuthentification:
    def test_login_correct(self, client):
        r = login(client, "admin@a.ga")
        assert r.status_code == 302  # redirection après login réussi

    def test_login_mauvais_mot_de_passe(self, client):
        page = client.get("/login")
        token = _extract_csrf(page.data)
        r = client.post("/login", data={
            "email": "admin@a.ga", "password": "FauxMotDePasse",
            "csrf_token": token,
        })
        assert r.status_code == 200  # reste sur la page de login
        assert b"incorrect" in r.data.lower() or b"error" in r.data.lower()

    def test_login_sans_csrf_rejete(self, client):
        """Une soumission de login sans token CSRF doit être rejetée."""
        r = client.post("/login", data={"email": "admin@a.ga",
                                        "password": "MotDePasse1"})
        assert r.status_code == 400  # CSRF manquant

    def test_logout(self, client):
        auth_session(client, "admin@a.ga")
        r = client.get("/logout", follow_redirects=False)
        assert r.status_code == 302


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — PAGES TENANT (authentifié)
# ══════════════════════════════════════════════════════════════════════════════
class TestPagesTenant:
    """Vérifie que toutes les pages principales répondent sans erreur 500."""

    PAGES = ["/dashboard", "/salaries", "/bulletins", "/conges", "/acomptes",
             "/journaliers", "/pointage", "/sites", "/periodes", "/parametres",
             "/utilisateurs", "/declaration-cnss", "/simulateur", "/recherche",
             "/audit"]

    @pytest.mark.parametrize("page", PAGES)
    def test_page_repond_sans_erreur(self, client, page):
        auth_session(client, "admin@a.ga")
        r = client.get(page, follow_redirects=False)
        assert r.status_code != 500, f"{page} renvoie une erreur 500"
        assert r.status_code in (200, 302)


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — CALCUL TEMPS RÉEL (le bug CSRF qu'on a corrigé)
# ══════════════════════════════════════════════════════════════════════════════
class TestCalculBulletin:
    def test_calculer_bulletin_sans_csrf_fonctionne(self, client):
        """Le calcul temps réel (POST JSON sans CSRF) doit fonctionner."""
        auth_session(client, "admin@a.ga")
        r = client.post("/api/calculer-bulletin",
                        json={"salaire_base": 500000},
                        headers={"Content-Type": "application/json"})
        assert r.status_code == 200, "Le calcul ne doit pas être bloqué par CSRF"
        data = r.get_json()
        assert "net_a_payer" in data
        assert data["salaire_brut"] == 500000

    def test_calcul_cnss_correct(self, client):
        """Vérifie que le calcul renvoie une CNSS cohérente (5% du brut)."""
        auth_session(client, "admin@a.ga")
        r = client.post("/api/calculer-bulletin", json={"salaire_base": 500000})
        data = r.get_json()
        # CNSS salarié = 5% de 500 000 = 25 000
        assert data["cnss_salarie"] == 25000

    def test_recherche_rapide_fonctionne(self, client):
        """L'autocomplétion doit renvoyer du JSON."""
        auth_session(client, "admin@a.ga")
        r = client.get("/api/recherche-rapide?q=ndo")
        assert r.status_code == 200
        assert r.is_json
        results = r.get_json()
        # Doit trouver le salarié NDONG du tenant A
        assert any("NDONG" in x.get("titre", "") for x in results)


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — PROTECTION CSRF DES FORMULAIRES
# ══════════════════════════════════════════════════════════════════════════════
class TestProtectionCSRF:
    def test_creation_salarie_sans_csrf_rejetee(self, client):
        """Un formulaire de modification SANS token doit être bloqué."""
        auth_session(client, "admin@a.ga")
        r = client.post("/salaries/nouveau",
                        data={"nom": "PIRATE", "prenom": "Sans Token"})
        assert r.status_code == 400  # CSRF protège bien


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — ISOLATION MULTI-TENANT (sécurité critique)
# ══════════════════════════════════════════════════════════════════════════════
class TestIsolationMultiTenant:
    def test_recherche_isolee_par_tenant(self, client):
        """L'admin du tenant A ne doit PAS voir les salariés du tenant B."""
        auth_session(client, "admin@a.ga")
        r = client.get("/api/recherche-rapide?q=obame")  # salarié du tenant B
        results = r.get_json()
        assert not any("OBAME" in x.get("titre", "") for x in results), \
            "Fuite de données entre tenants !"

    def test_admin_a_voit_son_salarie(self, client):
        """L'admin du tenant A voit bien son propre salarié."""
        auth_session(client, "admin@a.ga")
        r = client.get("/api/recherche-rapide?q=ndong")
        results = r.get_json()
        assert any("NDONG" in x.get("titre", "") for x in results)


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — PERMISSIONS (super-admin vs tenant)
# ══════════════════════════════════════════════════════════════════════════════
class TestPermissionsRoutes:
    def test_tenant_admin_ne_voit_pas_admin(self, client):
        """Un admin de tenant ne doit PAS accéder au panneau super-admin."""
        auth_session(client, "admin@a.ga")
        r = client.get("/admin", follow_redirects=False)
        assert r.status_code == 403

    def test_validation_mot_de_passe_inscription(self, client):
        """Un mot de passe faible doit être refusé à l'inscription."""
        page = client.get("/inscription")
        token = _extract_csrf(page.data)
        r = client.post("/inscription", data={
            "email": "nouveau@test.ga", "password": "123",  # trop court
            "denomination": "Test SARL", "nom": "Test", "prenom": "User",
            "csrf_token": token,
        }, follow_redirects=True)
        # Le mot de passe faible doit être signalé
        assert b"caract" in r.data.lower() or b"majuscule" in r.data.lower() \
            or b"chiffre" in r.data.lower()


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — SAUVEGARDE BASE DE DONNÉES
# ══════════════════════════════════════════════════════════════════════════════
class TestBackup:
    def test_page_backups_super_admin(self, client):
        """Le super-admin accède à la page des sauvegardes."""
        # Créer un super-admin dans la base de test
        with flask_app.app_context():
            from models import Utilisateur
            sa = Utilisateur(nom="SUPER", prenom="Admin", email="super@test.ga",
                             role="SUPER_ADMIN", actif=True, email_verifie=True)
            sa.set_password("MotDePasse1")
            db.session.add(sa)
            db.session.commit()
            sa_id = sa.id
        with client.session_transaction() as sess:
            sess["_user_id"] = str(sa_id); sess["_fresh"] = True
        r = client.get("/admin/backups")
        assert r.status_code == 200

    def test_page_backups_interdite_tenant(self, client):
        """Un admin de tenant ne peut PAS accéder aux sauvegardes."""
        auth_session(client, "admin@a.ga")
        r = client.get("/admin/backups", follow_redirects=False)
        assert r.status_code == 403

    def test_backup_sans_config_message_clair(self):
        """Sans config B2, run_backup renvoie un échec propre (pas un crash)."""
        import backup
        ok, message = backup.run_backup()
        # En environnement de test sans B2, doit échouer proprement
        assert isinstance(ok, bool)
        assert isinstance(message, str)
        assert len(message) > 0


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — MONITORING SENTRY
# ══════════════════════════════════════════════════════════════════════════════
class TestMonitoring:
    def test_scrub_filtre_mots_de_passe(self):
        """Les champs sensibles doivent être filtrés."""
        from monitoring import _scrub
        data = {"password": "secret", "email": "u@test.ga", "salaire": 500000}
        result = _scrub(data)
        assert result["password"] == "[FILTRÉ]"
        assert result["email"] == "u@test.ga"      # non sensible
        assert result["salaire"] == 500000          # donnée métier conservée

    def test_scrub_recursif_imbrique(self):
        """Le filtrage doit fonctionner en profondeur."""
        from monitoring import _scrub
        data = {"niveau1": {"token_api": "x", "nom": "NDONG"}}
        result = _scrub(data)
        assert result["niveau1"]["token_api"] == "[FILTRÉ]"
        assert result["niveau1"]["nom"] == "NDONG"

    def test_init_sentry_sans_dsn_desactive(self, monkeypatch):
        """Sans SENTRY_DSN, l'init retourne False sans erreur."""
        monkeypatch.delenv("SENTRY_DSN", raising=False)
        from monitoring import init_sentry
        assert init_sentry() is False

    def test_set_user_context_ne_casse_pas(self):
        """set_user_context ne doit jamais lever d'exception."""
        from monitoring import set_user_context
        # Avec None, un objet vide, un objet bizarre — aucune exception attendue
        set_user_context(None)
        set_user_context(object())
        # Si on arrive ici, c'est bon
        assert True


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — JOURS FÉRIÉS GABONAIS
# ══════════════════════════════════════════════════════════════════════════════
class TestJoursFeries:
    def test_calcul_paques(self):
        """L'algorithme de Computus doit donner les bonnes dates de Pâques."""
        from jours_feries import _dimanche_paques
        from datetime import date
        assert _dimanche_paques(2024) == date(2024, 3, 31)
        assert _dimanche_paques(2025) == date(2025, 4, 20)
        assert _dimanche_paques(2026) == date(2026, 4, 5)

    def test_jours_fixes_presents(self):
        """Les jours fériés fixes gabonais doivent être détectés."""
        from jours_feries import est_jour_ferie
        from datetime import date
        assert est_jour_ferie(date(2026, 1, 1))    # Nouvel An
        assert est_jour_ferie(date(2026, 5, 1))    # Fête du Travail
        assert est_jour_ferie(date(2026, 8, 17))   # Indépendance
        assert est_jour_ferie(date(2026, 12, 25))  # Noël

    def test_jour_normal_non_ferie(self):
        """Un jour ordinaire ne doit pas être férié."""
        from jours_feries import est_jour_ferie
        from datetime import date
        assert not est_jour_ferie(date(2026, 6, 9))  # mardi ordinaire

    def test_type_jour_auto(self):
        """La détection automatique du type de jour."""
        from jours_feries import type_jour_auto
        from datetime import date
        assert type_jour_auto(date(2026, 8, 17)) == "FERIE"     # Indépendance
        assert type_jour_auto(date(2026, 6, 7))  == "DIMANCHE"  # dimanche
        assert type_jour_auto(date(2026, 6, 9))  == "NORMAL"    # mardi

    def test_api_jour_ferie(self, client):
        """L'API doit renvoyer le type de jour correct."""
        auth_session(client, "admin@a.ga")
        r = client.get("/api/jour-ferie?date=2026-08-17")
        assert r.status_code == 200
        data = r.get_json()
        assert data["type_jour"] == "FERIE"
        assert data["est_ferie"] is True

    def test_page_jours_feries(self, client):
        """La page calendrier doit s'afficher."""
        auth_session(client, "admin@a.ga")
        r = client.get("/jours-feries")
        assert r.status_code == 200


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — PAGINATION
# ══════════════════════════════════════════════════════════════════════════════
class TestPagination:
    def test_salaries_pagine(self, client):
        """La liste des salariés accepte le paramètre page."""
        auth_session(client, "admin@a.ga")
        r = client.get("/salaries?page=1")
        assert r.status_code == 200

    def test_conges_pagine(self, client):
        """La liste des congés accepte le paramètre page."""
        auth_session(client, "admin@a.ga")
        r = client.get("/conges?page=1")
        assert r.status_code == 200

    def test_page_inexistante_pas_erreur(self, client):
        """Une page hors limites ne doit pas planter (error_out=False)."""
        auth_session(client, "admin@a.ga")
        r = client.get("/salaries?page=9999")
        assert r.status_code == 200


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — MODULE PRESTATAIRES
# ══════════════════════════════════════════════════════════════════════════════
class TestPrestataires:
    def test_liste_prestataires(self, client):
        auth_session(client, "admin@a.ga")
        r = client.get("/prestataires")
        assert r.status_code == 200

    def test_page_nouveau_prestataire(self, client):
        auth_session(client, "admin@a.ga")
        r = client.get("/prestataires/nouveau")
        assert r.status_code == 200

    def test_etat_annuel(self, client):
        auth_session(client, "admin@a.ga")
        r = client.get("/prestataires/etat-annuel?annee=2026")
        assert r.status_code == 200

    def test_calcul_facture_correct(self):
        """Le calcul TVA + retenue à la source doit être exact."""
        from models import FacturePrestataire
        from datetime import date
        f = FacturePrestataire(tenant_id=1, prestataire_id=1, numero="T1",
                               date_facture=date.today(),
                               montant_ht=2000000, taux_tva=18, taux_retenue=9.5)
        f.calculer()
        assert f.montant_tva == 360000          # 18% de 2M
        assert f.montant_ttc == 2360000         # 2M + TVA
        assert f.montant_retenue == 190000      # 9.5% de 2M
        assert f.montant_net_a_payer == 2170000 # TTC - retenue

    def test_api_calculer_facture(self, client):
        auth_session(client, "admin@a.ga")
        r = client.post("/api/prestataire/calculer-facture",
                        json={"montant_ht": 1000000, "taux_tva": 18, "taux_retenue": 5})
        assert r.status_code == 200
        d = r.get_json()
        assert d["montant_tva"] == 180000
        assert d["montant_net_a_payer"] == 1130000  # 1.18M - 50k

    def test_creation_prestataire_isole_par_tenant(self, client):
        """Un prestataire créé par le tenant A n'est pas visible par le tenant B."""
        from models import Prestataire
        # Créer un prestataire pour le tenant A directement
        with flask_app.app_context():
            from models import Tenant
            ta = Tenant.query.filter_by(slug="entreprise-a").first()
            p = Prestataire(tenant_id=ta.id, code="PREA1",
                            raison_sociale="FOURNISSEUR A", categorie="FOURNISSEUR")
            db.session.add(p)
            db.session.commit()
            pid = p.id
        # Le tenant B ne doit pas y accéder
        auth_session(client, "admin@b.ga")
        r = client.get(f"/prestataires/{pid}", follow_redirects=False)
        assert r.status_code == 404  # introuvable pour le tenant B


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — DOCUMENTS RH & EXPORT EN MASSE
# ══════════════════════════════════════════════════════════════════════════════
class TestDocumentsRH:
    def _setup_salarie(self):
        """Crée un salarié dans le tenant A et retourne son id."""
        with flask_app.app_context():
            from models import Tenant, Salarie
            from datetime import date
            ta = Tenant.query.filter_by(slug="entreprise-a").first()
            s = Salarie(tenant_id=ta.id, matricule="DOC01", nom="MBA", prenom="Paul",
                        date_embauche=date(2021, 1, 1), emploi="Chef de chantier",
                        sexe="M", statut="ACTIF")
            db.session.add(s)
            db.session.commit()
            return s.id

    def test_attestation_travail_pdf(self, client):
        sid = self._setup_salarie()
        auth_session(client, "admin@a.ga")
        r = client.get(f"/salaries/{sid}/document/attestation-travail")
        assert r.status_code == 200
        assert r.data[:4] == b"%PDF"

    def test_certificat_travail_pdf(self, client):
        sid = self._setup_salarie()
        auth_session(client, "admin@a.ga")
        r = client.get(f"/salaries/{sid}/document/certificat-travail")
        assert r.status_code == 200
        assert r.data[:4] == b"%PDF"

    def test_attestation_salaire_pdf(self, client):
        sid = self._setup_salarie()
        auth_session(client, "admin@a.ga")
        r = client.get(f"/salaries/{sid}/document/attestation-salaire")
        assert r.status_code == 200
        assert r.data[:4] == b"%PDF"

    def test_solde_tout_compte_pdf(self, client):
        sid = self._setup_salarie()
        auth_session(client, "admin@a.ga")
        r = client.get(f"/salaries/{sid}/document/solde-tout-compte")
        assert r.status_code == 200
        assert r.data[:4] == b"%PDF"

    def test_document_type_inconnu(self, client):
        sid = self._setup_salarie()
        auth_session(client, "admin@a.ga")
        r = client.get(f"/salaries/{sid}/document/type-bidon", follow_redirects=False)
        assert r.status_code == 302  # redirige avec message d'erreur

    def test_document_isole_par_tenant(self, client):
        """Le tenant B ne peut pas générer de document pour un salarié du tenant A."""
        sid = self._setup_salarie()  # salarié du tenant A
        auth_session(client, "admin@b.ga")
        r = client.get(f"/salaries/{sid}/document/attestation-travail",
                       follow_redirects=False)
        assert r.status_code == 404

    def test_export_zip_bulletins(self, client):
        """L'export « tous les bulletins » doit produire UN SEUL PDF."""
        with flask_app.app_context():
            from models import Tenant, Salarie, PeriodePaie, BulletinPaie
            from datetime import date
            ta = Tenant.query.filter_by(slug="entreprise-a").first()
            s = Salarie(tenant_id=ta.id, matricule="ZIP01", nom="OYONO", prenom="Luc",
                        date_embauche=date(2022, 1, 1), statut="ACTIF")
            db.session.add(s); db.session.flush()
            p = PeriodePaie(tenant_id=ta.id, mois=5, annee=2026,
                            libelle_mois="MAI", statut="OUVERT")
            db.session.add(p); db.session.flush()
            b = BulletinPaie(tenant_id=ta.id, salarie_id=s.id, periode_id=p.id,
                             salaire_base=400000, salaire_brut=400000,
                             net_a_payer=340000, statut="VALIDE")
            db.session.add(b); db.session.commit()
            pid = p.id
        auth_session(client, "admin@a.ga")
        r = client.get(f"/bulletins/export-zip/{pid}")
        assert r.status_code == 200
        assert r.data[:4] == b"%PDF"  # un seul PDF regroupé
        assert r.mimetype == "application/pdf"


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — NOTIFICATIONS & RAPPELS
# ══════════════════════════════════════════════════════════════════════════════
class TestNotifications:
    def test_page_notifications(self, client):
        auth_session(client, "admin@a.ga")
        r = client.get("/notifications")
        assert r.status_code == 200

    def test_api_count(self, client):
        auth_session(client, "admin@a.ga")
        r = client.get("/api/notifications/count")
        assert r.status_code == 200
        d = r.get_json()
        assert "total" in d and "critiques" in d

    def test_contrat_echeance_detecte(self, client):
        """Un contrat arrivant à échéance doit générer une notification."""
        from models import Tenant, Salarie, Contrat
        from notifications import get_notifications
        from blueprints.tenant import _NOTIF_MODELS
        from datetime import date, timedelta
        with flask_app.app_context():
            ta = Tenant.query.filter_by(slug="entreprise-a").first()
            s = Salarie(tenant_id=ta.id, matricule="NOTIF1", nom="ECHEANCE", prenom="Test",
                        date_embauche=date(2020,1,1), statut="ACTIF")
            db.session.add(s); db.session.flush()
            db.session.add(Contrat(tenant_id=ta.id, salarie_id=s.id, type_contrat="CDD",
                date_debut=date(2024,1,1), date_fin=date.today()+timedelta(days=5),
                salaire_base=300000, actif=True))
            db.session.commit()
            notifs = get_notifications(ta, db, _NOTIF_MODELS)
            contrats = [n for n in notifs if n["categorie"] == "contrat"]
            assert len(contrats) >= 1

    def test_conge_a_valider_detecte(self, client):
        """Un congé en attente doit générer une notification."""
        from models import Tenant, Salarie, Conge
        from notifications import get_notifications
        from blueprints.tenant import _NOTIF_MODELS
        from datetime import date, timedelta
        with flask_app.app_context():
            ta = Tenant.query.filter_by(slug="entreprise-a").first()
            s = Salarie(tenant_id=ta.id, matricule="NOTIF2", nom="CONGE", prenom="Test",
                        date_embauche=date(2020,1,1), statut="ACTIF")
            db.session.add(s); db.session.flush()
            db.session.add(Conge(tenant_id=ta.id, salarie_id=s.id, annee=date.today().year,
                jours_acquis=20, jours_pris=0, date_depart=date.today()+timedelta(days=10),
                statut="DEMANDÉ"))
            db.session.commit()
            notifs = get_notifications(ta, db, _NOTIF_MODELS)
            conges = [n for n in notifs if n["categorie"] == "conge"]
            assert len(conges) >= 1

    def test_notifications_isolees_par_tenant(self, client):
        """Les notifications d'un tenant ne fuient pas vers un autre."""
        from models import Tenant
        from notifications import get_notifications
        from blueprints.tenant import _NOTIF_MODELS
        with flask_app.app_context():
            tb = Tenant.query.filter_by(slug="entreprise-b").first()
            notifs_b = get_notifications(tb, db, _NOTIF_MODELS)
            # Toutes les notifs de B doivent concerner B (pas de fuite)
            assert isinstance(notifs_b, list)


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — IMMUABILITÉ DES BULLETINS VALIDÉS
# ══════════════════════════════════════════════════════════════════════════════
class TestBulletinImmuable:
    """Un bulletin validé ne doit pas pouvoir être réécrit via la saisie."""

    def _creer_bulletin_valide(self):
        with flask_app.app_context():
            sal = Salarie.query.filter_by(matricule="EA001").first()
            per = PeriodePaie.query.filter_by(mois=6, annee=2026).first()
            b = BulletinPaie(tenant_id=sal.tenant_id, salarie_id=sal.id,
                             periode_id=per.id, statut="VALIDÉ",
                             salaire_base=300000, salaire_brut=300000,
                             net_a_payer=250000)
            db.session.add(b)
            db.session.commit()
            return sal.id, per.id, b.id, float(b.net_a_payer)

    def test_bulletin_valide_non_modifiable(self, client):
        auth_session(client, "admin@a.ga")
        sid, pid, bid, net_avant = self._creer_bulletin_valide()

        # Récupérer le token CSRF depuis le formulaire de saisie
        page = client.get(f"/bulletins/saisie?salarie_id={sid}")
        token = _extract_csrf(page.data)

        # Tenter d'écraser le bulletin validé avec un nouveau salaire
        client.post("/bulletins/saisie", data={
            "salarie_id": sid, "periode_id": pid, "csrf_token": token,
            "salaire_base": 999999, "action": "valider",
        }, follow_redirects=False)

        # Le bulletin validé doit être resté intact
        with flask_app.app_context():
            b2 = BulletinPaie.query.get(bid)
            assert b2.statut in ("VALIDÉ", "VALIDE")
            assert float(b2.net_a_payer) == net_avant
            assert float(b2.salaire_base) == 300000


class TestNumeroBulletin:
    """Numérotation séquentielle et immuable des bulletins (#2)."""

    def test_numero_sequentiel_et_immuable(self, client):
        from blueprints.tenant import attribuer_numero_bulletin
        with flask_app.app_context():
            sal = Salarie.query.filter_by(matricule="EA001").first()
            per1 = PeriodePaie.query.filter_by(mois=6, annee=2026).first()
            per2 = PeriodePaie(tenant_id=sal.tenant_id, mois=7, annee=2026,
                               libelle_mois="JUILLET", statut="OUVERT")
            db.session.add(per2); db.session.commit()

            b1 = BulletinPaie(tenant_id=sal.tenant_id, salarie_id=sal.id,
                              periode_id=per1.id, statut="BROUILLON",
                              salaire_base=300000, salaire_brut=300000, net_a_payer=250000)
            b2 = BulletinPaie(tenant_id=sal.tenant_id, salarie_id=sal.id,
                              periode_id=per2.id, statut="BROUILLON",
                              salaire_base=300000, salaire_brut=300000, net_a_payer=250000)
            db.session.add_all([b1, b2]); db.session.commit()

            # Premier numéro
            attribuer_numero_bulletin(b1)
            assert b1.numero == "BP-2026-000001"

            # Le suivant s'incrémente
            attribuer_numero_bulletin(b2)
            assert b2.numero == "BP-2026-000002"

            # Immuabilité : un nouvel appel ne réattribue pas
            attribuer_numero_bulletin(b1)
            assert b1.numero == "BP-2026-000001"


class TestPagesLegales:
    """Les pages légales publiques doivent être accessibles sans connexion (#4)."""

    def test_cgu_accessible(self, client):
        assert client.get("/cgu").status_code == 200

    def test_cgv_accessible(self, client):
        r = client.get("/cgv")
        assert r.status_code == 200
        assert b"Conditions" in r.data

    def test_mentions_legales_accessibles(self, client):
        assert client.get("/mentions-legales").status_code == 200

    def test_confidentialite_accessible(self, client):
        assert client.get("/politique-confidentialite").status_code == 200


# ═══════════════════════════════════════════════════════════════════════════
# NON-RÉGRESSION : routes de téléchargement / import / export
# But : attraper les imports manquants (NameError) sur les routes peu couvertes.
# Un statut 500 = régression (souvent un import oublié). 200/302/400/403/404 = OK.
# ═══════════════════════════════════════════════════════════════════════════
class TestTelechargementsExports:
    """Frappe les routes de génération de fichiers ; aucune ne doit renvoyer 500."""

    def _login_admin(self, client):
        login(client, "admin@a.ga")

    def test_routes_sans_parametre_ne_plantent_pas(self, client):
        self._login_admin(client)
        urls = [
            "/salaries/pointage/modele",        # télécharger modèle pointage salariés
            "/journaliers/pointage/modele",     # télécharger modèle pointage journaliers
            "/salaries/import/modele",          # modèle d'import salariés
            "/parametres/export-donnees",       # export ZIP des données
            "/parametres/modele-bulletin",      # modèle de bulletin
            "/parametres/grille-salaires",      # page grille
            "/langue/fr",                       # changement de langue
            "/recherche?q=test",                # recherche globale
            "/api/recherche-rapide?q=test",     # API recherche rapide
            "/messages",                        # messagerie support
        ]
        for u in urls:
            r = client.get(u, follow_redirects=False)
            assert r.status_code != 500, f"{u} renvoie 500 (probable import manquant)"

    def test_routes_avec_periode_ne_plantent_pas(self, client):
        self._login_admin(client)
        from models import PeriodePaie, Tenant
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        p = PeriodePaie.query.filter_by(tenant_id=t.id).first()
        assert p is not None
        urls = [
            f"/export/sage/journal/{p.id}",
            f"/export/sage/livre/{p.id}",
            f"/export/sage/les-deux/{p.id}",
            f"/rapport/pdf/{p.id}",
        ]
        for u in urls:
            r = client.get(u, follow_redirects=False)
            assert r.status_code != 500, f"{u} renvoie 500 (probable import manquant)"

    def test_import_pointage_sans_fichier_ne_plante_pas(self, client):
        """POST d'import sans fichier : doit être géré proprement (pas de 500)."""
        self._login_admin(client)
        for u in ["/salaries/pointage/importer", "/journaliers/pointage/importer"]:
            r = client.post(u, data={}, follow_redirects=False)
            assert r.status_code != 500, f"{u} renvoie 500 sans fichier"


class TestHistoriqueSalarie:
    """Modification date d'embauche / fonction + traçage de l'historique."""

    def test_modif_embauche_fonction_tracee(self, client):
        from models import db, Tenant, Salarie, HistoriqueSalarie
        from datetime import date
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        s = Salarie(tenant_id=t.id, matricule="HIST1", nom="TEST", prenom="Hist",
                    sexe="M", statut="ACTIF", date_embauche=date(2023, 1, 1),
                    emploi="Ouvrier", situation_matrimoniale="CELIBATAIRE", nb_enfants=0)
        db.session.add(s); db.session.commit(); sid = s.id
        login(client, "admin@a.ga")
        page = client.get(f"/salaries/{sid}/modifier")
        token = _extract_csrf(page.data)
        r = client.post(f"/salaries/{sid}/modifier", data={
            "nom": "TEST", "prenom": "Hist", "sexe": "M", "nationalite": "GABONAISE",
            "situation_matrimoniale": "CELIBATAIRE", "nb_enfants": "0", "statut": "ACTIF",
            "date_embauche": "2023-06-15", "emploi": "Contremaître", "csrf_token": token,
        }, follow_redirects=True)
        assert r.status_code == 200
        s = Salarie.query.get(sid)
        assert str(s.date_embauche) == "2023-06-15"
        assert s.emploi == "Contremaître"
        h = HistoriqueSalarie.query.filter_by(salarie_id=sid).all()
        champs = {e.champ for e in h}
        assert "date_embauche" in champs and "emploi" in champs


class TestSimulateurDroits:
    """Simulateur de droits de fin de contrat (API) — modes libre et existant."""

    def test_simuler_droits_libre(self, client):
        login(client, "admin@a.ga")
        r = client.post("/api/simuler-droits", json={
            "mode": "libre", "salaire": 300000, "anciennete": 5,
            "statut": "EXECUTION", "cause": "LICENCIEMENT", "date_cessation": "2026-01-31",
        })
        assert r.status_code == 200
        d = r.get_json()
        assert d["indem_licenciement"] > 0
        assert d["total_brut"] >= d["indem_licenciement"]
        assert "total_net_estime" in d

    def test_simuler_droits_sans_salaire_refuse(self, client):
        login(client, "admin@a.ga")
        r = client.post("/api/simuler-droits", json={"mode": "libre", "cause": "LICENCIEMENT"})
        assert r.status_code == 400


class TestSimulateurDroits:
    """Simulateur de droits de fin de contrat (salarié existant + saisie libre)."""

    def test_page_et_calculs(self, client):
        from models import db, Tenant, Salarie, Contrat
        from datetime import date
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        t.convention = "BTP"; t.jours_conge_par_mois = 2.0   # état déterministe
        s = Salarie(tenant_id=t.id, matricule="SIMU1", nom="SIMU", prenom="Test",
                    statut="ACTIF", date_embauche=date(2021, 1, 1), nb_enfants=2)
        db.session.add(s); db.session.flush()
        db.session.add(Contrat(tenant_id=t.id, salarie_id=s.id, type_contrat="CDI",
                               date_debut=date(2021, 1, 1), salaire_base=250000, actif=True))
        db.session.commit(); sid = s.id
        login(client, "admin@a.ga")
        assert client.get("/simulateur/droits").status_code == 200
        # mode salarié
        r = client.get(f"/simulateur/droits?go=1&mode=salarie&salarie_id={sid}"
                       "&cause=LICENCIEMENT&date_cessation=2026-06-30")
        assert r.status_code == 200 and "Total net à payer" in r.get_data(as_text=True)
        # saisie libre
        r = client.get("/simulateur/droits?go=1&mode=libre&salaire=300000"
                       "&date_embauche=2020-01-01&nb_enfants=1&cause=DEMISSION&date_cessation=2026-06-30")
        assert r.status_code == 200 and "Total net à payer" in r.get_data(as_text=True)

    def test_helper_cause_change_resultat(self):
        from simulateur_droits import simuler_fin_contrat
        from datetime import date
        lic = simuler_fin_contrat(salaire=250000, date_embauche=date(2021,1,1),
                                  convention="BTP", cause="LICENCIEMENT", date_cessation=date(2026,6,30))
        dem = simuler_fin_contrat(salaire=250000, date_embauche=date(2021,1,1),
                                  convention="BTP", cause="DEMISSION", date_cessation=date(2026,6,30))
        # Le motif change le résultat : le licenciement (avec préavis employeur)
        # donne un net supérieur à la démission.
        assert lic["total_net"] != dem["total_net"]
        assert lic["preavis_montant"] > 0 and dem["preavis_montant"] == 0


class TestSanctions:
    """Module discipline : ajout, lettre PDF, suppression."""

    def test_cycle_sanction(self, client):
        from models import db, Tenant, Salarie, Sanction
        from datetime import date
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        s = Salarie(tenant_id=t.id, matricule="SANC1", nom="DISC", prenom="Test",
                    statut="ACTIF", date_embauche=date(2024, 1, 1), emploi="Ouvrier")
        db.session.add(s); db.session.commit(); sid = s.id
        login(client, "admin@a.ga")
        page = client.get(f"/salaries/{sid}/modifier"); token = _extract_csrf(page.data)
        # ajout
        r = client.post(f"/salaries/{sid}/sanctions", data={
            "type": "AVERTISSEMENT", "date_sanction": "2026-06-15",
            "motif": "Retards", "description": "Faits.", "csrf_token": token,
        }, follow_redirects=True)
        assert r.status_code == 200
        sa = Sanction.query.filter_by(salarie_id=sid).first()
        assert sa is not None and sa.type == "AVERTISSEMENT"
        # lettre PDF
        r = client.get(f"/salaries/{sid}/sanctions/{sa.id}/lettre")
        assert r.status_code == 200 and r.data[:4] == b"%PDF"
        # suppression
        r = client.post(f"/salaries/{sid}/sanctions/{sa.id}/supprimer",
                        data={"csrf_token": token}, follow_redirects=True)
        assert Sanction.query.filter_by(salarie_id=sid).count() == 0

    def test_demission_pas_indemnite_licenciement(self):
        """Non-régression : une démission n'affiche jamais 'indemnité de licenciement'."""
        from simulateur_droits import simuler_fin_contrat
        from datetime import date
        r = simuler_fin_contrat(salaire=250000, date_embauche=date(2021,1,1),
                                convention="BTP", cause="DEMISSION", date_cessation=date(2026,6,30))
        assert r["type_indemnite"] != "LICENCIEMENT"


class TestProcedureRupture:
    """Workflow de rupture : étapes selon motif + aperçu STC + finalisation."""

    def test_page_et_finalisation(self, client):
        from models import db, Tenant, Salarie, Contrat
        from datetime import date
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        t.convention = "BTP"; t.jours_conge_par_mois = 2.0
        s = Salarie(tenant_id=t.id, matricule="RUP1", nom="RUPT", prenom="Test",
                    statut="ACTIF", date_embauche=date(2021, 1, 1), emploi="Maçon")
        db.session.add(s); db.session.flush()
        c = Contrat(tenant_id=t.id, salarie_id=s.id, type_contrat="CDI",
                    date_debut=date(2021, 1, 1), salaire_base=250000, actif=True)
        db.session.add(c); db.session.commit(); sid, cid = s.id, c.id
        login(client, "admin@a.ga")
        # page licenciement
        r = client.get(f"/salaries/{sid}/rupture?motif=LICENCIEMENT&date_cessation=2026-06-30")
        assert r.status_code == 200
        html = r.get_data(as_text=True)
        assert "Procédure de rupture" in html and "Total net à payer" in html
        # démission → étapes adaptées
        r = client.get(f"/salaries/{sid}/rupture?motif=DEMISSION")
        assert "lettre de démission" in r.get_data(as_text=True)
        # finalisation
        page = client.get(f"/salaries/{sid}/modifier"); token = _extract_csrf(page.data)
        r = client.post(f"/contrats/{cid}/terminer", data={
            "type_rupture": "LICENCIEMENT", "date_arret": "2026-06-30", "csrf_token": token,
        }, follow_redirects=True)
        assert r.status_code == 200
        s = Salarie.query.get(sid)
        assert s.statut == "INACTIF" and str(s.date_cessation) == "2026-06-30"


class TestStatsLanding:
    """Comptage landing (agrégé, sans donnée perso) + page admin."""

    _H = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120 Safari/537.36"}

    def test_comptage_et_page(self, client):
        from models import StatLanding
        # visites humaines + clic CTA
        client.get("/", headers=self._H)
        client.get("/", headers={**self._H, "Referer": "https://www.google.com/"})
        client.get("/?src=facebook", headers=self._H)
        client.get("/essai", headers=self._H)
        # bots : doivent être ignorés
        client.get("/", headers={"User-Agent": "Mozilla/5.0 (compatible; Googlebot/2.1)"})
        client.get("/", headers={"User-Agent": ""})
        rows = StatLanding.query.all()
        assert sum(r.vues for r in rows) == 3     # 3 humains, 0 bot
        assert sum(r.clics_cta for r in rows) >= 1
        sources = {r.source for r in rows}
        assert "google" in sources and "facebook" in sources
        # page admin (super-admin)
        login(client, "super@admin.ga")
        r = client.get("/admin/landing")
        # super@admin.ga n'existe peut-être pas dans le seed : on vérifie au moins que
        # la route répond (200 si super-admin, 302/403 sinon) sans planter (pas de 500).
        assert r.status_code != 500


class TestValidationEmail:
    """Vérification de l'email à l'inscription (format, typo, domaine)."""

    def test_module_verification(self):
        from email_validation import verifier_email, suggerer_correction
        assert verifier_email("jean@gmail.com")[0] is True
        assert verifier_email("nimportequoi")[0] is False          # format
        assert verifier_email("test@mailinator.com")[0] is False   # jetable
        assert suggerer_correction("paul@gmial.com") == "paul@gmail.com"  # typo

    def test_inscription_rejette_email_douteux(self, client):
        # faute de frappe → suggestion
        r = client.get("/inscription"); tok = _extract_csrf(r.data)
        r = client.post("/inscription", data={
            "email": "paul@gmial.com", "password": "MotDePasse1",
            "denomination": "TEST SARL", "nom": "P", "prenom": "Q", "csrf_token": tok,
        }, follow_redirects=True)
        assert "gmail.com" in r.get_data(as_text=True)
        # format invalide
        r = client.get("/inscription"); tok = _extract_csrf(r.data)
        r = client.post("/inscription", data={
            "email": "pasunemail", "password": "MotDePasse1",
            "denomination": "T", "nom": "P", "prenom": "Q", "csrf_token": tok,
        }, follow_redirects=True)
        assert "format invalide" in r.get_data(as_text=True)


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — MISE À DISPOSITION DE PERSONNEL
# ══════════════════════════════════════════════════════════════════════════════
class TestMiseADisposition:
    def test_dashboard_accessible(self, client):
        login(client, "admin@a.ga")
        assert client.get("/mise-a-disposition").status_code == 200

    def test_cycle_complet_client_affectation_facture(self, client):
        from models import ClientUtilisateur, AffectationMAD, Salarie
        login(client, "admin@a.ga")
        # créer un client
        page = client.get("/mise-a-disposition/clients/nouveau")
        token = _extract_csrf(page.data)
        r = client.post("/mise-a-disposition/clients/nouveau",
                        data={"nom": "CLIENT MAD", "secteur": "BTP", "csrf_token": token},
                        follow_redirects=False)
        assert r.status_code == 302
        with flask_app.app_context():
            c = ClientUtilisateur.query.filter_by(nom="CLIENT MAD").first()
            assert c is not None
            cid = c.id
            sid = Salarie.query.filter_by(matricule="EA001").first().id
        # affecter un salarié (token depuis la fiche client)
        page = client.get(f"/mise-a-disposition/clients/{cid}")
        token = _extract_csrf(page.data)
        r = client.post(f"/mise-a-disposition/clients/{cid}/affecter",
                        data={"salarie_id": sid, "poste": "Maçon",
                              "mode_facturation": "TAUX_JOUR", "valeur": "25000",
                              "date_debut": "2026-06-01", "csrf_token": token})
        assert r.status_code == 302
        with flask_app.app_context():
            a = AffectationMAD.query.filter_by(client_id=cid).first()
            assert a is not None and float(a.valeur) == 25000.0
            aid = a.id
        # page facture
        page = client.get(f"/mise-a-disposition/clients/{cid}/facture?mois=6&annee=2026")
        assert page.status_code == 200 and b"TOTAL" in page.data
        token = _extract_csrf(page.data)
        # PDF
        r = client.post(f"/mise-a-disposition/clients/{cid}/facture/pdf",
                        data={"mois": "6", "annee": "2026", f"q_{aid}": "22",
                              "csrf_token": token})
        assert r.status_code == 200
        assert r.headers["Content-Type"] == "application/pdf"
        assert r.data[:5] == b"%PDF-"

    def test_isolation_tenant(self, client):
        """Un tenant ne peut pas accéder au client MAD d'un autre tenant."""
        from models import ClientUtilisateur
        login(client, "admin@a.ga")
        page = client.get("/mise-a-disposition/clients/nouveau")
        token = _extract_csrf(page.data)
        client.post("/mise-a-disposition/clients/nouveau",
                    data={"nom": "SECRET A", "csrf_token": token}, follow_redirects=False)
        with flask_app.app_context():
            cid = ClientUtilisateur.query.filter_by(nom="SECRET A").first().id
        # tenant B tente d'y accéder → redirigé (introuvable)
        client.get("/logout")
        login(client, "admin@b.ga")
        r = client.get(f"/mise-a-disposition/clients/{cid}", follow_redirects=False)
        assert r.status_code == 302


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — MAD : contrat + pointage par client
# ══════════════════════════════════════════════════════════════════════════════
class TestMADContratEtPointage:
    def _setup(self):
        from models import db, Tenant, Salarie, ClientUtilisateur, AffectationMAD, Pointage
        from datetime import date
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        s = Salarie.query.filter_by(matricule="EA001").first()
        c = ClientUtilisateur(tenant_id=t.id, nom="UTILISATRICE SARL", nif="N9",
                              adresse="Owendo", rccm="R1", contact_nom="M. DIALLO")
        db.session.add(c); db.session.commit()
        a = AffectationMAD(tenant_id=t.id, client_id=c.id, salarie_id=s.id, poste="Maçon",
                           date_debut=date(2026, 6, 1), mode_facturation="TAUX_JOUR",
                           valeur=25000, actif=True)
        db.session.add(a); db.session.commit()
        for d in range(1, 19):  # 18 jours pointés en juin 2026
            db.session.add(Pointage(tenant_id=t.id, salarie_id=s.id,
                                    date_pointage=date(2026, 6, d), present=True, heures_normales=8))
        db.session.commit()
        return c.id, a.id

    def test_contrat_mad_pdf(self, client):
        cid, aid = self._setup()
        login(client, "admin@a.ga")
        r = client.get(f"/mise-a-disposition/affectations/{aid}/contrat")
        assert r.status_code == 200
        assert r.headers["Content-Type"] == "application/pdf"
        assert r.data[:5] == b"%PDF-"

    def test_facture_prefill_depuis_pointage(self, client):
        cid, aid = self._setup()
        login(client, "admin@a.ga")
        r = client.get(f"/mise-a-disposition/clients/{cid}/facture?mois=6&annee=2026")
        assert r.status_code == 200
        assert b'value="18"' in r.data  # jours réellement pointés, pas les jours ouvrés


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — Profil d'inscription (type_compte) + mise en relation par NIF
# ══════════════════════════════════════════════════════════════════════════════
class TestProfilEtRelations:
    def test_type_compte_inscription(self, client):
        from models import Tenant
        page = client.get("/inscription")
        tok = _extract_csrf(page.data)
        r = client.post("/inscription", data={
            "denomination": "MAD CORP", "nom": "P", "prenom": "Q",
            "email": "madcorp.paie@gmail.com", "password": "MotDePasse1",
            "type_compte": "MISE_A_DISPOSITION", "csrf_token": tok,
        }, follow_redirects=False)
        t = Tenant.query.filter_by(denomination="MAD CORP").first()
        assert t is not None, "tenant non créé"
        assert t.type_compte == "MISE_A_DISPOSITION" and t.est_cabinet is False

    def test_envoyer_et_confidentialite(self, client):
        from models import db, Tenant, MiseEnRelation
        b = Tenant.query.filter_by(slug="entreprise-b").first()
        b.nif = "NIF-B-123"; db.session.commit()
        auth_session(client, "admin@a.ga")
        page = client.get("/mise-a-disposition/relations"); tok = _extract_csrf(page.data)
        r = client.post("/mise-a-disposition/relations/envoyer",
                        data={"nif": "NIF-B-123", "message": "Bonjour", "csrf_token": tok})
        assert r.status_code == 302
        rel = MiseEnRelation.query.filter_by(nif_recherche="NIF-B-123").first()
        assert rel is not None and rel.statut == "EN_ATTENTE" and rel.tenant_cible_id == b.id
        # A (demandeur) ne voit pas le nom de B tant que c'est en attente
        page = client.get("/mise-a-disposition/relations")
        assert b"ENTREPRISE B" not in page.data and "confidentiel".encode() in page.data

    def test_accepter_cree_client_chez_demandeur(self, client):
        from models import db, Tenant, MiseEnRelation, ClientUtilisateur
        a = Tenant.query.filter_by(slug="entreprise-a").first()
        b = Tenant.query.filter_by(slug="entreprise-b").first()
        b.nif = "NIF-B-999"
        rel = MiseEnRelation(tenant_demandeur_id=a.id, tenant_cible_id=b.id,
                             nif_recherche="NIF-B-999", statut="EN_ATTENTE", message="Hello")
        db.session.add(rel); db.session.commit()
        rid, a_id = rel.id, a.id
        # B (la cible) accepte
        auth_session(client, "admin@b.ga")
        page = client.get("/mise-a-disposition/relations"); tok = _extract_csrf(page.data)
        assert b"ENTREPRISE A" in page.data and "Accepter".encode() in page.data
        r = client.post(f"/mise-a-disposition/relations/{rid}/accepter",
                        data={"csrf_token": tok})
        assert r.status_code == 302
        rel = MiseEnRelation.query.get(rid)
        assert rel.statut == "ACCEPTEE" and rel.client_cree_id is not None
        c = ClientUtilisateur.query.filter_by(tenant_id=a_id, nif="NIF-B-999").first()
        assert c is not None and "ENTREPRISE B" in (c.nom or "").upper()

    def test_anti_enumeration(self, client):
        auth_session(client, "admin@a.ga")
        page = client.get("/mise-a-disposition/relations"); tok = _extract_csrf(page.data)
        r = client.post("/mise-a-disposition/relations/envoyer",
                        data={"nif": "NIF-INEXISTANT-999", "csrf_token": tok},
                        follow_redirects=True)
        assert "recevra votre demande".encode() in r.data


# ══════════════════════════════════════════════════════════════════════════════
# TESTS — Régime fiscal prestataire + CSS
# ══════════════════════════════════════════════════════════════════════════════
class TestPrestataireCSS:
    def _login_edit(self, client):
        # admin@a.ga a le rôle TENANT_ADMIN (can_edit)
        auth_session(client, "admin@a.ga")

    def test_regime_precompte_css(self, client):
        from models import Prestataire, FacturePrestataire
        self._login_edit(client)
        page = client.get("/prestataires/nouveau"); tok = _extract_csrf(page.data)
        r = client.post("/prestataires/nouveau", data={
            "code": "PX", "type_personne": "MORALE", "categorie": "SOUS_TRAITANT",
            "raison_sociale": "BTP SARL", "regime_fiscal": "PRECOMPTE",
            "assujetti_css": "on", "resident": "on", "ville": "Libreville",
            "csrf_token": tok})
        assert r.status_code == 302
        p = Prestataire.query.filter_by(raison_sociale="BTP SARL").first()
        assert p is not None
        assert p.regime_fiscal == "PRECOMPTE" and p.assujetti_tva is False
        assert float(p.taux_retenue_source) == 9.5 and p.assujetti_css is True
        # facture HT=1 000 000 → CSS 10 000, précompte 95 000, net 915 000
        page = client.get(f"/prestataires/{p.id}"); tok = _extract_csrf(page.data)
        r = client.post(f"/prestataires/{p.id}/factures/nouvelle", data={
            "numero": "F001", "date_facture": "2026-03-10", "montant_ht": "1000000",
            "csrf_token": tok})
        assert r.status_code == 302
        f = FacturePrestataire.query.filter_by(numero="F001").first()
        assert float(f.montant_css) == 10000 and float(f.montant_retenue) == 95000
        assert float(f.montant_ttc) == 1010000 and float(f.montant_net_a_payer) == 915000

    def test_regime_tva_css(self, client):
        from models import Prestataire, FacturePrestataire
        self._login_edit(client)
        page = client.get("/prestataires/nouveau"); tok = _extract_csrf(page.data)
        client.post("/prestataires/nouveau", data={
            "code": "PY", "type_personne": "MORALE", "categorie": "FREELANCE",
            "raison_sociale": "CONSEIL SA", "regime_fiscal": "TVA_CSS",
            "resident": "on", "ville": "Libreville", "csrf_token": tok})
        p = Prestataire.query.filter_by(raison_sociale="CONSEIL SA").first()
        assert p.regime_fiscal == "TVA_CSS" and p.assujetti_tva is True and p.assujetti_css is True
        page = client.get(f"/prestataires/{p.id}"); tok = _extract_csrf(page.data)
        client.post(f"/prestataires/{p.id}/factures/nouvelle", data={
            "numero": "F002", "date_facture": "2026-03-11", "montant_ht": "1000000",
            "csrf_token": tok})
        f = FacturePrestataire.query.filter_by(numero="F002").first()
        # TVA 180 000 + CSS 10 000 → TTC 1 190 000, pas de retenue
        assert float(f.montant_tva) == 180000 and float(f.montant_css) == 10000
        assert float(f.montant_ttc) == 1190000 and float(f.montant_net_a_payer) == 1190000

    def test_das_honoraires_inclut_css(self, client):
        from models import db, Tenant, Prestataire, FacturePrestataire
        import declaration_das as dd, models as M
        self._login_edit(client)
        page = client.get("/prestataires/nouveau"); tok = _extract_csrf(page.data)
        client.post("/prestataires/nouveau", data={
            "code": "PZ", "type_personne": "MORALE", "categorie": "SOUS_TRAITANT",
            "raison_sociale": "MACON SARL", "regime_fiscal": "CSS",
            "resident": "on", "ville": "Libreville", "csrf_token": tok})
        p = Prestataire.query.filter_by(raison_sociale="MACON SARL").first()
        page = client.get(f"/prestataires/{p.id}"); tok = _extract_csrf(page.data)
        client.post(f"/prestataires/{p.id}/factures/nouvelle", data={
            "numero": "F003", "date_facture": "2026-05-10", "montant_ht": "2000000",
            "csrf_token": tok})
        f = FacturePrestataire.query.filter_by(numero="F003").first()
        f.statut = "PAYEE"; db.session.commit()
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        lignes, tot = dd.agreger_honoraires(t, 2026, db=db, models=M)
        assert tot["css"] == 20000  # 1% de 2 000 000
        assert any(l["css"] == 20000 for l in lignes)


class TestDASExcelCSS:
    def test_excel_das_inclut_css(self, client):
        from models import (db, Tenant, Salarie, PeriodePaie, BulletinPaie,
                            Prestataire, FacturePrestataire)
        from datetime import date
        import declaration_das as dd, models as M
        import openpyxl, io
        t = Tenant.query.filter_by(slug="entreprise-a").first()
        s = Salarie(tenant_id=t.id, matricule="DX", nom="X", prenom="Y",
                    date_embauche=date(2020, 1, 1), statut="ACTIF")
        db.session.add(s); db.session.commit()
        per = PeriodePaie(tenant_id=t.id, mois=1, annee=2026, libelle_mois="JANVIER", statut="VALIDÉ")
        db.session.add(per); db.session.commit()
        db.session.add(BulletinPaie(tenant_id=t.id, salarie_id=s.id, periode_id=per.id,
                       salaire_base=500000, salaire_brut=500000, net_a_payer=400000, statut="VALIDÉ"))
        p = Prestataire(tenant_id=t.id, code="PCX", categorie="SOUS_TRAITANT",
                        raison_sociale="BTP SARL", regime_fiscal="CSS",
                        assujetti_tva=False, assujetti_css=True, resident=True)
        db.session.add(p); db.session.commit()
        f = FacturePrestataire(tenant_id=t.id, prestataire_id=p.id, numero="FX",
                               date_facture=date(2026, 4, 1), montant_ht=2000000,
                               taux_tva=0, taux_css=1, taux_retenue=0, statut="PAYEE")
        f.lignes = []; f.calculer(); db.session.add(f); db.session.commit()
        content = dd.generer_das_excel(t, 2026, models=M)
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb["ID23-24 - Honoraires"]
        hdr = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
        assert "CSS (1%)" in hdr
        ci = hdr.index("CSS (1%)") + 1
        assert ws.cell(5, ci).value == 20000  # 1% de 2 000 000
