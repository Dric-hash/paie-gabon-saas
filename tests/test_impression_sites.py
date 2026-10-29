"""Listes imprimées des salariés et des journaliers classées par site."""
import re
from datetime import date

from test_suppression_salarie import env, _cli  # noqa: F401
import models as M
from models import db


def _ordre(html, noms):
    pos = {n: html.find(n) for n in noms}
    assert all(v >= 0 for v in pos.values()), pos
    return [n for n, _ in sorted(pos.items(), key=lambda x: x[1])]


def _donnees(env):
    t = env["t"]
    owendo, akanda, ferme = M.Site(tenant_id=t, nom="Owendo"), M.Site(tenant_id=t, nom="Akanda"), \
        M.Site(tenant_id=t, nom="Zone fermée")
    db.session.add_all([owendo, akanda, ferme])
    db.session.flush()
    s0, s1 = db.session.get(M.Salarie, env["s"]), db.session.get(M.Salarie, env["autre"])
    s0.nom, s1.nom = "ZOLA", "ABAGA"
    s2 = M.Salarie(tenant_id=t, matricule="M9", nom="MBA", prenom="Awa", date_embauche=date(2024, 1, 1))
    j1 = M.Journalier(tenant_id=t, nom="NZE", prenom="Paul", taux_horaire=1000, statut="ACTIF")
    j2 = M.Journalier(tenant_id=t, nom="BIYOGO", prenom="Luc", taux_horaire=1000, statut="ACTIF")
    db.session.add_all([s2, j1, j2])
    db.session.flush()
    db.session.add_all([
        M.AffectationSite(tenant_id=t, site_id=owendo.id, salarie_id=s0.id, date_debut=date(2025, 1, 1)),
        # ABAGA : ancienne affectation à Owendo terminée, active à Akanda
        M.AffectationSite(tenant_id=t, site_id=owendo.id, salarie_id=s1.id, date_debut=date(2024, 1, 1), actif=False),
        M.AffectationSite(tenant_id=t, site_id=akanda.id, salarie_id=s1.id, date_debut=date(2025, 6, 1)),
        M.AffectationSite(tenant_id=t, site_id=owendo.id, journalier_id=j1.id, date_debut=date(2025, 1, 1)),
    ])
    db.session.commit()
    return {"owendo": owendo.id, "akanda": akanda.id}


def test_salaries_par_site(env):
    ids = _donnees(env)
    cli = _cli(env["admin"])
    h = cli.get("/salaries/imprimer").get_data(as_text=True)
    assert "classée par site" in h
    # sites par ordre alphabétique, sans site en dernier ; Zone fermée (vide) absente
    assert _ordre(h, ['class="site">Akanda', 'class="site">Owendo', 'class="site">Sans site affecté']) == \
        ['class="site">Akanda', 'class="site">Owendo', 'class="site">Sans site affecté']
    assert "Zone fermée" not in h.split("</form>")[-1]
    assert _ordre(h, ["ABAGA", "ZOLA", "MBA Awa"]) == ["ABAGA", "ZOLA", "MBA Awa"]
    assert h.count("<table>") == 3
    # un seul site, une page par site, ou classement alphabétique
    h = cli.get(f"/salaries/imprimer?site={ids['owendo']}").get_data(as_text=True).split("</form>")[-1]
    assert "ZOLA" in h and "ABAGA" not in h and "MBA Awa" not in h and "classée par site — Owendo" in h
    h = cli.get("/salaries/imprimer?site=0").get_data(as_text=True).split("</form>")[-1]
    assert "MBA Awa" in h and "ZOLA" not in h
    assert cli.get("/salaries/imprimer?saut=1").get_data(as_text=True).count("site-bloc saut") == 2
    h = cli.get("/salaries/imprimer?classement=alpha").get_data(as_text=True)
    assert 'class="site"' not in h and h.count("<table>") == 1
    assert _ordre(h, ["ABAGA", "MBA Awa", "ZOLA"]) == ["ABAGA", "MBA Awa", "ZOLA"]


def test_journaliers_par_site(env):
    _donnees(env)
    h = _cli(env["admin"]).get("/journaliers/imprimer").get_data(as_text=True)
    assert _ordre(h, ['class="site">Owendo', "NZE Paul", 'class="site">Sans site affecté', "BIYOGO Luc"]) == \
        ['class="site">Owendo', "NZE Paul", 'class="site">Sans site affecté', "BIYOGO Luc"]
    assert re.search(r"Owendo <span>· 1 journalier\(s\)", h)


def _paie(env):
    """Bulletins d'octobre et acomptes pour ZOLA (Owendo), ABAGA (Akanda), MBA (sans site)."""
    ids = _donnees(env)
    t = env["t"]
    per = M.PeriodePaie(tenant_id=t, annee=2026, mois=10, libelle_mois="OCTOBRE")
    db.session.add(per)
    db.session.flush()
    sals = {s.nom: s for s in M.Salarie.query.filter_by(tenant_id=t)}
    for nom, net in (("ZOLA", 300000), ("ABAGA", 200000), ("MBA", 150000)):
        s = sals[nom]
        db.session.add(M.BulletinPaie(tenant_id=t, salarie_id=s.id, periode_id=per.id, salaire_base=net,
                                      salaire_brut=net, net_a_payer=net, statut="VALIDÉ"))
        db.session.add(M.Acompte(tenant_id=t, salarie_id=s.id, montant=net // 10, date_acompte=date(2026, 10, 5),
                                 mois=10, annee=2026, statut="EN_ATTENTE"))
    s = sals["ZOLA"]
    s.date_embauche = date(2026, 10, 20)                     # embauché en dernier
    db.session.commit()
    return dict(ids, periode=per.id)


def test_bordereau_par_site(env):
    ids = _paie(env)
    cli = _cli(env["admin"])
    h = cli.get(f"/bulletins/bordereau?periode_id={ids['periode']}").get_data(as_text=True)
    assert _ordre(h, ["Akanda <span>", "ABAGA", "Sous-total Akanda (1)", "Owendo <span>", "ZOLA",
                      "Sous-total Owendo (1)", "Sans site affecté <span>", "MBA Awa"])[0] == "Akanda <span>"
    assert "Sous-total Owendo (1)</td><td class=\"num\">300000 FCFA" in h
    assert "TOTAL GÉNÉRAL (3)" in h and "650000 FCFA" in h
    h = cli.get(f"/bulletins/bordereau?periode_id={ids['periode']}&site_id={ids['owendo']}").get_data(as_text=True)
    assert "ZOLA" in h and "ABAGA" not in h and "TOTAL GÉNÉRAL (1)" in h
    h = cli.get(f"/bulletins/bordereau?periode_id={ids['periode']}&classement=alpha").get_data(as_text=True)
    assert "Sous-total" not in h and _ordre(h, ["ABAGA", "MBA", "ZOLA"]) == ["ABAGA", "MBA", "ZOLA"]


def test_journal_excel_par_site(env):
    import io
    from openpyxl import load_workbook
    ids = _paie(env)
    cli = _cli(env["admin"])
    ws = load_workbook(io.BytesIO(cli.get(f"/bulletins/export/{ids['periode']}").data)).active
    lignes = [[c.value for c in r] for r in ws.iter_rows(min_row=3)]
    assert lignes[0][:3] == ["Site", "Matricule", "Nom"]
    col_net = lignes[0].index("Net à Payer")
    resume = [(l[0], l[2], l[col_net]) for l in lignes[1:] if l[0]]
    assert resume == [("Akanda", "ABAGA", 200000), ("Sous-total Akanda (1)", None, 200000),
                      ("Owendo", "ZOLA", 300000), ("Sous-total Owendo (1)", None, 300000),
                      ("Sans site affecté", "MBA", 150000), ("Sous-total Sans site affecté (1)", None, 150000),
                      ("TOTAL GÉNÉRAL (3)", None, 650000)]
    ws = load_workbook(io.BytesIO(cli.get(f"/bulletins/export/{ids['periode']}?site_id={ids['akanda']}").data)).active
    noms = [r[2].value for r in ws.iter_rows(min_row=4) if r[2].value]
    assert noms == ["ABAGA"]


def test_bulletins_pdf_groupe_par_site(env, monkeypatch):
    import pdf_bulletin
    ids = _paie(env)
    ordre = []
    monkeypatch.setattr(pdf_bulletin, "generer_bulletins_pdf",
                        lambda buls, t, modele=None: ordre.extend(b.salarie.nom for b in buls) or b"%PDF-1.4")
    cli = _cli(env["admin"])
    assert cli.get(f"/bulletins/export-zip/{ids['periode']}").status_code == 200
    assert ordre == ["ABAGA", "ZOLA", "MBA"]
    ordre.clear()
    cli.get(f"/bulletins/export-zip/{ids['periode']}?site_id={ids['owendo']}")
    assert ordre == ["ZOLA"]
    ordre.clear()
    cli.get(f"/bulletins/export-zip/{ids['periode']}?classement=alpha")
    assert ordre == ["ABAGA", "MBA", "ZOLA"]


def test_acomptes_par_site(env):
    ids = _paie(env)
    a_ids = [str(a.id) for a in M.Acompte.query.all()]
    h = _cli(env["admin"]).post("/acomptes/imprimer", data={"acompte_ids": a_ids, "mois": "10", "annee": "2026"}
                                ).get_data(as_text=True)
    assert _ordre(h, ["Akanda <span>", "Owendo <span>", "Sans site affecté <span>"]) == \
        ["Akanda <span>", "Owendo <span>", "Sans site affecté <span>"]
    assert "Sous-total Owendo (hors annulés)</td><td class=\"num\">30000" in h


def test_registre_chronologique_ou_par_site(env):
    _paie(env)
    cli = _cli(env["admin"])
    h = cli.get("/salaries/registre").get_data(as_text=True)
    assert "site-row" not in h.split("</form>")[-1]
    # chronologique : ZOLA (embauché en 2026) en dernier, n° 3
    assert _ordre(h, ["ZOLA", "ABAGA"]) == ["ABAGA", "ZOLA"]
    h = cli.get("/salaries/registre?classement=site").get_data(as_text=True)
    corps = h.split("</form>")[-1]
    assert _ordre(corps, ["Akanda <span>", "Owendo <span>", "Sans site affecté <span>"])[0] == "Akanda <span>"
    # le n° d'inscription reste celui de l'ordre chronologique
    import re as _re
    zola = _re.search(r'<td class="c mono">(\d+)</td>\s*<td class="mono">M0</td>', corps)
    assert zola and zola.group(1) == "3"
