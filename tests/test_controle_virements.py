"""Contrôle de la paie avant validation, coordonnées bancaires, ordres de virement et paie en espèces."""
from datetime import date

from test_suppression_salarie import env, _cli, _flash  # noqa: F401
import models as M
from models import db
import controle_paie as cp

AUJ = date.today()


def _periode(t, a, m):
    p = M.PeriodePaie(tenant_id=t, annee=a, mois=m, libelle_mois=M.PeriodePaie.MOIS_NOMS[m])
    db.session.add(p)
    db.session.flush()
    return p


def _bul(t, s, p, net, statut="VALIDÉ", hs=0):
    b = M.BulletinPaie(tenant_id=t, salarie_id=s, periode_id=p.id, salaire_base=net, salaire_brut=net, net_a_payer=net,
                       heures_sup_10=hs, statut=statut)
    db.session.add(b)
    db.session.flush()
    return b


def _scenario(env):
    """s : 200 000 nets et 10 000 d'heures sup en juillet-septembre, puis 300 000 et 50 000 d'heures sup en octobre ;
    autre : embauché en octobre, sans bulletin ; s3 : net négatif, contrat à 100 000, virement sans RIB."""
    t = env["t"]
    s, autre = db.session.get(M.Salarie, env["s"]), db.session.get(M.Salarie, env["autre"])
    s.date_embauche = date(2024, 1, 1)
    s3 = M.Salarie(tenant_id=t, matricule="M3", nom="TROIS", prenom="T", date_embauche=date(2024, 1, 1),
                   mode_paiement="VIREMENT")
    db.session.add(s3)
    db.session.flush()
    db.session.add(M.Contrat(tenant_id=t, salarie_id=s3.id, date_debut=date(2024, 1, 1), salaire_base=100000, actif=True))
    for m in (7, 8, 9):
        _bul(t, s.id, _periode(t, 2026, m), 200000, hs=10000)
    oct_ = _periode(t, 2026, 10)
    b_s = _bul(t, s.id, oct_, 300000, statut="BROUILLON", hs=50000)
    b_3 = _bul(t, s3.id, oct_, -5000, statut="BROUILLON")
    db.session.commit()
    return {"oct": oct_.id, "s": s.id, "autre": autre.id, "s3": s3.id, "b_s": b_s.id, "b_3": b_3.id}


def _codes(anomalies):
    return {(a["salarie"].id if a["salarie"] else None, a["code"]) for a in anomalies}


def test_controle_de_la_paie(env):
    ids = _scenario(env)
    t = db.session.get(M.Tenant, env["t"])
    t.controle_salaire_minimum = 150000
    db.session.commit()
    p = db.session.get(M.PeriodePaie, ids["oct"])
    an = cp.controler(t, p)
    codes = _codes(an)
    assert {(ids["s"], "VARIATION"), (ids["s"], "HEURES_SUP"), (ids["autre"], "SANS_BULLETIN"),
            (ids["s3"], "NET_NEGATIF"), (ids["s3"], "SOUS_MINIMUM"), (ids["s3"], "SANS_RIB")} <= codes
    assert (ids["autre"], "ENTREE") not in codes                 # pas de bulletin : seule l'absence est signalée
    assert an[0]["niveau"] == "bloquant" and cp.resume(an)["bloquant"] == 1
    variation = next(a for a in an if a["code"] == "VARIATION")
    assert "hausse de 50 %" in variation["titre"]
    cli = _cli(env["admin"])
    page = cli.get(f"/paie/controle?periode_id={ids['oct']}").get_data(as_text=True)
    assert "Net à payer négatif" in page and "Heures supplémentaires inhabituelles" in page and "paramètre de l&#39;entreprise" in page
    assert "Contrôler la paie (" in cli.get(f"/bulletins?periode_id={ids['oct']}").get_data(as_text=True)
    # validation : le net négatif est refusé, le reste passe
    cli.post("/bulletins/valider-lot", data={"periode_id": ids["oct"], "action_lot": "valider", "tout_valider": "1"})
    assert "1 bulletin(s) non validé(s) (net négatif" in _flash(cli)
    assert db.session.get(M.BulletinPaie, ids["b_s"]).statut == "VALIDÉ"
    assert db.session.get(M.BulletinPaie, ids["b_3"]).statut == "BROUILLON"
    cli.post(f"/bulletins/{ids['b_3']}/valider")
    assert "Net à payer négatif" in _flash(cli) and db.session.get(M.BulletinPaie, ids["b_3"]).statut == "BROUILLON"
    # réglages : seuil 60 % et minimum 90 000 → variation et minimum ne sont plus signalés
    cli.post("/paie/controle/parametres", data={"seuil": "60", "minimum": "90000", "periode_id": ids["oct"]})
    t = db.session.get(M.Tenant, env["t"])
    codes = _codes(cp.controler(t, db.session.get(M.PeriodePaie, ids["oct"])))
    assert (ids["s"], "VARIATION") not in codes and (ids["s3"], "SOUS_MINIMUM") not in codes
    cli.post("/paie/controle/parametres", data={"seuil": "150", "periode_id": ids["oct"]})
    assert "entre 0 et 100" in _flash(cli)
    # sans réglage : valeur du moteur de paie, signalée comme à confirmer
    t.controle_salaire_minimum = None
    db.session.commit()
    assert cp.salaire_minimum(t)[1] == "valeur du moteur de paie (à confirmer)"


def test_fiche_salarie_paiement(env):
    cli = _cli(env["admin"])
    sid = env["s"]
    base = {"nom": "NOM0", "prenom": "P", "date_embauche": "2026-10-10", "statut": "ACTIF"}
    cli.post(f"/salaries/{sid}/modifier", data=dict(base, mode_paiement="VIREMENT", banque="BGFI Bank",
                                                    rib="  GA21  4000 3000  123 ", titulaire_compte=""))
    s = db.session.get(M.Salarie, sid)
    assert (s.mode_paiement, s.banque, s.rib, s.titulaire_compte) == ("VIREMENT", "BGFI Bank", "GA21 4000 3000 123", None)
    # un enregistrement sans la section paiement ne remet plus le salarié en espèces
    cli.post(f"/salaries/{sid}/modifier", data=base)
    assert db.session.get(M.Salarie, sid).mode_paiement == "VIREMENT"
    assert "Paiement du salaire" in cli.get(f"/salaries/{sid}/modifier").get_data(as_text=True)


def test_ordre_de_virement_et_especes(env, monkeypatch):
    import interco_caisse
    envois = []
    monkeypatch.setattr(interco_caisse, "proposer_ecriture", lambda t, **k: envois.append(k) or True)
    t = env["t"]
    s, autre = db.session.get(M.Salarie, env["s"]), db.session.get(M.Salarie, env["autre"])
    s.mode_paiement, s.banque, s.rib = "VIREMENT", "BGFI Bank", "40003 04100 12345678901 12"
    autre.mode_paiement = "VIREMENT"                     # sans RIB
    s3 = M.Salarie(tenant_id=t, matricule="M3", nom="TROIS", prenom="T", date_embauche=date(2024, 1, 1))
    db.session.add(s3)
    db.session.flush()
    p = _periode(t, 2026, 10)
    b1 = _bul(t, s.id, p, 300000).id
    _bul(t, autre.id, p, 250000)
    _bul(t, s3.id, p, 100000)
    db.session.commit()
    cli = _cli(env["admin"])
    cli.post("/virements/comptes", data={"banque": "UGB", "rib": "40001 00001 00000000001 01", "intitule": "Salaires"})
    cb = M.CompteBancaire.query.filter_by(tenant_id=t).one().id
    page = cli.get(f"/virements?periode_id={p.id}").get_data(as_text=True)
    assert "Bulletins à virer (1)" in page and "RIB manquant (1)" in page and "NOM1 P" in page
    r = cli.post("/virements/ordres", data={"periode_id": p.id, "compte_id": cb, "bulletin_id": [b1],
                                            "date_execution": AUJ.isoformat()})
    o = M.OrdreVirement.query.filter_by(tenant_id=t).one()
    assert (o.numero, o.total, o.lignes[0].rib, o.lignes[0].nom) == (f"OV-{AUJ.year}-0001", 300000,
                                                                     "40003 04100 12345678901 12", "NOM0 P")
    # le bulletin n'est plus proposé, ne se paie pas à la main, ne repasse pas en brouillon
    assert "Bulletins à virer (0)" in cli.get(f"/virements?periode_id={p.id}").get_data(as_text=True)
    cli.post(f"/bulletins/{b1}/payer", data={})
    assert "OV-" in _flash(cli) and db.session.get(M.BulletinPaie, b1).statut == "VALIDÉ"
    cli.post("/bulletins/valider-lot", data={"periode_id": p.id, "action_lot": "annuler_validation", "bulletin_ids": str(b1)})
    assert db.session.get(M.BulletinPaie, b1).statut == "VALIDÉ"
    # lettre et Excel
    lettre = cli.get(f"/virements/ordres/{o.id}/lettre").get_data(as_text=True)
    assert "ORDRE DE VIREMENT N° OV-".lower() in lettre.lower() and "trois cent mille" in lettre and "40001 00001" in lettre
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(cli.get(f"/virements/ordres/{o.id}/excel").data)).active
    assert [c.value for c in ws[6]][2:] == ["NOM0 P", "BGFI Bank", "40003 04100 12345678901 12", 300000]
    # exécution : bulletin payé par virement, proposé à la Caisse
    cli.post(f"/virements/ordres/{o.id}/executer", data={"date_execution": AUJ.isoformat()})
    b = db.session.get(M.BulletinPaie, b1)
    assert (b.statut, b.mode_paiement, b.date_paiement) == ("PAYÉ", "VIREMENT", AUJ)
    assert envois and envois[0]["source_ref"] == f"bulletin-{b1}" and envois[0]["montant"] == 300000
    cli.post(f"/virements/ordres/{o.id}/annuler", data={"motif": "x"})
    assert "ne s'annule pas" in _flash(cli)
    # espèces : s3 (fiche en espèces), pas les salariés en virement
    esp = cli.get(f"/virements/especes?periode_id={p.id}").get_data(as_text=True)
    assert "TROIS T" in esp and "NOM0 P" not in esp and "cent mille" in esp


def test_annulation_et_cloisonnement(env):
    t = env["t"]
    s = db.session.get(M.Salarie, env["s"])
    s.mode_paiement, s.rib = "VIREMENT", "123"
    p = _periode(t, 2026, 10)
    b1 = _bul(t, s.id, p, 100000).id
    db.session.add(M.CompteBancaire(tenant_id=t, banque="UGB", rib="1"))
    db.session.commit()
    cli = _cli(env["admin"])
    cb = M.CompteBancaire.query.filter_by(tenant_id=t).one().id
    cli.post("/virements/ordres", data={"periode_id": p.id, "compte_id": cb, "bulletin_id": [b1]})
    o = M.OrdreVirement.query.one()
    cli.post(f"/virements/ordres/{o.id}/annuler", data={"motif": ""})
    assert "motif" in _flash(cli)
    cli.post(f"/virements/ordres/{o.id}/annuler", data={"motif": "Erreur de compte"})
    assert db.session.get(M.OrdreVirement, o.id).statut == "ANNULE"
    cli.post("/virements/ordres", data={"periode_id": p.id, "compte_id": cb, "bulletin_id": [b1]})
    assert M.OrdreVirement.query.count() == 2 and M.OrdreVirement.query.order_by(M.OrdreVirement.id.desc()).first().numero \
        == f"OV-{AUJ.year}-0002"
    # une autre entreprise ne voit ni l'ordre ni ses comptes
    t2 = M.Tenant(slug="b", denomination="BETA", pays="Gabon", plan_id=db.session.get(M.Tenant, t).plan_id,
                  statut="ACTIF", token_api="b")
    db.session.add(t2)
    db.session.flush()
    u2 = M.Utilisateur(nom="B", prenom="B", email="b@b.ga", role="TENANT_ADMIN", tenant_id=t2.id, actif=True,
                       email_verifie=True)
    u2.set_password("x")
    db.session.add(u2)
    db.session.commit()
    cli2 = _cli(u2.id)
    from flask import g
    g.pop("_login_user", None)      # le contexte d'application du test est partagé : oublier l'utilisateur précédent
    for url in (f"/virements/ordres/{o.id}", f"/virements/ordres/{o.id}/lettre", f"/virements/ordres/{o.id}/excel"):
        assert cli2.get(url).status_code == 404, url
    assert cli2.post(f"/virements/ordres/{o.id}/executer", data={"date_execution": AUJ.isoformat()}).status_code == 404
    assert "UGB" not in cli2.get("/virements/comptes").get_data(as_text=True)
    assert cli2.post("/virements/ordres", data={"periode_id": p.id, "compte_id": cb, "bulletin_id": [b1]}).status_code == 404


def test_brouillon_d_un_ordre_annule_supprimable(env):
    """Bulletin viré puis ordre annulé, validation annulée, brouillon supprimé : l'historique de
    l'ordre reste (nom, RIB, montant) sans bloquer la suppression (contrainte de clé étrangère)."""
    t = env["t"]
    s = db.session.get(M.Salarie, env["s"])
    s.mode_paiement, s.rib = "VIREMENT", "RIB-1"
    p = _periode(t, 2026, 10)
    b1 = _bul(t, s.id, p, 100000).id
    db.session.add(M.CompteBancaire(tenant_id=t, banque="UGB", rib="1"))
    db.session.commit()
    cli = _cli(env["admin"])
    cb = M.CompteBancaire.query.filter_by(tenant_id=t).one().id
    cli.post("/virements/ordres", data={"periode_id": p.id, "compte_id": cb, "bulletin_id": [b1]})
    o = M.OrdreVirement.query.one()
    # ordre en cours : suppression refusée même si le bulletin était brouillon (cas forcé)
    db.session.get(M.BulletinPaie, b1).statut = "BROUILLON"
    db.session.commit()
    cli.post(f"/bulletins/{b1}/supprimer")
    assert "ordre de virement" in _flash(cli) and db.session.get(M.BulletinPaie, b1) is not None
    db.session.get(M.BulletinPaie, b1).statut = "VALIDÉ"
    db.session.commit()
    cli.post(f"/virements/ordres/{o.id}/annuler", data={"motif": "Erreur"})
    cli.post("/bulletins/valider-lot", data={"periode_id": p.id, "action_lot": "annuler_validation", "bulletin_ids": str(b1)})
    assert db.session.get(M.BulletinPaie, b1).statut == "BROUILLON"
    cli.post(f"/bulletins/{b1}/supprimer")
    db.session.expire_all()
    assert db.session.get(M.BulletinPaie, b1) is None
    l = M.LigneOrdreVirement.query.one()
    assert (l.bulletin_id, l.nom, l.rib, float(l.montant)) == (None, "NOM0 P", "RIB-1", 100000)
