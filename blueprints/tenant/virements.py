"""blueprints/tenant/virements.py — Contrôle de la paie, ordres de virement et paie en espèces.

  /paie/controle            anomalies de la période avant validation (controle_paie.py)
  /virements                bulletins validés à payer par virement, ordres de la période
  /virements/ordres/<id>    ordre de virement : lettre à la banque, export Excel, exécution
  /virements/comptes        comptes bancaires de l'entreprise (donneurs d'ordre)
  /virements/especes        état de paie en espèces à faire émarger
"""
from datetime import date, datetime
from io import BytesIO

from flask import render_template, request, redirect, url_for, flash, abort, send_file, current_app
from flask_login import login_required, current_user
from sqlalchemy import func

from blueprints.tenant import bp
from core import tenant_required, get_tenant
from audit import log_action
from models import (db, utcnow, BulletinPaie, PeriodePaie, Salarie, CompteBancaire, OrdreVirement,
                    LigneOrdreVirement)
import controle_paie as cp

STATUTS_OV = {"PREPARE": ("Préparé", "#1e40af", "#dbeafe"), "EXECUTE": ("Exécuté", "#065f46", "#d1fae5"),
              "ANNULE": ("Annulé", "#b91c1c", "#fee2e2")}
VALIDES = ("VALIDÉ", "VALIDE")


def _periode(t, pid=None):
    """Période demandée, sinon la plus récente."""
    q = PeriodePaie.query.filter_by(tenant_id=t.id)
    if pid:
        return q.filter_by(id=pid).first_or_404()
    return q.order_by(PeriodePaie.annee.desc(), PeriodePaie.mois.desc()).first()


def _periodes(t):
    return PeriodePaie.query.filter_by(tenant_id=t.id).order_by(PeriodePaie.annee.desc(), PeriodePaie.mois.desc()).all()


def _auteur():
    return getattr(current_user, "nom_complet", None) or current_user.email


# ═══════════════════════════ Contrôle de la paie ═══════════════════════════
@bp.route("/paie/controle")
@login_required
@tenant_required
def controle_paie():
    t = get_tenant()
    p = _periode(t, request.args.get("periode_id", type=int))
    anomalies = cp.controler(t, p) if p else []
    nb_bulletins = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=p.id).count() if p else 0
    minimum, source = cp.salaire_minimum(t)
    return render_template("tenant/controle_paie.html", periode=p, periodes=_periodes(t), anomalies=anomalies,
                           resume=cp.resume(anomalies), niveaux=cp.NIVEAUX, nb_bulletins=nb_bulletins,
                           minimum=minimum, source_minimum=source, seuil=cp.seuil_variation(t), tenant=t)


@bp.route("/paie/controle/parametres", methods=["POST"])
@login_required
@tenant_required
def controle_paie_parametres():
    if not current_user.is_tenant_admin:
        abort(403)
    t = get_tenant()
    try:
        seuil = float((request.form.get("seuil") or "").replace(",", ".")) if request.form.get("seuil") else None
        minimum = float((request.form.get("minimum") or "").replace(" ", "").replace(",", ".")) \
            if request.form.get("minimum") else None
    except ValueError:
        flash("Valeurs invalides.", "error")
        return redirect(url_for("tenant.controle_paie", periode_id=request.form.get("periode_id")))
    if seuil is not None and not 0 < seuil <= 100:
        flash("Le seuil de variation doit être entre 0 et 100 %.", "error")
    elif minimum is not None and minimum < 0:
        flash("Le minimum ne peut pas être négatif.", "error")
    else:
        t.controle_seuil_variation, t.controle_salaire_minimum = seuil, minimum
        db.session.commit()
        log_action("UPDATE", "parametres", t.id, "Réglages du contrôle de la paie",
                   apres={"seuil_variation": seuil, "salaire_minimum": minimum})
        flash("Réglages du contrôle enregistrés.", "success")
    return redirect(url_for("tenant.controle_paie", periode_id=request.form.get("periode_id")))


# ═══════════════════════════ Comptes bancaires de l'entreprise ═══════════════════════════
@bp.route("/virements/comptes", methods=["GET", "POST"])
@login_required
@tenant_required
def virements_comptes():
    t = get_tenant()
    if request.method == "POST":
        if not current_user.is_tenant_admin:
            abort(403)
        banque = (request.form.get("banque") or "").strip()[:100]
        rib = " ".join((request.form.get("rib") or "").split())[:50]
        if not banque or not rib:
            flash("Indiquez la banque et le RIB du compte.", "error")
        else:
            db.session.add(CompteBancaire(tenant_id=t.id, banque=banque, rib=rib,
                                          intitule=(request.form.get("intitule") or "").strip()[:150] or None,
                                          agence=(request.form.get("agence") or "").strip()[:150] or None))
            db.session.commit()
            log_action("CREATE", "compte_bancaire", None, f"Compte bancaire ajouté : {banque}")
            flash(f"Compte {banque} ajouté.", "success")
        return redirect(url_for("tenant.virements_comptes"))
    comptes = CompteBancaire.query.filter_by(tenant_id=t.id).order_by(CompteBancaire.actif.desc(), CompteBancaire.banque).all()
    return render_template("tenant/virements_comptes.html", comptes=comptes)


@bp.route("/virements/comptes/<int:cbid>/basculer", methods=["POST"])
@login_required
@tenant_required
def virements_compte_basculer(cbid):
    if not current_user.is_tenant_admin:
        abort(403)
    t = get_tenant()
    c = CompteBancaire.query.filter_by(id=cbid, tenant_id=t.id).first_or_404()
    c.actif = not (c.actif is not False)
    db.session.commit()
    flash(f"Compte {c.banque} {'réactivé' if c.actif else 'désactivé'}.", "success")
    return redirect(url_for("tenant.virements_comptes"))


# ═══════════════════════════ Virements ═══════════════════════════
def bulletins_en_ordre(tid):
    """{bulletin_id: ordre} pour les ordres non annulés."""
    res = {}
    q = (db.session.query(LigneOrdreVirement.bulletin_id, OrdreVirement).join(OrdreVirement)
         .filter(OrdreVirement.tenant_id == tid, OrdreVirement.statut != "ANNULE"))
    for bid, o in q.all():
        res[bid] = o
    return res


def detacher_bulletins(tid, bulletin_ids):
    """Avant de supprimer des bulletins (brouillons) : les lignes d'ordres ANNULÉS qui les citent
    gardent leur historique sans le lien. Renvoie False si un bulletin est sur un ordre en cours
    ou exécuté (suppression à refuser)."""
    ids = [i for i in bulletin_ids if i]
    if not ids:
        return True
    lignes = (LigneOrdreVirement.query.join(OrdreVirement)
              .filter(LigneOrdreVirement.tenant_id == tid, LigneOrdreVirement.bulletin_id.in_(ids)).all())
    if any(l.ordre.statut != "ANNULE" for l in lignes):
        return False
    for l in lignes:
        l.bulletin_id = None
    db.session.flush()
    return True


def _a_virer(t, p):
    """(prêts, sans RIB) : bulletins validés de la période, salarié payé par virement, hors ordre en cours."""
    deja = bulletins_en_ordre(t.id)
    prets, sans_rib = [], []
    q = (BulletinPaie.query.join(Salarie).filter(BulletinPaie.tenant_id == t.id, BulletinPaie.periode_id == p.id,
                                                  BulletinPaie.statut.in_(VALIDES), Salarie.mode_paiement == "VIREMENT")
         .order_by(Salarie.nom, Salarie.prenom))
    for b in q.all():
        if b.id in deja or float(b.net_a_payer or 0) <= 0:
            continue
        (prets if (b.salarie.rib or "").strip() else sans_rib).append(b)
    return prets, sans_rib


def _numero_ov(tid, jour):
    prefixe = f"OV-{jour.year}-"
    n = 0
    for (num,) in db.session.query(OrdreVirement.numero).filter(OrdreVirement.tenant_id == tid,
                                                                 OrdreVirement.numero.like(prefixe + "%")):
        try:
            n = max(n, int(num[len(prefixe):]))
        except ValueError:
            pass
    return f"{prefixe}{n + 1:04d}"


@bp.route("/virements")
@login_required
@tenant_required
def virements():
    t = get_tenant()
    p = _periode(t, request.args.get("periode_id", type=int))
    prets, sans_rib, ordres, nb_brouillons = [], [], [], 0
    if p:
        prets, sans_rib = _a_virer(t, p)
        ordres = (OrdreVirement.query.filter_by(tenant_id=t.id, periode_id=p.id)
                  .order_by(OrdreVirement.id.desc()).all())
        nb_brouillons = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=p.id, statut="BROUILLON").count()
    return render_template("tenant/virements.html", periode=p, periodes=_periodes(t), prets=prets, sans_rib=sans_rib,
                           ordres=ordres, statuts=STATUTS_OV, nb_brouillons=nb_brouillons, aujourdhui=date.today(),
                           comptes=CompteBancaire.query.filter_by(tenant_id=t.id).filter(
                               CompteBancaire.actif.isnot(False)).order_by(CompteBancaire.banque).all())


@bp.route("/virements/ordres", methods=["POST"])
@login_required
@tenant_required
def virements_ordre_creer():
    if not current_user.can_edit:
        abort(403)
    t = get_tenant()
    p = PeriodePaie.query.filter_by(id=request.form.get("periode_id", type=int), tenant_id=t.id).first_or_404()
    retour = redirect(url_for("tenant.virements", periode_id=p.id))
    compte = CompteBancaire.query.filter_by(id=request.form.get("compte_id", type=int), tenant_id=t.id).first()
    if compte is None or compte.actif is False:
        flash("Choisissez le compte de l'entreprise à débiter.", "error")
        return retour
    try:
        execution = datetime.strptime((request.form.get("date_execution") or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        execution = None
    ids = {int(x) for x in request.form.getlist("bulletin_id") if str(x).isdigit()}
    prets, _ = _a_virer(t, p)
    choisis = [b for b in prets if b.id in ids]
    if not choisis:
        flash("Cochez au moins un bulletin à virer.", "error")
        return retour
    o = OrdreVirement(tenant_id=t.id, periode_id=p.id, compte_id=compte.id, numero=_numero_ov(t.id, date.today()),
                      date_ordre=date.today(), date_execution=execution, auteur=_auteur())
    db.session.add(o)
    db.session.flush()
    for b in choisis:
        s = b.salarie
        db.session.add(LigneOrdreVirement(tenant_id=t.id, ordre_id=o.id, bulletin_id=b.id, salarie_id=s.id,
                                          matricule=s.matricule, nom=(s.titulaire_compte or s.nom_complet)[:250],
                                          banque=s.banque, rib=s.rib.strip(), montant=b.net_a_payer))
    db.session.commit()
    log_action("CREATE", "ordre_virement", o.id, f"Ordre de virement {o.numero} — {len(choisis)} salarié(s), "
               f"{o.total:,.0f} FCFA".replace(",", " "))
    flash(f"Ordre {o.numero} préparé : imprimez la lettre pour la banque, puis marquez-le exécuté une fois "
          f"le virement passé.", "success")
    return redirect(url_for("tenant.virements_ordre", ovid=o.id))


def _ordre(ovid):
    t = get_tenant()
    return t, OrdreVirement.query.filter_by(id=ovid, tenant_id=t.id).first_or_404()


@bp.route("/virements/ordres/<int:ovid>")
@login_required
@tenant_required
def virements_ordre(ovid):
    t, o = _ordre(ovid)
    return render_template("tenant/virement_ordre.html", o=o, statuts=STATUTS_OV, aujourdhui=date.today())


@bp.route("/virements/ordres/<int:ovid>/lettre")
@login_required
@tenant_required
def virements_ordre_lettre(ovid):
    t, o = _ordre(ovid)
    from montant_lettres import montant_en_lettres
    return render_template("tenant/virement_lettre_print.html", o=o, t=t,
                           total_lettres=montant_en_lettres(round(o.total)), now=utcnow())


@bp.route("/virements/ordres/<int:ovid>/excel")
@login_required
@tenant_required
def virements_ordre_excel(ovid):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    t, o = _ordre(ovid)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = o.numero
    ws.append([f"Ordre de virement {o.numero} — {t.denomination}"])
    ws.append([f"Compte à débiter : {o.compte.banque} — {o.compte.rib}"])
    ws.append([f"Salaires {o.periode.libelle_mois} {o.periode.annee}"])
    ws.append([])
    ws.append(["N°", "Matricule", "Bénéficiaire", "Banque", "RIB / IBAN", "Montant (FCFA)"])
    for c in ws[5]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="0F3D36")
        c.alignment = Alignment(horizontal="center")
    for i, l in enumerate(o.lignes, 1):
        ws.append([i, l.matricule, l.nom, l.banque or "", l.rib, round(float(l.montant or 0))])
        ws.cell(row=ws.max_row, column=5).number_format = "@"
        ws.cell(row=ws.max_row, column=6).number_format = "#,##0"
    ws.append(["", "", "TOTAL", "", "", round(o.total)])
    ws.cell(row=ws.max_row, column=6).number_format = "#,##0"
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    ws["A1"].font = Font(bold=True, size=13)
    for col, w in zip("ABCDEF", (6, 14, 34, 22, 34, 16)):
        ws.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{o.numero}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp.route("/virements/ordres/<int:ovid>/executer", methods=["POST"])
@login_required
@tenant_required
def virements_ordre_executer(ovid):
    """Le virement est passé : les bulletins de l'ordre sont marqués payés (mode virement, date
    d'exécution) et proposés à la Caisse, comme un paiement normal."""
    if not current_user.can_edit:
        abort(403)
    t, o = _ordre(ovid)
    retour = redirect(url_for("tenant.virements_ordre", ovid=o.id))
    if o.statut != "PREPARE":
        flash("Cet ordre n'est plus en attente.", "warning")
        return retour
    try:
        jour = datetime.strptime((request.form.get("date_execution") or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        jour = None
    if jour is None or jour > date.today() or jour < o.date_ordre:
        flash("Date d'exécution invalide (entre la date de l'ordre et aujourd'hui).", "error")
        return retour
    payes = []
    for l in o.lignes:
        b = l.bulletin
        if b is None or b.statut not in VALIDES:
            continue
        b.statut, b.date_paiement, b.mode_paiement = "PAYÉ", jour, "VIREMENT"
        payes.append(b)
    o.statut, o.date_execution = "EXECUTE", jour
    db.session.commit()
    log_action("PAY", "ordre_virement", o.id, f"Ordre {o.numero} exécuté — {len(payes)} bulletin(s) payé(s)")
    for b in payes:
        try:
            from interco_caisse import proposer_ecriture
            nom = b.salarie.nom_complet if b.salarie else "salarié"
            proposer_ecriture(t, source_ref=f"bulletin-{b.id}", montant=float(b.net_a_payer or 0),
                              motif=f"Salaire {nom} — {b.periode.libelle_mois} {b.periode.annee}",
                              compte_suggere="6611", date_operation=jour, mode_paiement="VIREMENT")
        except Exception as e:  # ne bloque jamais
            current_app.logger.warning(f"[COMPTA] Proposition d'écriture échouée (bulletin {b.id}) : {e}")
    ignores = len(o.lignes) - len(payes)
    flash(f"Ordre {o.numero} exécuté : {len(payes)} bulletin(s) payé(s) le {jour:%d/%m/%Y}."
          + (f" {ignores} bulletin(s) ignoré(s) (plus au statut validé)." if ignores else ""), "success")
    return retour


@bp.route("/virements/ordres/<int:ovid>/annuler", methods=["POST"])
@login_required
@tenant_required
def virements_ordre_annuler(ovid):
    if not current_user.can_edit:
        abort(403)
    t, o = _ordre(ovid)
    motif = (request.form.get("motif") or "").strip()[:300]
    if o.statut != "PREPARE":
        flash("Un ordre exécuté ne s'annule pas ici : les bulletins sont déjà payés.", "warning")
    elif not motif:
        flash("Indiquez le motif de l'annulation.", "error")
    else:
        o.statut, o.motif_annulation = "ANNULE", motif
        db.session.commit()
        log_action("CANCEL", "ordre_virement", o.id, f"Ordre {o.numero} annulé : {motif}")
        flash(f"Ordre {o.numero} annulé : ses bulletins peuvent être repris dans un nouvel ordre.", "success")
    return redirect(url_for("tenant.virements_ordre", ovid=o.id))


# ═══════════════════════════ Paie en espèces ═══════════════════════════
@bp.route("/virements/especes")
@login_required
@tenant_required
def virements_especes():
    """État de paie en espèces de la période, à faire émarger par chaque salarié."""
    t = get_tenant()
    p = _periode(t, request.args.get("periode_id", type=int))
    if p is None:
        abort(404)
    buls = (BulletinPaie.query.join(Salarie)
            .filter(BulletinPaie.tenant_id == t.id, BulletinPaie.periode_id == p.id,
                    BulletinPaie.statut.in_(VALIDES + ("PAYÉ", "PAYE")))
            .order_by(Salarie.nom, Salarie.prenom).all())
    # payé : mode figé au paiement ; validé : mode de la fiche salarié
    especes = [b for b in buls if ((b.mode_paiement if b.statut in ("PAYÉ", "PAYE") else b.salarie.mode_paiement)
                                   or "ESPECES") == "ESPECES" and float(b.net_a_payer or 0) > 0]
    from montant_lettres import montant_en_lettres
    total = sum(float(b.net_a_payer or 0) for b in especes)
    return render_template("tenant/paie_especes_print.html", t=t, periode=p, bulletins=especes, total=total,
                           total_lettres=montant_en_lettres(round(total)), now=utcnow())
