# -*- coding: utf-8 -*-
"""Journaliers : fiche, pointage, paie journalière, avances, feuilles, imports/exports —
extrait de tenant.py (même blueprint « tenant »)."""
from datetime import datetime, date, timedelta
from flask import (render_template, request, redirect, url_for, flash, session,
                   current_app, abort, send_file)
from flask_login import login_required, current_user
from sqlalchemy.orm import joinedload
from sqlalchemy import desc, or_, func
from blueprints.tenant import bp, _pointages_mois_contexte, _resoudre_mois_annee
from core import tenant_required, get_tenant, _pd, _parse_date
from audit import log_action
from models import (db, Journalier, Pointage, FeuillePaieJournalier, AvanceJournalier,
                    Salarie, Site, AffectationSite, Contrat, CategorieEmploi)


@bp.route("/journaliers/<int:id>/convertir", methods=["GET", "POST"])
@login_required
def journalier_convertir(id):
    """Transforme un journalier en salarié (mensuel).

    Principe : on CRÉE un nouveau salarié à partir des données du journalier, on
    ARCHIVE le journalier (statut CONVERTI) pour préserver tout son historique de
    paie journalière (pointages/avances/paies passés restent attachés au
    journalier), et on reporte son affectation de site active. L'opération est
    neutre pour le quota (−1 journalier actif, +1 salarié actif)."""
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))
    if not current_user.can_edit:
        abort(403)
    j = Journalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if (j.statut or "").upper() != "ACTIF":
        flash("Ce journalier n'est pas actif — il a peut-être déjà été converti.", "error")
        return redirect(url_for("tenant.journalier_detail", id=j.id))
    cats = CategorieEmploi.query.filter_by(tenant_id=t.id).all()

    # Matricule suggéré (préfixe sur le nom + séquence, garanti unique).
    import re as _re
    base = (_re.sub(r"[^A-Z]", "", (j.nom or "").upper())[:4]) or "SAL"
    _n = 1
    while Salarie.query.filter_by(tenant_id=t.id, matricule=f"{base}{_n:03d}").first():
        _n += 1
    suggestion = f"{base}{_n:03d}"

    if request.method == "POST":
        matricule = (request.form.get("matricule") or "").strip().upper()
        if not matricule:
            flash("Le matricule est obligatoire.", "error")
            return redirect(url_for("tenant.journalier_convertir", id=j.id))
        if Salarie.query.filter_by(tenant_id=t.id, matricule=matricule).first():
            flash(f"Le matricule « {matricule} » existe déjà. Choisissez-en un autre.", "error")
            return redirect(url_for("tenant.journalier_convertir", id=j.id))
        sb = float(request.form.get("salaire_base") or 0)

        s = Salarie(
            tenant_id=t.id, matricule=matricule,
            categorie_id=request.form.get("categorie_id") or None,
            nom=(request.form.get("nom") or j.nom).strip().upper(),
            prenom=(request.form.get("prenom") or j.prenom).strip(),
            telephone=request.form.get("telephone") or j.telephone,
            email=(request.form.get("email") or "").strip() or None,
            nationalite=request.form.get("nationalite") or j.nationalite or "GABONAISE",
            sexe=request.form.get("sexe"),
            date_naissance=_pd(request.form.get("date_naissance")),
            date_embauche=_pd(request.form.get("date_embauche")) or j.date_embauche or date.today(),
            situation_matrimoniale=request.form.get("situation_matrimoniale"),
            nb_enfants=int(request.form.get("nb_enfants") or 0),
            nb_enfants_moins_16ans=int(request.form.get("nb_enfants_moins_16ans") or 0),
            nombre_parts=float(request.form.get("nombre_parts") or 1),
            numero_cnss=request.form.get("numero_cnss"),
            numero_cnamgs=request.form.get("numero_cnamgs"),
            emploi=request.form.get("emploi") or j.profession,
            assujetti_cnamgs=request.form.get("assujetti_cnamgs") == "OUI",
            nif=request.form.get("nif"),
            niveau=request.form.get("niveau"),
            code_emploi=request.form.get("code_emploi"),
            statut="ACTIF",
        )
        db.session.add(s)
        if sb:
            db.session.add(Contrat(
                tenant_id=t.id, salarie=s,
                type_contrat=request.form.get("type_contrat", "CDI"),
                date_debut=s.date_embauche, salaire_base=sb, poste=s.emploi, actif=True))
        # Archiver le journalier (historique préservé, retiré des listes actives).
        j.statut = "CONVERTI"
        db.session.flush()
        # Reporter l'affectation de site active vers le nouveau salarié.
        aff = AffectationSite.query.filter_by(
            tenant_id=t.id, journalier_id=j.id, actif=True).first()
        if aff:
            db.session.add(AffectationSite(
                tenant_id=t.id, site_id=aff.site_id, salarie_id=s.id, actif=True))
            aff.actif = False
        log_action("UPDATE", "journalier", j.id,
                   f"Journalier {j.nom_complet} converti en salarié (matricule {s.matricule})")
        log_action("CREATE", "salarie", s.id,
                   f"Salarié issu de la conversion du journalier {j.nom_complet}",
                   apres={"nom": s.nom, "prenom": s.prenom, "matricule": s.matricule,
                          "salaire_base": sb, "origine": "journalier"})
        db.session.commit()
        flash(f"{s.nom_complet} est désormais salarié (matricule {s.matricule}). "
              f"L'historique du journalier est conservé.", "success")
        return redirect(url_for("tenant.salarie_detail", id=s.id))

    return render_template("tenant/journalier_convertir.html",
                           journalier=j, categories=cats, tenant=t, suggestion=suggestion)




@bp.route("/journaliers")
@login_required
def journaliers():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    q    = request.args.get("q", "")
    page = request.args.get("page", 1, type=int)
    query = Journalier.query.filter_by(tenant_id=t.id)
    if q: query = query.filter(db.or_(Journalier.nom.ilike(f"%{q}%"), Journalier.prenom.ilike(f"%{q}%"), Journalier.profession.ilike(f"%{q}%")))
    pagination = query.order_by(Journalier.nom).paginate(page=page, per_page=25, error_out=False)
    _args = {k: v for k, v in request.args.items() if k != 'page'}
    _base = request.path + '?' + '&'.join(f'{k}={v}' for k, v in _args.items())
    _sep  = '&' if _args else '?'
    return render_template("tenant/journaliers.html",
        tenant=t, journaliers=pagination.items, pagination=pagination, q=q,
        pagination_base=_base + _sep)

@bp.route("/journaliers/nouveau", methods=["GET","POST"])
@login_required
def journalier_nouveau():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    # Vérifier quota dès le GET
    q = t.quota_employes_info
    if q["max"] and q["plein"]:
        flash(
            f"Limite atteinte — Plan « {t.plan.nom} » : {q['max']} employé(s) maximum "
            f"({q['salaries']} salarié(s) + {q['journaliers']} journalier(s)). "
            f"Passez au plan supérieur.", "error"
        )
        return redirect(url_for("tenant.journaliers"))
    if request.method == "POST":
        if not t.peut_ajouter_employe:
            flash(f"Limite atteinte ({t.plan.max_salaries} employés). Passez au plan supérieur.","error")
            return redirect(url_for("tenant.journaliers"))
        j = Journalier(tenant_id=t.id,
            nom=request.form["nom"].strip().upper(),
            prenom=request.form["prenom"].strip(),
            telephone=request.form.get("telephone","").strip(),
            profession=request.form.get("profession","").strip().upper(),
            taux_horaire=float(request.form.get("taux_horaire",0) or 0),
            type_paie=("MENSUEL" if request.form.get("type_paie")=="MENSUEL" else "JOURNALIER"),
            date_embauche=_parse_date(request.form.get("date_embauche")),
            date_debut=   _parse_date(request.form.get("date_debut")),
            date_fin=     _parse_date(request.form.get("date_fin")),
            nationalite=  request.form.get("nationalite","").strip() or None,
            mode_paiement=(request.form.get("mode_paiement","ESPECES") or "ESPECES").strip(),
            statut="ACTIF")
        db.session.add(j); db.session.commit()
        log_action("CREATE", "journalier", j.id,
                   f"Création journalier {j.nom_complet} ({j.type_paie})", apres=j.to_dict())
        db.session.commit()
        flash(f"Journalier {j.nom_complet} créé.", "success")
        return redirect(url_for("tenant.journaliers"))
    return render_template("tenant/journalier_form.html", tenant=t, journalier=None)

@bp.route("/journaliers/<int:id>")
@login_required
def journalier_detail(id):
    """Fiche détail d'un journalier avec historique de pointage."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    j = Journalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()

    # Feuilles de paie
    feuilles = FeuillePaieJournalier.query.filter_by(
        journalier_id=id, tenant_id=t.id
    ).order_by(FeuillePaieJournalier.date_fin.desc()).all()
    total_percu = sum(float(f.montant_brut or 0) for f in feuilles if f.statut == "PAYÉ")

    # ── Avances du journalier + déduction sur les feuilles ────────────────────
    avances = (AvanceJournalier.query.filter_by(journalier_id=id, tenant_id=t.id)
               .order_by(AvanceJournalier.date_avance.desc()).all())
    total_avances = round(sum(float(a.montant or 0) for a in avances), 2)
    avances_reste = round(sum(a.reste_a_regulariser for a in avances), 2)
    imput_av = _imputer_avances_journalier(feuilles, {id: avances})

    # Affectation site courante
    aff = AffectationSite.query.filter_by(
        journalier_id=id, tenant_id=t.id, actif=True).first()

    # ── Historique des pointages ──────────────────────────────────────────────
    nb_jours = request.args.get("nb_jours", type=int, default=30)
    nb_jours = min(max(nb_jours, 7), 90)
    date_fin   = datetime.now().date()
    date_debut = date_fin - timedelta(days=nb_jours - 1)

    pts_hist = Pointage.query.filter_by(tenant_id=t.id, journalier_id=id)        .filter(Pointage.date_pointage >= date_debut,
                Pointage.date_pointage <= date_fin)        .order_by(Pointage.date_pointage.desc()).all()

    nb_presences   = sum(1 for p in pts_hist if p.present)
    nb_absences    = sum(1 for p in pts_hist if p.absent)
    nb_non_pointes = nb_jours - len(pts_hist)
    h_normales_tot = round(sum(float(p.heures_normales or 0) for p in pts_hist if p.present), 1)
    h_sup_tot      = round(sum(float(p.heures_sup or 0) for p in pts_hist if p.present), 1)
    taux_presence  = round(nb_presences / (nb_presences + nb_absences) * 100
                           ) if (nb_presences + nb_absences) > 0 else 0

    return render_template("tenant/journalier_detail.html",
        journalier=j, tenant=t, feuilles=feuilles,
        total_percu=total_percu, aff=aff,
        avances=avances, total_avances=total_avances, avances_reste=avances_reste,
        imput_av=imput_av,
        pts_hist=pts_hist, nb_jours=nb_jours,
        nb_presences=nb_presences, nb_absences=nb_absences,
        nb_non_pointes=nb_non_pointes,
        h_normales_tot=h_normales_tot, h_sup_tot=h_sup_tot,
        taux_presence=taux_presence,
        date_debut_hist=date_debut, date_fin_hist=date_fin)

@bp.route("/journaliers/<int:id>/modifier", methods=["GET","POST"])
@login_required
def journalier_modifier(id):
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    j = Journalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if request.method == "POST":
        j.nom=request.form["nom"].strip().upper(); j.prenom=request.form["prenom"].strip()
        j.telephone=request.form.get("telephone","").strip()
        j.profession=request.form.get("profession","").strip().upper()
        j.taux_horaire=float(request.form.get("taux_horaire",0) or 0)
        j.type_paie=("MENSUEL" if request.form.get("type_paie")=="MENSUEL" else "JOURNALIER")
        j.date_embauche=_parse_date(request.form.get("date_embauche"))
        j.date_debut=   _parse_date(request.form.get("date_debut"))
        j.date_fin=     _parse_date(request.form.get("date_fin"))
        j.nationalite=  request.form.get("nationalite","").strip() or None
        j.mode_paiement=(request.form.get("mode_paiement","ESPECES") or "ESPECES").strip()
        j.statut=request.form.get("statut","ACTIF")
        # ── Affectation site ──────────────────────────────────────────────
        site_id = request.form.get("site_id", type=int)
        if site_id:
            aff_prev = AffectationSite.query.filter_by(
                journalier_id=j.id, tenant_id=t.id, actif=True).first()
            if aff_prev and aff_prev.site_id != site_id:
                aff_prev.actif    = False
                aff_prev.date_fin = date.today()
                aff_prev.motif    = "Réaffecté via formulaire journalier"
            if not aff_prev or aff_prev.site_id != site_id:
                db.session.add(AffectationSite(
                    tenant_id=t.id, site_id=site_id, journalier_id=j.id,
                    date_debut=date.today(), actif=True,
                    cree_par=current_user.email))
        elif request.form.get("retirer_site"):
            aff = AffectationSite.query.filter_by(
                journalier_id=j.id, tenant_id=t.id, actif=True).first()
            if aff:
                aff.actif    = False
                aff.date_fin = date.today()
                aff.motif    = "Retiré via formulaire journalier"
        db.session.commit()
        log_action("UPDATE", "journalier", j.id,
                   f"Modification journalier {j.nom_complet}", apres=j.to_dict())
        db.session.commit()
        flash("Journalier mis à jour.", "success")
        return redirect(url_for("tenant.journaliers"))
    aff_actuelle = AffectationSite.query.filter_by(
        journalier_id=id, tenant_id=t.id, actif=True).first()
    sites = Site.query.filter_by(tenant_id=t.id, actif=True).order_by(Site.nom).all()
    return render_template("tenant/journalier_form.html", tenant=t, journalier=j,
        sites=sites, aff_actuelle=aff_actuelle)

# ── Pointage ──────────────────────────────────────────────────────────────────
@bp.route("/pointage")
@login_required
def pointage():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    now = datetime.now()
    date_str = request.args.get("date", now.strftime("%Y-%m-%d"))
    try: date_sel = datetime.strptime(date_str, "%Y-%m-%d").date()
    except: date_sel = now.date()
    # ── Filtre par site ───────────────────────────────────────────────────────
    site_filtre_id = request.args.get("site_id", type=int)
    sites_list = Site.query.filter_by(tenant_id=t.id, actif=True).order_by(Site.nom).all()
    site_filtre = Site.query.filter_by(id=site_filtre_id, tenant_id=t.id).first() if site_filtre_id else None

    if site_filtre_id:
        # Salariés affectés à ce site
        ids_sal = [a.salarie_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_filtre_id, actif=True
        ).filter(AffectationSite.salarie_id.isnot(None)).all()]
        ids_jour = [a.journalier_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_filtre_id, actif=True
        ).filter(AffectationSite.journalier_id.isnot(None)).all()]
        salaries_list   = Salarie.query.filter(
            Salarie.tenant_id==t.id, Salarie.statut=="ACTIF",
            Salarie.id.in_(ids_sal)
        ).order_by(Salarie.nom).all()
        journaliers_list = Journalier.query.filter(
            Journalier.tenant_id==t.id, Journalier.statut=="ACTIF",
            Journalier.id.in_(ids_jour)
        ).order_by(Journalier.nom).all()
    else:
        salaries_list    = Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Salarie.nom).all()
        journaliers_list = Journalier.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Journalier.nom).all()

    pts_salaries    = {p.salarie_id:    p for p in Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_sel).filter(Pointage.salarie_id.isnot(None)).all()}
    pts_journaliers = {p.journalier_id: p for p in Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_sel).filter(Pointage.journalier_id.isnot(None)).all()}
    nb_presents_sal  = sum(1 for p in pts_salaries.values()    if p.present)
    nb_presents_jour = sum(1 for p in pts_journaliers.values() if p.present)
    nb_absents       = sum(1 for p in list(pts_salaries.values())+list(pts_journaliers.values()) if p.absent)
    lundi   = date_sel - timedelta(days=date_sel.weekday())
    semaine = [lundi + timedelta(days=i) for i in range(6)]

    # Affectation site de chaque travailleur pour affichage dans le pointage
    aff_sal  = {a.salarie_id:    a.site for a in AffectationSite.query.filter_by(tenant_id=t.id, actif=True).filter(AffectationSite.salarie_id.isnot(None)).all()}
    aff_jour = {a.journalier_id: a.site for a in AffectationSite.query.filter_by(tenant_id=t.id, actif=True).filter(AffectationSite.journalier_id.isnot(None)).all()}

    return render_template("tenant/pointage.html",
        tenant=t, date_sel=date_sel, semaine=semaine,
        date_hier=(date_sel  - timedelta(days=1)).strftime("%Y-%m-%d"),
        date_demain=(date_sel + timedelta(days=1)).strftime("%Y-%m-%d"),
        salaries=salaries_list, journaliers=journaliers_list,
        pts_salaries=pts_salaries, pts_journaliers=pts_journaliers,
        nb_presents_sal=nb_presents_sal, nb_presents_jour=nb_presents_jour,
        nb_absents=nb_absents, now=now,
        sites=sites_list, site_filtre=site_filtre,
        aff_sal=aff_sal, aff_jour=aff_jour)

@bp.route("/pointage/individuel", methods=["GET","POST"])
@login_required
def pointage_individuel():
    """Pointage d'un seul salarié ou journalier."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))

    date_str  = request.args.get("date", datetime.now().strftime("%Y-%m-%d"))
    type_w    = request.args.get("type", "sal")   # "sal" ou "jour"
    worker_id = request.args.get("id", type=int)

    try:   date_sel = datetime.strptime(date_str, "%Y-%m-%d").date()
    except: date_sel = datetime.now().date()

    if request.method == "POST":
        # Sauvegarder le pointage individuel
        date_p = datetime.strptime(
            request.form.get("date_pointage", date_str), "%Y-%m-%d").date()
        wtype  = request.form.get("worker_type", "sal")
        wid    = request.form.get("worker_id", type=int)

        # Validation stricte : le travailleur doit exister ET appartenir au tenant
        if wtype == "sal":
            worker_obj = Salarie.query.filter_by(id=wid, tenant_id=t.id).first()
        else:
            worker_obj = Journalier.query.filter_by(id=wid, tenant_id=t.id).first()
        if not worker_obj:
            flash("Travailleur introuvable ou non autorisé.", "error")
            return redirect(url_for("tenant.pointage"))

        present = request.form.get("present") == "1"
        absent  = not present
        def _hm(val):
            """Valider et retourner un horaire HH:MM ou None."""
            v = (val or "").strip()
            if not v: return None
            import re
            return v if re.match(r"^\d{1,2}:\d{2}$", v) else None

        def _diff_hm(debut, fin):
            """Calculer la différence en heures entre deux horaires HH:MM."""
            try:
                h1,m1 = map(int, debut.split(":")); h2,m2 = map(int, fin.split(":"))
                diff = (h2*60+m2 - h1*60-m1) / 60
                return max(0.0, round(diff, 2))
            except: return 0.0

        # Horaires saisis
        em = _hm(request.form.get("entree_matin"))
        sm = _hm(request.form.get("sortie_matin"))
        ea = _hm(request.form.get("entree_apmidi"))
        sa = _hm(request.form.get("sortie_apmidi"))
        es = _hm(request.form.get("entree_sup"))
        ss = _hm(request.form.get("sortie_sup"))

        # Calcul auto des heures normales depuis les horaires
        h_normales_auto = 0.0
        if em and sm: h_normales_auto += _diff_hm(em, sm)
        if ea and sa: h_normales_auto += _diff_hm(ea, sa)
        h_normales_man = float(request.form.get("heures_normales", 0) or 0)
        # Priorité aux horaires si saisis, sinon saisie manuelle
        heures_normales_final = round(h_normales_auto, 2) if h_normales_auto > 0 else h_normales_man or 8

        # Heures sup depuis horaires
        h_sup_horaire = _diff_hm(es, ss) if (es and ss) else 0.0

        type_jour = request.form.get("type_jour", "NORMAL")

        # Reclasser les heures selon le type de jour
        h_sup_10_man = float(request.form.get("heures_sup_10",0) or 0)
        h_sup_30_man = float(request.form.get("heures_sup_30",0) or 0)
        h_sup_30b_man = float(request.form.get("heures_sup_30b",0) or 0)
        h_sup_40_man = float(request.form.get("heures_sup_40",0) or 0)
        h_sup_70_man = float(request.form.get("heures_sup_70",0) or 0)
        h_sup_30b_final = 0
        _conv_t = (t.convention or "").upper()

        if type_jour == "DIMANCHE":
            if _conv_t in ("PETROLE", "INDUSTRIE"):
                # Pétrole / Industrie : dimanche de jour → case 30b ; la nuit est
                # recalculée par la ventilation mensuelle.
                h_sup_30b_final = round(h_sup_horaire + heures_normales_final, 2)
                h_sup_70_final = 0
            else:
                # Dimanche travaillé : intégralité en +70% (réglementation BTP Gabon)
                h_sup_70_final = round(h_sup_horaire + heures_normales_final, 2)
            h_sup_10_final = 0; h_sup_30_final = 0; h_sup_40_final = 0
            heures_normales_final = 0
        elif type_jour == "FERIE":
            if _conv_t in ("PETROLE", "INDUSTRIE"):
                h_sup_30b_final = round(h_sup_horaire + heures_normales_final, 2)
                h_sup_70_final = 0
            else:
                # Tout va en +70% (jour férié)
                h_sup_70_final = round(h_sup_horaire + heures_normales_final, 2)
            h_sup_10_final = 0; h_sup_30_final = 0; h_sup_40_final = 0
            heures_normales_final = 0
        elif type_jour in ("CHOME_PAYE", "CHOME_RECUPERABLE"):
            # Présent mais jour chômé → heures normales conservées, pas de sup
            h_sup_10_final = 0; h_sup_30_final = 0; h_sup_40_final = 0; h_sup_70_final = 0
        else:
            # NORMAL
            h_sup_10_final = round(h_sup_horaire, 2) if h_sup_horaire > 0 else h_sup_10_man
            h_sup_30_final = h_sup_30_man
            h_sup_30b_final = h_sup_30b_man
            h_sup_40_final = h_sup_40_man
            h_sup_70_final = h_sup_70_man

        kwargs = dict(
            heures_normales = heures_normales_final,
            motif_absence   = request.form.get("motif_absence","") if absent else None,
            entree_matin    = em, sortie_matin  = sm,
            entree_apmidi   = ea, sortie_apmidi = sa,
            entree_sup      = es, sortie_sup    = ss,
            type_jour       = type_jour,
        )
        if wtype == "sal":
            kwargs.update(dict(
                heures_sup_10 = h_sup_10_final,
                heures_sup_30 = h_sup_30_final,
                heures_sup_30b = h_sup_30b_final,
                heures_sup_40 = h_sup_40_final,
                heures_sup_70 = h_sup_70_final,
            ))
            pt = Pointage.query.filter_by(
                tenant_id=t.id, date_pointage=date_p, salarie_id=wid).first()
            if not pt:
                pt = Pointage(tenant_id=t.id, date_pointage=date_p, salarie_id=wid)
                db.session.add(pt)
        else:
            # Journalier : gérer heures_sup selon type_jour
            if type_jour in ("DIMANCHE", "FERIE"):
                kwargs["heures_sup"]      = round(h_sup_horaire + heures_normales_final, 2)
                kwargs["heures_normales"] = 0
            else:
                kwargs["heures_sup"] = float(request.form.get("heures_sup",0) or 0)
            pt = Pointage.query.filter_by(
                tenant_id=t.id, date_pointage=date_p, journalier_id=wid).first()
            if not pt:
                pt = Pointage(tenant_id=t.id, date_pointage=date_p, journalier_id=wid)
                db.session.add(pt)
        pt.present = present
        pt.absent  = absent
        for k, v in kwargs.items():
            setattr(pt, k, v)
        db.session.commit()
        worker_name = worker_obj.nom_complet
        flash(f"✅ Pointage de {worker_name} enregistré.", "success")
        # Rester sur la même page pour pointer la personne suivante
        redir = request.form.get("next_url") or f"/pointage/individuel?date={date_p}&type={wtype}&id={wid}"
        return redirect(redir)

    # GET : charger le travailleur sélectionné
    worker = pt_existant = None
    historique_30j = []
    stats_30j = {"presences": 0, "absences": 0, "h_normales": 0.0,
                 "h_sup": 0.0, "taux": 0}

    if worker_id:
        if type_w == "sal":
            worker = Salarie.query.filter_by(id=worker_id, tenant_id=t.id).first()
            pt_existant = Pointage.query.filter_by(
                tenant_id=t.id, date_pointage=date_sel, salarie_id=worker_id).first()
            if worker:
                date_debut_30 = date_sel - timedelta(days=29)
                historique_30j = Pointage.query.filter_by(
                    tenant_id=t.id, salarie_id=worker_id
                ).filter(
                    Pointage.date_pointage >= date_debut_30,
                    Pointage.date_pointage <= date_sel
                ).order_by(Pointage.date_pointage.desc()).all()
        else:
            worker = Journalier.query.filter_by(id=worker_id, tenant_id=t.id).first()
            pt_existant = Pointage.query.filter_by(
                tenant_id=t.id, date_pointage=date_sel, journalier_id=worker_id).first()
            if worker:
                date_debut_30 = date_sel - timedelta(days=29)
                historique_30j = Pointage.query.filter_by(
                    tenant_id=t.id, journalier_id=worker_id
                ).filter(
                    Pointage.date_pointage >= date_debut_30,
                    Pointage.date_pointage <= date_sel
                ).order_by(Pointage.date_pointage.desc()).all()

        # Stats sur les 30 jours
        if historique_30j:
            nb_p = sum(1 for p in historique_30j if p.present)
            nb_a = sum(1 for p in historique_30j if p.absent)
            hn   = round(sum(float(p.heures_normales or 0) for p in historique_30j if p.present), 1)
            if type_w == "sal":
                hs = round(sum(
                    float(p.heures_sup_10 or 0) + float(p.heures_sup_30 or 0) +
                    float(p.heures_sup_40 or 0) + float(p.heures_sup_70 or 0)
                    for p in historique_30j if p.present), 1)
            else:
                hs = round(sum(float(p.heures_sup or 0) for p in historique_30j if p.present), 1)
            total_ptg = nb_p + nb_a
            stats_30j = {
                "presences": nb_p, "absences": nb_a,
                "h_normales": hn, "h_sup": hs,
                "taux": round(nb_p / total_ptg * 100) if total_ptg > 0 else 0
            }

    # Listes pour la recherche
    salaries_list    = Salarie.query.filter_by(
        tenant_id=t.id, statut="ACTIF").order_by(Salarie.nom).all()
    journaliers_list = Journalier.query.filter_by(
        tenant_id=t.id, statut="ACTIF").order_by(Journalier.nom).all()

    return render_template("tenant/pointage_individuel.html",
        tenant=t, date_sel=date_sel,
        date_hier=(date_sel - timedelta(days=1)).strftime("%Y-%m-%d"),
        date_demain=(date_sel + timedelta(days=1)).strftime("%Y-%m-%d"),
        type_w=type_w, worker=worker, pt_existant=pt_existant,
        historique_30j=historique_30j, stats_30j=stats_30j,
        salaries=salaries_list, journaliers=journaliers_list,
        now=datetime.now())

@bp.route("/pointage/supprimer/<int:ptg_id>", methods=["POST"])
@login_required
def pointage_supprimer(ptg_id):
    """Supprimer un pointage individuel."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    pt = Pointage.query.filter_by(id=ptg_id, tenant_id=t.id).first_or_404()
    # Mémoriser le contexte pour rediriger au bon endroit
    next_url = request.form.get("next_url", "/pointage")
    date_str = str(pt.date_pointage)
    type_w   = "sal" if pt.salarie_id else "jour"
    wid      = pt.salarie_id or pt.journalier_id
    db.session.delete(pt)
    db.session.commit()
    log_action("DELETE", "pointage", ptg_id, f"Suppression pointage du {date_str}")
    db.session.commit()
    flash("🗑️ Pointage supprimé.", "success")
    return redirect(next_url or f"/pointage/individuel?date={date_str}&type={type_w}&id={wid}")

@bp.route("/pointage/sauvegarder", methods=["POST"])
@login_required
def pointage_sauvegarder():
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    date_str = request.form.get("date_pointage")
    try: date_p = datetime.strptime(date_str, "%Y-%m-%d").date()
    except: flash("Date invalide.", "error"); return redirect(url_for("tenant.pointage"))
    nb = 0
    sel_sal  = {v for k,v in request.form.items() if k.startswith("sel_sal_")}
    sel_jour = {v for k,v in request.form.items() if k.startswith("sel_jour_")}
    for key, val in request.form.items():
        if key.startswith("sal_present_"):
            sid_str = key.replace("sal_present_","")
            if sid_str not in sel_sal: continue
            sid = int(sid_str); present = val == "1"; absent = not present
            pt = Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_p, salarie_id=sid).first()
            if not pt: pt = Pointage(tenant_id=t.id, date_pointage=date_p, salarie_id=sid); db.session.add(pt)
            pt.present=present; pt.absent=absent
            pt.heures_normales = float(request.form.get(f"sal_heures_{sid}", 8) or 8)
            pt.heures_sup_10   = float(request.form.get(f"sal_sup10_{sid}", 0) or 0)
            pt.heures_sup_30   = float(request.form.get(f"sal_sup30_{sid}", 0) or 0)
            pt.heures_sup_30b  = float(request.form.get(f"sal_sup30b_{sid}", 0) or 0)
            pt.heures_sup_40   = float(request.form.get(f"sal_sup40_{sid}", 0) or 0)
            pt.heures_sup_70   = float(request.form.get(f"sal_sup70_{sid}", 0) or 0)
            pt.motif_absence   = request.form.get(f"sal_motif_{sid}", "") if absent else None
            nb += 1
        if key.startswith("jour_present_"):
            jid_str = key.replace("jour_present_","")
            if jid_str not in sel_jour: continue
            jid = int(jid_str); present = val == "1"; absent = not present
            pt = Pointage.query.filter_by(tenant_id=t.id, date_pointage=date_p, journalier_id=jid).first()
            if not pt: pt = Pointage(tenant_id=t.id, date_pointage=date_p, journalier_id=jid); db.session.add(pt)
            pt.present=present; pt.absent=absent
            pt.heures_normales = float(request.form.get(f"jour_heures_{jid}", 8) or 8)
            pt.heures_sup      = float(request.form.get(f"jour_sup_{jid}", 0) or 0)
            pt.motif_absence   = request.form.get(f"jour_motif_{jid}", "") if absent else None
            nb += 1
    db.session.commit()
    if nb:
        log_action("UPDATE", "pointage", None,
                   f"Pointage du {date_p.strftime('%d/%m/%Y')} enregistré ({nb} ligne(s))")
        db.session.commit()
    flash(f"Pointage du {date_p.strftime('%d/%m/%Y')} sauvegardé ({nb} lignes).", "success")
    return redirect(url_for("tenant.pointage", date=date_str))

# ── Paie journaliers ──────────────────────────────────────────────────────────
@bp.route("/journaliers/paie")
@login_required
def journaliers_paie():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))

    # ── Filtre par site ──────────────────────────────────────────────────────
    site_filtre_id = request.args.get("site_id", type=int)
    sites_list     = Site.query.filter_by(tenant_id=t.id, actif=True).order_by(Site.nom).all()
    site_filtre    = Site.query.filter_by(id=site_filtre_id, tenant_id=t.id).first() if site_filtre_id else None
    statut_filtre  = request.args.get("statut", "")

    if site_filtre_id:
        ids_jour = [a.journalier_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_filtre_id, actif=True
        ).filter(AffectationSite.journalier_id.isnot(None)).all()]
        journaliers_list = Journalier.query.filter(
            Journalier.tenant_id==t.id, Journalier.statut=="ACTIF",
            Journalier.id.in_(ids_jour)
        ).order_by(Journalier.nom).all()
    else:
        journaliers_list = Journalier.query.filter_by(tenant_id=t.id, statut="ACTIF").order_by(Journalier.nom).all()

    # ── Feuilles filtrées ────────────────────────────────────────────────────
    q_feuilles = FeuillePaieJournalier.query.filter_by(tenant_id=t.id)
    if site_filtre_id:
        ids_jour_all = [a.journalier_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_filtre_id
        ).filter(AffectationSite.journalier_id.isnot(None)).all()]
        q_feuilles = q_feuilles.filter(FeuillePaieJournalier.journalier_id.in_(ids_jour_all))
    if statut_filtre:
        q_feuilles = q_feuilles.filter_by(statut=statut_filtre)
    page_f      = request.args.get("page", 1, type=int)
    # KPIs sur toutes les feuilles (sans pagination)
    feuilles_tous    = q_feuilles.order_by(FeuillePaieJournalier.date_fin.desc()).all()
    total_en_attente = sum(float(f.montant_brut or 0) for f in feuilles_tous if f.statut == "EN_ATTENTE")
    total_paye       = sum(float(f.montant_brut or 0) for f in feuilles_tous if f.statut == "PAYÉ")
    nb_en_attente    = sum(1 for f in feuilles_tous if f.statut == "EN_ATTENTE")
    q_feuilles = q_feuilles.options(joinedload(FeuillePaieJournalier.journalier))
    pagination_f     = q_feuilles.order_by(FeuillePaieJournalier.date_fin.desc()).paginate(page=page_f, per_page=25, error_out=False)
    feuilles         = pagination_f.items

    # Déduction des avances pour les feuilles affichées
    imput_av = _imputer_avances_journalier(
        feuilles, _avances_par_journalier(t.id, {f.journalier_id for f in feuilles}))

    # Affectation site de chaque journalier (pour affichage dans la liste)
    aff_jour = {a.journalier_id: a.site for a in AffectationSite.query.filter_by(
        tenant_id=t.id, actif=True
    ).filter(AffectationSite.journalier_id.isnot(None)).all()}

    _args = {k: v for k, v in request.args.items() if k != 'page'}
    _base = request.path + '?' + '&'.join(f'{k}={v}' for k, v in _args.items())
    _sep  = '&' if _args else '?'
    return render_template("tenant/journaliers_paie.html",
        tenant=t, feuilles=feuilles, journaliers=journaliers_list,
        sites=sites_list, site_filtre=site_filtre, statut_filtre=statut_filtre,
        total_en_attente=total_en_attente, total_paye=total_paye,
        nb_en_attente=nb_en_attente, aff_jour=aff_jour, imput_av=imput_av,
        pagination=pagination_f, pagination_base=_base + _sep,
        now=datetime.now(), today=date.today().isoformat())

@bp.route("/journaliers/paie/generer", methods=["POST"])
@login_required
def journaliers_paie_generer():
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    date_debut = _parse_date(request.form.get("date_debut"))
    date_fin   = _parse_date(request.form.get("date_fin"))
    if not date_debut or not date_fin: flash("Dates invalides.", "error"); return redirect(url_for("tenant.journaliers_paie"))
    site_id      = request.form.get("site_id", type=int)
    taux_custom  = {}  # taux personnalisés par journalier
    heures_custom = {} # heures manuelles par journalier
    ids_coches   = request.form.getlist("journalier_ids")

    for key, val in request.form.items():
        if key.startswith("taux_") and val:
            try: taux_custom[int(key[5:])] = float(val)
            except: pass
        if key.startswith("heures_") and val:
            try: heures_custom[int(key[7:])] = float(val)
            except: pass

    if ids_coches:
        journaliers_a_payer = Journalier.query.filter(
            Journalier.tenant_id==t.id,
            Journalier.id.in_([int(i) for i in ids_coches])
        ).all()
    elif site_id:
        ids_site = [a.journalier_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_id, actif=True
        ).filter(AffectationSite.journalier_id.isnot(None)).all()]
        journaliers_a_payer = Journalier.query.filter(
            Journalier.tenant_id==t.id, Journalier.statut=="ACTIF",
            Journalier.id.in_(ids_site)
        ).all()
    else:
        journaliers_a_payer = Journalier.query.filter_by(tenant_id=t.id, statut="ACTIF").all()

    nb = 0
    for j in journaliers_a_payer:
        if str(j.id) not in ids_coches and ids_coches:
            continue
        taux = taux_custom.get(j.id, float(j.taux_horaire or 0))
        if j.id in heures_custom:
            total_h  = heures_custom[j.id]
            nb_jours = 1  # heures manuelles = considéré comme 1 entrée
            h_norm   = total_h; h_sup = 0.0   # pas de détail sur saisie manuelle
        else:
            pts = Pointage.query.filter_by(tenant_id=t.id, journalier_id=j.id)                  .filter(Pointage.date_pointage>=date_debut, Pointage.date_pointage<=date_fin,
                          Pointage.present==True).all()
            h_norm   = sum(float(p.heures_normales or 0) for p in pts)
            h_sup    = sum(float(p.heures_sup or 0) for p in pts)
            total_h  = h_norm + h_sup
            nb_jours = len(pts)
        if total_h <= 0 and nb_jours == 0: continue
        if FeuillePaieJournalier.query.filter_by(
            tenant_id=t.id, journalier_id=j.id,
            date_debut=date_debut, date_fin=date_fin).first(): continue
        db.session.add(FeuillePaieJournalier(
            tenant_id=t.id, journalier_id=j.id,
            date_debut=date_debut, date_fin=date_fin,
            nb_jours=nb_jours, total_heures=total_h,
            heures_normales=h_norm, heures_sup=h_sup,
            taux_horaire=taux, montant_brut=round(total_h*taux, 2),
            statut="EN_ATTENTE"))
        nb += 1
    db.session.commit()
    flash(f"{nb} feuille(s) générée(s).", "success")
    redirect_url = url_for("tenant.journaliers_paie")
    if site_id:
        redirect_url += f"?site_id={site_id}"
    return redirect(redirect_url)

@bp.route("/journaliers/paie/generer-mois", methods=["POST"])
@login_required
def journaliers_paie_generer_mois():
    """Génère la paie de FIN DE MOIS pour tous les journaliers de type MENSUEL.

    La période couvre le mois entier (1er → dernier jour). Le montant est calculé
    à partir des pointages présents du mois × taux horaire.
    """
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    import calendar as _cal
    mois  = request.form.get("mois", type=int)
    annee = request.form.get("annee", type=int)
    if not mois or not annee:
        flash("Mois/année manquant.", "error")
        return redirect(url_for("tenant.journaliers_paie"))
    date_debut = date(annee, mois, 1)
    date_fin   = date(annee, mois, _cal.monthrange(annee, mois)[1])

    mensuels = Journalier.query.filter_by(
        tenant_id=t.id, statut="ACTIF", type_paie="MENSUEL").all()
    nb = 0; nb_existant = 0
    for j in mensuels:
        if FeuillePaieJournalier.query.filter_by(
            tenant_id=t.id, journalier_id=j.id,
            date_debut=date_debut, date_fin=date_fin).first():
            nb_existant += 1
            continue
        pts = Pointage.query.filter_by(tenant_id=t.id, journalier_id=j.id).filter(
            Pointage.date_pointage >= date_debut,
            Pointage.date_pointage <= date_fin,
            Pointage.present == True).all()
        total_h  = sum(float(p.heures_normales or 0) + float(p.heures_sup or 0) for p in pts)
        h_norm   = sum(float(p.heures_normales or 0) for p in pts)
        h_sup    = sum(float(p.heures_sup or 0) for p in pts)
        nb_jours = len(pts)
        if total_h <= 0 and nb_jours == 0:
            continue
        taux = float(j.taux_horaire or 0)
        montant_exact = total_h * taux
        # Arrondi au millier supérieur UNIQUEMENT si l'utilisateur l'a demandé.
        if request.form.get("arrondi_millier") in ("1", "on", "true"):
            from calculs_paie import arrondi_millier_superieur
            brut = arrondi_millier_superieur(montant_exact)
        else:
            brut = round(montant_exact)   # montant exact au franc — l'utilisateur ajuste s'il veut
        db.session.add(FeuillePaieJournalier(
            tenant_id=t.id, journalier_id=j.id,
            date_debut=date_debut, date_fin=date_fin,
            nb_jours=nb_jours, total_heures=total_h,
            heures_normales=h_norm, heures_sup=h_sup,
            taux_horaire=taux, montant_brut=brut,
            statut="EN_ATTENTE"))
        nb += 1
    db.session.commit()
    if nb:
        log_action("CREATE", "feuille_journalier", None,
                   f"Génération paie mensuelle : {nb} feuille(s) ({mois:02d}/{annee})")
        db.session.commit()
    msg = f"{nb} feuille(s) mensuelle(s) générée(s) pour {mois:02d}/{annee}."
    if nb_existant:
        msg += f" {nb_existant} déjà existante(s) ignorée(s)."
    if not mensuels:
        msg = "Aucun journalier de type « Mensuel » n'est défini."
    flash(msg, "success" if mensuels else "error")
    return redirect(url_for("tenant.journaliers_paie"))


def _avances_par_journalier(tenant_id, journalier_ids=None):
    """Charge les avances et les regroupe par journalier : {jid: [AvanceJournalier]}."""
    from collections import defaultdict
    q = AvanceJournalier.query.filter_by(tenant_id=tenant_id)
    if journalier_ids is not None:
        if not journalier_ids:
            return {}
        q = q.filter(AvanceJournalier.journalier_id.in_(list(journalier_ids)))
    par_j = defaultdict(list)
    for a in q.order_by(AvanceJournalier.date_avance).all():
        par_j[a.journalier_id].append(a)
    return par_j


def _imputer_avances_journalier(feuilles, avances_par_journalier):
    """Déduit les avances des feuilles de paie, par journalier.

    Les feuilles PAYÉ utilisent le montant d'avance figé (`avance_deduite`).
    Pour les feuilles EN_ATTENTE, l'encours d'avances non régularisées du
    journalier est réparti sur ses feuilles (la plus ancienne d'abord), plafonné
    au montant de chaque feuille.

    Renvoie {feuille_id: {"avance": x, "net": y}} (en XAF).
    """
    from collections import defaultdict
    encours = {}
    for jid, avs in avances_par_journalier.items():
        encours[jid] = round(sum(max(0.0, float(a.montant or 0) - float(a.montant_regularise or 0))
                                 for a in avs), 2)
    res = {}
    par_j = defaultdict(list)
    for f in feuilles:
        par_j[f.journalier_id].append(f)
    for jid, fs in par_j.items():
        rem = encours.get(jid, 0.0)
        for f in fs:                       # feuilles déjà payées : déduction figée
            if f.statut == "PAYÉ":
                ded = float(f.avance_deduite or 0)
                res[f.id] = {"avance": ded, "net": round(float(f.montant_a_payer) - ded, 2)}
        for f in sorted([f for f in fs if f.statut != "PAYÉ"],
                        key=lambda x: (x.date_fin, x.id)):
            base = float(f.montant_a_payer)
            ded  = round(min(rem, base), 2)
            rem  = round(rem - ded, 2)
            res[f.id] = {"avance": ded, "net": round(base - ded, 2)}
    return res


@bp.route("/journaliers/<int:id>/avances/nouvelle", methods=["POST"])
@login_required
def journalier_avance_nouvelle(id):
    """Enregistre une avance versée à un journalier."""
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    j = Journalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    try:
        montant = float(request.form.get("montant", 0) or 0)
    except (TypeError, ValueError):
        montant = 0
    if montant <= 0:
        flash("Le montant de l'avance doit être supérieur à zéro.", "error")
        return redirect(url_for("tenant.journalier_detail", id=id))
    aff = AffectationSite.query.filter_by(journalier_id=id, tenant_id=t.id, actif=True).first()
    av = AvanceJournalier(
        tenant_id=t.id, journalier_id=j.id,
        site_id=(aff.site_id if aff else None),
        montant=montant,
        date_avance=_parse_date(request.form.get("date_avance", "")) or datetime.now().date(),
        mode_paiement=(request.form.get("mode_paiement", "ESPECES") or "ESPECES").strip(),
        reference=(request.form.get("reference", "") or "").strip() or None,
        motif=(request.form.get("motif", "") or "").strip() or None)
    db.session.add(av); db.session.commit()
    log_action("CREATE", "avance_journalier", av.id,
               f"Avance {montant:,.0f} F — {j.nom_complet}".replace(",", " "),
               apres=av.to_dict())
    db.session.commit()
    flash(f"Avance de {montant:,.0f} F enregistrée pour {j.nom_complet}.".replace(",", " "), "success")
    return redirect(url_for("tenant.journalier_detail", id=id))


@bp.route("/journaliers/avances/<int:aid>/modifier", methods=["POST"])
@login_required
def journalier_avance_modifier(aid):
    """Modifie une avance, tant qu'elle n'est ni validée ni déjà déduite."""
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    av = AvanceJournalier.query.filter_by(id=aid, tenant_id=t.id).first_or_404()
    jid = av.journalier_id
    if not av.est_modifiable:
        flash("Cette avance est validée ou déjà déduite : elle n'est plus modifiable.", "error")
        return redirect(url_for("tenant.journalier_detail", id=jid))
    try:
        montant = float(request.form.get("montant", av.montant) or av.montant)
    except (TypeError, ValueError):
        montant = float(av.montant)
    if montant <= 0:
        flash("Le montant de l'avance doit être supérieur à zéro.", "error")
        return redirect(url_for("tenant.journalier_detail", id=jid))
    av.montant = montant
    d = _parse_date(request.form.get("date_avance", ""))
    if d:
        av.date_avance = d
    av.motif = (request.form.get("motif", "") or "").strip() or None
    db.session.commit()
    log_action("UPDATE", "avance_journalier", av.id,
               f"Modification avance — {av.journalier.nom_complet}", apres=av.to_dict())
    db.session.commit()
    flash("Avance modifiée.", "success")
    return redirect(url_for("tenant.journalier_detail", id=jid))


@bp.route("/journaliers/avances/<int:aid>/valider", methods=["POST"])
@login_required
def journalier_avance_valider(aid):
    """Valide une avance : elle devient figée (plus modifiable ni supprimable)."""
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    av = AvanceJournalier.query.filter_by(id=aid, tenant_id=t.id).first_or_404()
    av.statut = "VALIDEE"
    db.session.commit()
    log_action("VALIDATE", "avance_journalier", av.id,
               f"Validation avance {float(av.montant):,.0f} F — {av.journalier.nom_complet}".replace(",", " "))
    db.session.commit()
    flash("Avance validée — elle est désormais figée.", "success")
    return redirect(url_for("tenant.journalier_detail", id=av.journalier_id))


@bp.route("/journaliers/avances/<int:aid>/supprimer", methods=["POST"])
@login_required
def journalier_avance_supprimer(aid):
    """Supprime une avance (uniquement si non validée et non encore déduite)."""
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    av = AvanceJournalier.query.filter_by(id=aid, tenant_id=t.id).first_or_404()
    jid = av.journalier_id
    if not av.est_modifiable:
        flash("Avance validée ou déjà déduite : suppression impossible.", "error")
        return redirect(url_for("tenant.journalier_detail", id=jid))
    db.session.delete(av); db.session.commit()
    log_action("DELETE", "avance_journalier", aid,
               f"Suppression avance {float(av.montant):,.0f} F".replace(",", " "), avant=av.to_dict())
    db.session.commit()
    flash("Avance supprimée.", "success")
    return redirect(url_for("tenant.journalier_detail", id=jid))


@bp.route("/journaliers/paie/<int:id>/payer", methods=["POST"])
@login_required
def journalier_payer(id):
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    f = FeuillePaieJournalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    f.montant_brut = f.montant_a_payer   # = montant enregistré (arrondi éventuel déjà appliqué à la génération)
    # Déduction des avances non régularisées du journalier (plus ancienne d'abord)
    avs = (AvanceJournalier.query.filter_by(tenant_id=t.id, journalier_id=f.journalier_id)
           .order_by(AvanceJournalier.date_avance).all())
    a_deduire = round(min(
        sum(a.reste_a_regulariser for a in avs),
        float(f.montant_a_payer)), 2)
    f.avance_deduite = a_deduire
    reste = a_deduire
    for a in avs:
        if reste <= 0:
            break
        dispo = a.reste_a_regulariser
        if dispo <= 0:
            continue
        pris = min(dispo, reste)
        a.montant_regularise = float(a.montant_regularise or 0) + pris
        reste = round(reste - pris, 2)
    mode = (request.form.get("mode_paiement") or "").strip()
    if mode not in ("ESPECES", "VIREMENT"):
        mode = (f.journalier.mode_paiement if f.journalier else "ESPECES") or "ESPECES"
    f.mode_paiement = mode
    # Date de paiement : celle choisie dans le formulaire, sinon aujourd'hui.
    date_pmt = date.today()
    _df = (request.form.get("date_paiement") or "").strip()
    if _df:
        try:
            date_pmt = datetime.strptime(_df, "%Y-%m-%d").date()
        except ValueError:
            date_pmt = date.today()
    f.statut = "PAYÉ"; f.date_paiement = date_pmt; db.session.commit()
    net = float(f.montant_brut) - a_deduire
    log_action("PAY", "feuille_journalier", f.id,
               f"Paiement {f.journalier.nom_complet} — net {net:,.0f} F"
               + (f" (avances {a_deduire:,.0f} F)" if a_deduire else ""),
               apres={"montant_brut": float(f.montant_brut), "avance_deduite": a_deduire,
                      "net": net, "periode": f"{f.date_debut}→{f.date_fin}"})
    db.session.commit()
    # Interconnexion Caisse : proposer la sortie correspondante (ne bloque jamais).
    try:
        from interco_caisse import proposer_ecriture
        nom = f.journalier.nom_complet if f.journalier else "journalier"
        proposer_ecriture(
            t, source_ref=f"journalier-{f.id}",
            montant=float(net),
            motif=f"Paie journalier {nom} — {f.date_debut}→{f.date_fin}",
            compte_suggere="6611", date_operation=f.date_paiement)
    except Exception as _e:
        current_app.logger.warning(f"[COMPTA] Proposition d'écriture échouée (journalier {f.id}) : {_e}")
    if a_deduire > 0:
        flash(f"Paiement de {f.journalier.nom_complet} enregistré "
              f"(net {net:,.0f} F après {a_deduire:,.0f} F d'avances).".replace(",", " "), "success")
    else:
        flash(f"Paiement de {f.journalier.nom_complet} enregistré.", "success")
    return redirect(url_for("tenant.journaliers_paie"))

@bp.route("/journaliers/paie/<int:id>/modifier", methods=["POST"])
@login_required
def journalier_feuille_modifier(id):
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    f = FeuillePaieJournalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    f.montant_brut = float(request.form.get("montant_brut", f.montant_brut) or f.montant_brut)
    f.observation  = request.form.get("observation", "").strip()
    db.session.commit(); flash("Feuille modifiée.", "success")
    return redirect(url_for("tenant.journaliers_paie"))

@bp.route("/journaliers/paie/<int:id>/supprimer", methods=["POST"])
@login_required
def journalier_feuille_supprimer(id):
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    f = FeuillePaieJournalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    db.session.delete(f); db.session.commit()
    flash("Feuille supprimée.", "success")
    return redirect(url_for("tenant.journaliers_paie"))

@bp.route("/journaliers/pointage/modele")
@login_required
def journaliers_pointage_modele():
    """Génère un modèle Excel (grille mensuelle) pré-rempli pour l'import de
    pointages journaliers. Une ligne par journalier du site choisi, une colonne
    par jour du mois. Dimanches et fériés grisés. L'utilisateur saisit les
    heures travaillées ; il téléverse ensuite le fichier via /importer.
    Params : site_id (optionnel), mois (YYYY-MM, défaut = mois en cours).
    """
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))

    import calendar as _cal
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from io import BytesIO

    # ── Période demandée ───────────────────────────────────────────────
    mois_str = request.args.get("mois", "")  # "2026-07"
    try:
        annee, mois = (int(x) for x in mois_str.split("-"))
        date(annee, mois, 1)  # validation
    except (ValueError, TypeError):
        today = datetime.now()
        annee, mois = today.year, today.month
    nb_jours = _cal.monthrange(annee, mois)[1]

    MOIS_FR = ["", "Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet",
               "Août", "Septembre", "Octobre", "Novembre", "Décembre"]
    mois_nom = f"{MOIS_FR[mois]} {annee}"

    # ── Journaliers du site (ou tous les actifs si aucun site) ─────────
    site_id = request.args.get("site_id", type=int)
    site_obj = None
    if site_id:
        site_obj = Site.query.filter_by(id=site_id, tenant_id=t.id).first()
        ids_site = [a.journalier_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_id, actif=True
        ).filter(AffectationSite.journalier_id.isnot(None)).all()]
        journaliers = Journalier.query.filter(
            Journalier.tenant_id == t.id, Journalier.statut == "ACTIF",
            Journalier.id.in_(ids_site)
        ).order_by(Journalier.nom).all() if ids_site else []
    else:
        journaliers = Journalier.query.filter_by(
            tenant_id=t.id, statut="ACTIF").order_by(Journalier.nom).all()

    # ── Styles ─────────────────────────────────────────────────────────
    VERT, BLANC, GRIS_WE = "0F3D36", "FFFFFF", "F3F4F6"
    thin = Side(style="thin", color="D1D5DB")
    bord = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_titre  = Font(name="Arial", size=12, bold=True, color=BLANC)
    f_entete = Font(name="Arial", size=9,  bold=True, color=BLANC)
    f_normal = Font(name="Arial", size=10)
    fill_vert = PatternFill("solid", fgColor=VERT)
    fill_we   = PatternFill("solid", fgColor=GRIS_WE)
    center = Alignment(horizontal="center", vertical="center")

    wb = Workbook()
    ws = wb.active
    ws.title = f"Pointage {MOIS_FR[mois][:3]} {annee}"
    ws.sheet_view.showGridLines = False

    n_cols = 2 + nb_jours  # ID + Nom + jours
    # Titre
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n_cols)
    tc = ws.cell(row=1, column=1,
                 value=f"POINTAGE JOURNALIERS — {mois_nom}"
                       + (f"  ·  {site_obj.nom}" if site_obj else "  ·  Tous sites"))
    tc.font = f_titre; tc.fill = fill_vert
    tc.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[1].height = 24

    # Légende
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n_cols)
    lc = ws.cell(row=2, column=1,
        value="Saisissez le nombre d'HEURES travaillées par jour. "
              "Vide = jour ignoré (non modifié). 0 = absent. "
              "Au-delà de 8h : le surplus compte en heures supplémentaires. "
              "Ne modifiez PAS la colonne ID.")
    lc.font = Font(name="Arial", size=9, italic=True, color="6B7280")
    lc.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 40

    # En-têtes (ligne 3)
    for col, lbl in ((1, "ID"), (2, "Nom journalier")):
        c = ws.cell(row=3, column=col, value=lbl)
        c.font = f_entete; c.fill = fill_vert; c.alignment = center; c.border = bord
    for j in range(1, nb_jours + 1):
        d = date(annee, mois, j)
        c = ws.cell(row=3, column=2 + j, value=j)
        c.font = f_entete; c.fill = fill_vert; c.alignment = center; c.border = bord
        # Dimanche ou férié → en-tête rouge (repère visuel)
        if type_jour_auto(d) in ("DIMANCHE", "FERIE"):
            c.fill = PatternFill("solid", fgColor="B91C1C")
    ws.row_dimensions[3].height = 18

    # Lignes journaliers
    for r, j in enumerate(journaliers, start=4):
        cid = ws.cell(row=r, column=1, value=j.id)
        cid.font = f_normal; cid.alignment = center; cid.border = bord
        cnom = ws.cell(row=r, column=2, value=j.nom_complet)
        cnom.font = f_normal; cnom.border = bord
        cnom.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for jour in range(1, nb_jours + 1):
            d = date(annee, mois, jour)
            c = ws.cell(row=r, column=2 + jour)
            c.font = f_normal; c.alignment = center; c.border = bord
            if type_jour_auto(d) in ("DIMANCHE", "FERIE"):
                c.fill = fill_we  # grisé : repère (mais saisissable, ex. dimanche travaillé)

    # Largeurs + gel des volets
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 24
    for j in range(1, nb_jours + 1):
        ws.column_dimensions[get_column_letter(2 + j)].width = 4.5
    ws.freeze_panes = "C4"

    # ── Sortie fichier ─────────────────────────────────────────────────
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"pointage_{annee}-{mois:02d}"
    if site_obj:
        nom_propre = "".join(c if c.isalnum() else "_" for c in site_obj.nom)
        fname += "_" + nom_propre
    return send_file(
        buf, as_attachment=True, download_name=f"{fname}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

@bp.route("/journaliers/pointage/importer", methods=["POST"])
@login_required
def journaliers_pointage_importer():
    """Étape 1/2 — APERÇU. Lit la grille mensuelle téléversée, applique la règle
    journalier (≤8h normal / surplus en sup / dimanche-férié → tout en sup),
    et affiche ce qui SERA créé/modifié. N'écrit RIEN en base. Les données
    validées repartent dans un champ caché vers /confirmer.
    """
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))

    import openpyxl, json
    import calendar as _cal

    fichier = request.files.get("fichier")
    if not fichier or not fichier.filename.endswith((".xlsx", ".xls")):
        flash("❌ Fichier invalide. Utilisez le modèle Excel fourni (.xlsx).", "error")
        return redirect(url_for("tenant.pointage"))

    # Garde anti-DoS : 5 Mo max (même règle que l'import salariés).
    fichier.seek(0, 2); _taille = fichier.tell(); fichier.seek(0)
    if _taille > 5_000_000:
        flash("❌ Fichier trop volumineux (max 5 Mo).", "error")
        return redirect(url_for("tenant.pointage"))

    try:
        wb = openpyxl.load_workbook(fichier, data_only=True)
        ws = wb.active
    except Exception as e:
        flash(f"❌ Erreur lecture fichier : {e}", "error")
        return redirect(url_for("tenant.pointage"))

    # ── Retrouver l'en-tête : la ligne qui contient "ID" en col.1 ──────
    header_row = None
    for r in range(1, 8):
        v = str(ws.cell(r, 1).value or "").upper().strip()
        if v == "ID":
            header_row = r
            break
    if not header_row:
        flash("❌ En-tête introuvable. Utilisez le modèle téléchargé (colonne ID).", "error")
        return redirect(url_for("tenant.pointage"))

    # ── Déduire l'année/mois à partir des colonnes de jours ────────────
    # Les colonnes 3..N portent les numéros de jour. On récupère aussi le
    # mois/année depuis le TITRE (ligne 1) si présent, sinon mois courant.
    titre = str(ws.cell(1, 1).value or "")
    annee = mois = None
    import re
    MOIS_MAP = {"JANVIER":1,"FÉVRIER":2,"FEVRIER":2,"MARS":3,"AVRIL":4,"MAI":5,
                "JUIN":6,"JUILLET":7,"AOÛT":8,"AOUT":8,"SEPTEMBRE":9,
                "OCTOBRE":10,"NOVEMBRE":11,"DÉCEMBRE":12,"DECEMBRE":12}
    for nom_m, num in MOIS_MAP.items():
        if nom_m in titre.upper():
            mois = num
            break
    m_an = re.search(r"(20\d{2})", titre)
    if m_an:
        annee = int(m_an.group(1))
    if not (annee and mois):
        now = datetime.now(); annee, mois = now.year, now.month
    nb_jours_mois = _cal.monthrange(annee, mois)[1]

    # ── Mapper les colonnes-jours : {col: numéro_de_jour} ──────────────
    col_jour = {}
    for c in range(3, ws.max_column + 1):
        val = ws.cell(header_row, c).value
        try:
            j = int(val)
            if 1 <= j <= nb_jours_mois:
                col_jour[c] = j
        except (TypeError, ValueError):
            continue
    if not col_jour:
        flash("❌ Aucune colonne de jour détectée. Utilisez le modèle téléchargé.", "error")
        return redirect(url_for("tenant.pointage"))

    # ── IDs de journaliers valides pour ce tenant (sécurité) ───────────
    ids_valides = {j.id: j for j in Journalier.query.filter_by(
        tenant_id=t.id).all()}

    # ── Parcours des lignes → construction de l'aperçu ─────────────────
    a_appliquer = []   # [{jid, nom, date, heures, hn, hs, present, type_jour, action}]
    lignes_ignorees = 0
    erreurs = []

    for r in range(header_row + 1, ws.max_row + 1):
        id_val = ws.cell(r, 1).value
        if id_val in (None, ""):
            continue
        try:
            jid = int(id_val)
        except (TypeError, ValueError):
            erreurs.append(f"Ligne {r} : ID « {id_val} » invalide.")
            continue
        j_obj = ids_valides.get(jid)
        if not j_obj:
            erreurs.append(f"Ligne {r} : journalier ID {jid} inconnu (ignoré).")
            continue

        for c, jour in col_jour.items():
            case = ws.cell(r, c).value
            # Case VIDE → on ignore ce jour (ne touche à rien).
            if case in (None, ""):
                continue
            try:
                heures = float(str(case).replace(",", "."))
            except (TypeError, ValueError):
                erreurs.append(f"Ligne {r}, jour {jour} : valeur « {case} » non numérique.")
                continue
            if heures < 0:
                erreurs.append(f"Ligne {r}, jour {jour} : heures négatives.")
                continue

            d = date(annee, mois, jour)
            tj = type_jour_auto(d)  # NORMAL | DIMANCHE | FERIE

            # ── Règle journalier ──────────────────────────────────────
            if heures == 0:
                present = False; hn = 0.0; hs = 0.0
            else:
                present = True
                if tj in ("DIMANCHE", "FERIE"):
                    hn = 0.0; hs = round(heures, 2)         # tout majoré
                elif heures > 8:
                    hn = 8.0; hs = round(heures - 8, 2)     # surplus en sup
                else:
                    hn = round(heures, 2); hs = 0.0

            # Existe déjà ? (pour l'étiquette créé / mis à jour)
            existe = Pointage.query.filter_by(
                tenant_id=t.id, date_pointage=d, journalier_id=jid).first() is not None

            a_appliquer.append({
                "jid": jid, "nom": j_obj.nom_complet,
                "date": d.strftime("%Y-%m-%d"), "jour": jour,
                "heures": heures, "hn": hn, "hs": hs,
                "present": present, "type_jour": tj,
                "action": "maj" if existe else "creer",
            })

    if not a_appliquer:
        flash("Aucune donnée exploitable dans le fichier (toutes les cases sont vides ?).", "warning")
        return redirect(url_for("tenant.pointage"))

    nb_creer = sum(1 for x in a_appliquer if x["action"] == "creer")
    nb_maj   = sum(1 for x in a_appliquer if x["action"] == "maj")

    return render_template("tenant/pointage_import_apercu.html",
        tenant=t, lignes=a_appliquer, nb_creer=nb_creer, nb_maj=nb_maj,
        nb_total=len(a_appliquer), erreurs=erreurs,
        annee=annee, mois=mois,
        payload=json.dumps(a_appliquer))

@bp.route("/journaliers/pointage/importer/confirmer", methods=["POST"])
@login_required
def journaliers_pointage_importer_confirmer():
    """Étape 2/2 — ENREGISTREMENT. Reçoit le JSON validé (champ caché de
    l'aperçu) et applique update-or-create sur chaque Pointage.
    """
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))

    import json
    try:
        lignes = json.loads(request.form.get("payload", "[]"))
    except (ValueError, TypeError):
        flash("❌ Données d'import illisibles. Recommencez le téléversement.", "error")
        return redirect(url_for("tenant.pointage"))

    if not lignes:
        flash("Aucune donnée à enregistrer.", "warning")
        return redirect(url_for("tenant.pointage"))

    # IDs valides (re-vérif sécurité : on ne fait jamais confiance au client).
    ids_valides = {j.id for j in Journalier.query.filter_by(tenant_id=t.id).all()}

    nb_creer = nb_maj = 0
    for x in lignes:
        try:
            jid = int(x["jid"])
            d   = _parse_date(x["date"])
            hn  = float(x["hn"]); hs = float(x["hs"])
            present = bool(x["present"]); tj = str(x["type_jour"])
        except (KeyError, ValueError, TypeError):
            continue
        if jid not in ids_valides or not d:
            continue

        pt = Pointage.query.filter_by(
            tenant_id=t.id, date_pointage=d, journalier_id=jid).first()
        if pt:
            nb_maj += 1
        else:
            pt = Pointage(tenant_id=t.id, date_pointage=d, journalier_id=jid)
            db.session.add(pt)
            nb_creer += 1

        pt.present         = present
        pt.absent          = not present
        pt.heures_normales = hn
        pt.heures_sup      = hs
        pt.type_jour       = tj

    db.session.commit()
    log_action("IMPORT", "pointage", None,
               f"Import pointages journaliers : {nb_creer} créé(s), {nb_maj} mis à jour")
    flash(f"✅ Import terminé : {nb_creer} pointage(s) créé(s), {nb_maj} mis à jour.", "success")
    return redirect(url_for("tenant.pointage"))

@bp.route("/journaliers/paie/imprimer-sites")
@login_required
def journaliers_paie_imprimer_sites():
    """Impression de la paie journalier GROUPÉE PAR SITE (une page par site).

    Pensée pour distribuer aux chefs de chantier : chaque site a sa propre page
    (saut de page), avec le détail des journaliers, une colonne signature et un
    sous-total. Une page récapitulative tous sites termine le document.
    Filtre par période obligatoire côté formulaire : date_debut / date_fin.
    """
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    statut_f   = request.args.get("statut", "")
    date_debut = _parse_date(request.args.get("date_debut", ""))
    date_fin   = _parse_date(request.args.get("date_fin", ""))

    # Feuilles de la période (et statut éventuel). On retient toute feuille qui
    # CHEVAUCHE la période choisie (plus tolérant qu'une inclusion stricte).
    q = FeuillePaieJournalier.query.filter_by(tenant_id=t.id)
    if statut_f:
        q = q.filter_by(statut=statut_f)
    if date_debut:
        q = q.filter(FeuillePaieJournalier.date_fin >= date_debut)
    if date_fin:
        q = q.filter(FeuillePaieJournalier.date_debut <= date_fin)
    feuilles = q.options(joinedload(FeuillePaieJournalier.journalier)).order_by(
        FeuillePaieJournalier.date_fin.desc()).all()
    # ── Sélection optionnelle des journaliers à imprimer ──────────────
    # Le front peut transmettre ?ids=3,7,12 pour ne sortir qu'une partie des
    # journaliers. Absent ou vide → comportement historique : tout le monde.
    ids_param = request.args.get("ids", "").strip()
    if ids_param:
        try:
            ids_voulus = {int(x) for x in ids_param.split(",") if x.strip().isdigit()}
        except ValueError:
            ids_voulus = set()
        if ids_voulus:
            feuilles = [f for f in feuilles if f.journalier_id in ids_voulus]

    # Site de chaque journalier : on prend TOUTES les affectations, en préférant
    # l'affectation active, puis la plus récente. Ainsi un journalier dont
    # l'affectation n'est pas marquée "active" n'est pas perdu, et seuls les
    # journaliers réellement sans aucune affectation tombent dans "Sans site".
    site_par_journalier = {}
    affs = (AffectationSite.query
            .filter_by(tenant_id=t.id)
            .filter(AffectationSite.journalier_id.isnot(None))
            .order_by(AffectationSite.actif.desc(),
                      AffectationSite.date_debut.desc())
            .all())
    for a in affs:
        if a.journalier_id not in site_par_journalier and a.site is not None:
            site_par_journalier[a.journalier_id] = a.site

    sites_list = Site.query.filter_by(tenant_id=t.id, actif=True).order_by(Site.nom).all()

    # ── Sélection optionnelle des SITES à imprimer ────────────────────
    # Le front peut transmettre ?sites=1,4,7 pour ne sortir que certains sites.
    # La valeur 0 (ou "sans") inclut les journaliers sans aucune affectation.
    # Absent ou vide → comportement historique : tous les sites.
    sites_param = request.args.get("sites", "").strip()
    sites_voulus = None
    if sites_param:
        sites_voulus = set()
        for x in sites_param.split(","):
            x = x.strip()
            if x.isdigit():
                sites_voulus.add(int(x))
            elif x.lower() == "sans":
                sites_voulus.add(0)
        if sites_voulus:
            feuilles = [f for f in feuilles
                        if (site_par_journalier.get(f.journalier_id).id
                            if site_par_journalier.get(f.journalier_id) else 0) in sites_voulus]

    # Déduction des avances (par journalier) sur les feuilles affichées
    imput_av = _imputer_avances_journalier(
        feuilles, _avances_par_journalier(t.id, {f.journalier_id for f in feuilles}))

    # Regroupement : { site_id : {"site": Site|None, "feuilles": [...], "total": x} }
    groupes = {}
    SANS_SITE = 0
    for f in feuilles:
        s = site_par_journalier.get(f.journalier_id)
        key = s.id if s else SANS_SITE
        if key not in groupes:
            groupes[key] = {"site": s, "feuilles": [], "total": 0.0,
                            "total_brut": 0.0, "total_avance": 0.0,
                            "recap_mode": {"ESPECES": {"total": 0.0, "nb": 0},
                                           "VIREMENT": {"total": 0.0, "nb": 0}}}
        groupes[key]["feuilles"].append(f)
        net = imput_av.get(f.id, {}).get("net", float(f.montant_a_payer or 0))
        groupes[key]["total"]       += net
        groupes[key]["total_brut"]  += float(f.montant_a_payer or 0)
        groupes[key]["total_avance"]+= imput_av.get(f.id, {}).get("avance", 0.0)
        mode = (f.mode_paiement if f.statut == "PAYÉ" and f.mode_paiement
                else (f.journalier.mode_paiement if f.journalier else "ESPECES")) or "ESPECES"
        if mode not in ("ESPECES", "VIREMENT"):
            mode = "ESPECES"
        groupes[key]["recap_mode"][mode]["total"] += net
        groupes[key]["recap_mode"][mode]["nb"]    += 1

    # Ordonner : sites par nom (selon sites_list), puis "Sans site" à la fin
    groupes_ordonnes = []
    for s in sites_list:
        if s.id in groupes:
            groupes_ordonnes.append(groupes[s.id])
    if SANS_SITE in groupes:
        groupes_ordonnes.append(groupes[SANS_SITE])

    total_general = sum(g["total"] for g in groupes_ordonnes)
    nb_total = sum(len(g["feuilles"]) for g in groupes_ordonnes)
    # Récap global par mode (somme des récaps de chaque site)
    recap_global = {"ESPECES": {"total": 0.0, "nb": 0}, "VIREMENT": {"total": 0.0, "nb": 0}}
    for g in groupes_ordonnes:
        for m in ("ESPECES", "VIREMENT"):
            recap_global[m]["total"] += g["recap_mode"][m]["total"]
            recap_global[m]["nb"]    += g["recap_mode"][m]["nb"]
    return render_template("tenant/journaliers_paie_sites_print.html",
        tenant=t, groupes=groupes_ordonnes, total_general=total_general,
        recap_global=recap_global,
        nb_total=nb_total, statut=statut_f, date_debut=date_debut, date_fin=date_fin,
        imput_av=imput_av, now=datetime.now())


@bp.route("/journaliers/paie/imprimer")
@login_required
def journaliers_paie_imprimer():
    """Page imprimable des feuilles de paie journalier (avec colonne signature).

    Filtres optionnels : site_id, statut, date_debut, date_fin.
    """
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    site_id    = request.args.get("site_id", type=int)
    statut_f   = request.args.get("statut", "")
    date_debut = _parse_date(request.args.get("date_debut", ""))
    date_fin   = _parse_date(request.args.get("date_fin", ""))
    site       = Site.query.filter_by(id=site_id, tenant_id=t.id).first() if site_id else None

    q = FeuillePaieJournalier.query.filter_by(tenant_id=t.id)
    if site_id:
        ids_j = [a.journalier_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_id).filter(
            AffectationSite.journalier_id.isnot(None)).all()]
        q = q.filter(FeuillePaieJournalier.journalier_id.in_(ids_j))
    if statut_f:
        q = q.filter_by(statut=statut_f)
    if date_debut:
        q = q.filter(FeuillePaieJournalier.date_debut >= date_debut)
    if date_fin:
        q = q.filter(FeuillePaieJournalier.date_fin <= date_fin)
    feuilles = q.options(joinedload(FeuillePaieJournalier.journalier)).order_by(
        FeuillePaieJournalier.date_fin.desc()).all()
    imput_av = _imputer_avances_journalier(
        feuilles, _avances_par_journalier(t.id, {f.journalier_id for f in feuilles}))
    total = sum(imput_av.get(f.id, {}).get("net", float(f.montant_a_payer or 0)) for f in feuilles)
    # ── Sous-totaux par mode de paiement (net réellement décaissé) ──────
    recap_mode = {"ESPECES": {"total": 0.0, "nb": 0}, "VIREMENT": {"total": 0.0, "nb": 0}}
    for f in feuilles:
        mode = (f.mode_paiement if f.statut == "PAYÉ" and f.mode_paiement
                else (f.journalier.mode_paiement if f.journalier else "ESPECES")) or "ESPECES"
        if mode not in recap_mode:
            mode = "ESPECES"
        net = imput_av.get(f.id, {}).get("net", float(f.montant_a_payer or 0))
        recap_mode[mode]["total"] += net
        recap_mode[mode]["nb"]    += 1
    return render_template("tenant/journaliers_paie_print.html",
        tenant=t, feuilles=feuilles, site=site, statut=statut_f,
        date_debut=date_debut, date_fin=date_fin, total=total,
        recap_mode=recap_mode,
        imput_av=imput_av, now=datetime.now())


@bp.route("/journaliers/paie/export")
@login_required
def journaliers_paie_export():
    """Export Excel des feuilles de paie journalier — filtré par site et/ou période."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
    import io, calendar

    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))

    # ── Paramètres de filtre ──────────────────────────────────────────────────
    site_id    = request.args.get("site_id",    type=int)
    statut_f   = request.args.get("statut",     "")
    date_debut = _parse_date(request.args.get("date_debut", ""))
    date_fin   = _parse_date(request.args.get("date_fin",   ""))
    site       = Site.query.filter_by(id=site_id, tenant_id=t.id).first() if site_id else None

    # ── Requête ───────────────────────────────────────────────────────────────
    q = FeuillePaieJournalier.query.filter_by(tenant_id=t.id)
    if site_id:
        ids_j = [a.journalier_id for a in AffectationSite.query.filter_by(
            tenant_id=t.id, site_id=site_id
        ).filter(AffectationSite.journalier_id.isnot(None)).all()]
        q = q.filter(FeuillePaieJournalier.journalier_id.in_(ids_j))
    if date_debut:
        q = q.filter(FeuillePaieJournalier.date_debut >= date_debut)
    if date_fin:
        q = q.filter(FeuillePaieJournalier.date_fin   <= date_fin)
    if statut_f:
        q = q.filter_by(statut=statut_f)
    feuilles = q.order_by(
        FeuillePaieJournalier.date_fin.desc(),
        FeuillePaieJournalier.journalier_id
    ).all()

    # ── Affectation site de chaque journalier ─────────────────────────────────
    aff_map = {}
    for a in AffectationSite.query.filter_by(tenant_id=t.id).filter(
            AffectationSite.journalier_id.isnot(None)).all():
        if a.journalier_id not in aff_map:
            aff_map[a.journalier_id] = a.site.nom if a.site else "—"

    # ── Styles communs ────────────────────────────────────────────────────────
    HDR_FONT   = Font(bold=True, color="FFFFFF", size=9)
    HDR_FILL   = PatternFill("solid", fgColor="1a2332")
    HDR_ALIGN  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    BODY_FONT  = Font(size=9)
    EVEN_FILL  = PatternFill("solid", fgColor="F7F8FA")
    TOTAL_FONT = Font(bold=True, size=10, color="FFFFFF")
    TOTAL_FILL = PatternFill("solid", fgColor="1a2332")
    MONEY_FMT  = '#,##0'
    thin       = Side(style="thin", color="D1D5DB")
    BORDER     = Border(left=thin, right=thin, top=thin, bottom=thin)
    CENTER     = Alignment(horizontal="center")
    RIGHT      = Alignment(horizontal="right")

    wb = Workbook()

    # ══════════════════════════════════════════════════════════════════════════
    # ONGLET 1 — Détail complet
    # ══════════════════════════════════════════════════════════════════════════
    ws = wb.active
    ws.title = "Détail"
    ws.freeze_panes = "A4"

    # Titre
    titre = (f"PAIE JOURNALIERS — {t.denomination}"
             + (f" — {site.nom}" if site else "")
             + (f" — {date_debut.strftime('%d/%m/%Y')} au {date_fin.strftime('%d/%m/%Y')}" if date_debut and date_fin else "")
             + f" — Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}")
    ws.merge_cells("A1:K1")
    ws["A1"] = titre
    ws["A1"].font = Font(bold=True, size=12, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1a2332")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 22

    ws.append([])  # ligne vide

    # En-têtes
    hdrs = ["Journalier","Profession","Site","Période du","au",
            "Nb jours","Total heures","Taux/h (FCFA)","Montant brut (FCFA)","Statut","Date paiement"]
    ws.append(hdrs)
    for c_idx, h in enumerate(hdrs, 1):
        cell = ws.cell(row=3, column=c_idx, value=h)
        cell.font  = HDR_FONT
        cell.fill  = HDR_FILL
        cell.alignment = HDR_ALIGN
        cell.border = BORDER
    ws.row_dimensions[3].height = 18

    # Données
    total_montant  = 0
    total_jours    = 0
    total_heures   = 0
    for row_idx, f in enumerate(feuilles, 4):
        site_nom  = aff_map.get(f.journalier_id, "—")
        montant   = float(f.montant_brut  or 0)
        heures    = float(f.total_heures  or 0)
        jours     = int(f.nb_jours or 0)
        total_montant += montant
        total_jours   += jours
        total_heures  += heures
        row_data = [
            csv_safe(f.journalier.nom_complet),
            csv_safe(f.journalier.profession or "—"),
            csv_safe(site_nom),
            f.date_debut.strftime("%d/%m/%Y") if f.date_debut else "",
            f.date_fin.strftime("%d/%m/%Y")   if f.date_fin   else "",
            jours,
            round(heures, 2),
            float(f.taux_horaire or 0),
            montant,
            csv_safe(f.statut),
            f.date_paiement.strftime("%d/%m/%Y") if f.date_paiement else "",
        ]
        ws.append(row_data)
        is_even = (row_idx % 2 == 0)
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=c_idx)
            cell.font   = BODY_FONT
            cell.border = BORDER
            if is_even: cell.fill = EVEN_FILL
            # Formats numériques
            if c_idx in (6, 7):   cell.alignment = CENTER
            if c_idx in (8, 9):
                cell.number_format = MONEY_FMT
                cell.alignment     = RIGHT
            # Statut coloré
            if c_idx == 10:
                cell.alignment = CENTER
                if val == "PAYÉ":
                    cell.font = Font(bold=True, color="065F46", size=9)
                else:
                    cell.font = Font(bold=True, color="92400E", size=9)

    # Ligne totaux
    ws.append([])
    tr = ws.max_row + 1
    totals = ["", "", "", "", "TOTAL", total_jours, round(total_heures,2),
              "", total_montant, "", ""]
    ws.append(totals)
    for c_idx, val in enumerate(totals, 1):
        cell = ws.cell(row=tr, column=c_idx)
        cell.font   = TOTAL_FONT
        cell.fill   = TOTAL_FILL
        cell.border = BORDER
        if c_idx in (8, 9):
            cell.number_format = MONEY_FMT
            cell.alignment     = RIGHT
        if c_idx == 5:
            cell.alignment = RIGHT

    # Largeurs colonnes
    col_widths = [28, 18, 20, 13, 13, 10, 13, 16, 20, 14, 15]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ══════════════════════════════════════════════════════════════════════════
    # ONGLET 2 — Récap par site
    # ══════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Récap par site")
    ws2.freeze_panes = "A3"

    ws2.merge_cells("A1:F1")
    ws2["A1"] = f"RÉCAPITULATIF PAR SITE — {t.denomination}"
    ws2["A1"].font = Font(bold=True, size=12, color="FFFFFF")
    ws2["A1"].fill = PatternFill("solid", fgColor="374151")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 20

    hdrs2 = ["Site","Nb journaliers","Nb jours total","Total heures","Montant brut (FCFA)","Statut"]
    ws2.append(hdrs2)
    for c_idx, h in enumerate(hdrs2, 1):
        cell = ws2.cell(row=2, column=c_idx, value=h)
        cell.font  = HDR_FONT
        cell.fill  = PatternFill("solid", fgColor="374151")
        cell.alignment = HDR_ALIGN
        cell.border = BORDER

    # Grouper par site
    from collections import defaultdict
    by_site = defaultdict(lambda: {"journaliers": set(), "jours": 0, "heures": 0.0,
                                    "montant": 0.0, "nb_payes": 0, "nb_total": 0})
    for f in feuilles:
        s_nom = aff_map.get(f.journalier_id, "Sans site")
        by_site[s_nom]["journaliers"].add(f.journalier_id)
        by_site[s_nom]["jours"]    += int(f.nb_jours or 0)
        by_site[s_nom]["heures"]   += float(f.total_heures or 0)
        by_site[s_nom]["montant"]  += float(f.montant_brut or 0)
        by_site[s_nom]["nb_total"] += 1
        if f.statut == "PAYÉ": by_site[s_nom]["nb_payes"] += 1

    grand_total = 0
    for row_idx, (s_nom, data) in enumerate(sorted(by_site.items()), 3):
        pct_paye = int(data["nb_payes"] / data["nb_total"] * 100) if data["nb_total"] else 0
        statut_txt = f"{data['nb_payes']}/{data['nb_total']} payé(s) ({pct_paye}%)"
        row_data = [
            csv_safe(s_nom),
            len(data["journaliers"]),
            data["jours"],
            round(data["heures"], 2),
            round(data["montant"], 2),
            statut_txt,
        ]
        ws2.append(row_data)
        grand_total += data["montant"]
        is_even = (row_idx % 2 == 0)
        for c_idx, val in enumerate(row_data, 1):
            cell = ws2.cell(row=row_idx, column=c_idx)
            cell.font   = BODY_FONT
            cell.border = BORDER
            if is_even: cell.fill = EVEN_FILL
            if c_idx == 5:
                cell.number_format = MONEY_FMT
                cell.alignment     = RIGHT
            if c_idx in (2, 3, 4):
                cell.alignment = CENTER

    # Total récap
    ws2.append([])
    tr2 = ws2.max_row + 1
    ws2.append(["TOTAL GÉNÉRAL", "", "", "", grand_total, ""])
    for c_idx in range(1, 7):
        cell = ws2.cell(row=tr2, column=c_idx)
        cell.font   = TOTAL_FONT
        cell.fill   = PatternFill("solid", fgColor="374151")
        cell.border = BORDER
        if c_idx == 5:
            cell.number_format = MONEY_FMT
            cell.alignment     = RIGHT

    for i, w in enumerate([28, 16, 14, 14, 22, 22], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # ══════════════════════════════════════════════════════════════════════════
    # ONGLET 3 — Récap par journalier
    # ══════════════════════════════════════════════════════════════════════════
    ws3 = wb.create_sheet("Récap par journalier")
    ws3.freeze_panes = "A3"

    ws3.merge_cells("A1:G1")
    ws3["A1"] = f"RÉCAPITULATIF PAR JOURNALIER — {t.denomination}"
    ws3["A1"].font = Font(bold=True, size=12, color="FFFFFF")
    ws3["A1"].fill = PatternFill("solid", fgColor="065F46")
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 20

    hdrs3 = ["Journalier","Profession","Site","Nb périodes","Nb jours","Total heures","Total perçu (FCFA)"]
    ws3.append(hdrs3)
    for c_idx, h in enumerate(hdrs3, 1):
        cell = ws3.cell(row=2, column=c_idx, value=h)
        cell.font  = HDR_FONT
        cell.fill  = PatternFill("solid", fgColor="065F46")
        cell.alignment = HDR_ALIGN
        cell.border = BORDER

    by_jour = defaultdict(lambda: {"nom":"","profession":"","site":"","periodes":0,"jours":0,"heures":0.0,"montant":0.0})
    for f in feuilles:
        jid = f.journalier_id
        by_jour[jid]["nom"]       = f.journalier.nom_complet
        by_jour[jid]["profession"]= f.journalier.profession or "—"
        by_jour[jid]["site"]      = aff_map.get(jid, "—")
        by_jour[jid]["periodes"]  += 1
        by_jour[jid]["jours"]     += int(f.nb_jours or 0)
        by_jour[jid]["heures"]    += float(f.total_heures or 0)
        by_jour[jid]["montant"]   += float(f.montant_brut or 0)

    grand_total3 = 0
    for row_idx, (jid, d) in enumerate(sorted(by_jour.items(), key=lambda x: x[1]["nom"]), 3):
        row_data = [csv_safe(d["nom"]), csv_safe(d["profession"]), csv_safe(d["site"]),
                    d["periodes"], d["jours"], round(d["heures"],2), round(d["montant"],2)]
        ws3.append(row_data)
        grand_total3 += d["montant"]
        is_even = (row_idx % 2 == 0)
        for c_idx, val in enumerate(row_data, 1):
            cell = ws3.cell(row=row_idx, column=c_idx)
            cell.font   = BODY_FONT
            cell.border = BORDER
            if is_even: cell.fill = EVEN_FILL
            if c_idx == 7:
                cell.number_format = MONEY_FMT; cell.alignment = RIGHT
            if c_idx in (4, 5, 6): cell.alignment = CENTER

    ws3.append([])
    tr3 = ws3.max_row + 1
    ws3.append(["TOTAL GÉNÉRAL","","","","","", grand_total3])
    for c_idx in range(1, 8):
        cell = ws3.cell(row=tr3, column=c_idx)
        cell.font = TOTAL_FONT
        cell.fill = PatternFill("solid", fgColor="065F46")
        cell.border = BORDER
        if c_idx == 7:
            cell.number_format = MONEY_FMT; cell.alignment = RIGHT

    for i, w in enumerate([28, 18, 20, 12, 11, 13, 22], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    # ── Export ────────────────────────────────────────────────────────────────
    out = io.BytesIO()
    wb.save(out); out.seek(0)

    parts = ["Paie_Journaliers"]
    if site:       parts.append(site.nom.replace(" ", "_"))
    if date_debut: parts.append(date_debut.strftime("%Y%m%d"))
    if date_fin:   parts.append("au" + date_fin.strftime("%Y%m%d"))
    parts.append(datetime.now().strftime("%Y%m%d"))
    fname = "_".join(parts) + ".xlsx"

    return send_file(out,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True, download_name=fname)

@bp.route("/journaliers/paie/payer-selection", methods=["POST"])
@login_required
def journaliers_payer_selection():
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    ids = [int(i) for i in request.form.get("feuille_ids","").split(",") if i.strip().isdigit()]
    nb = 0
    for fid in ids:
        f = FeuillePaieJournalier.query.filter_by(id=fid, tenant_id=t.id, statut="EN_ATTENTE").first()
        if f: f.statut="PAYÉ"; f.date_paiement=datetime.now().date(); nb+=1
    db.session.commit()
    flash(f"{nb} journalier(s) payé(s).", "success")
    return redirect(url_for("tenant.journaliers_paie"))



@bp.route("/journaliers/<int:id>/pointages/imprimer")
@login_required
def journalier_pointages_imprimer(id):
    """Relevé mensuel imprimable des pointages d'un journalier (totaux d'heures)."""
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))
    j = Journalier.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    mois, annee = _resoudre_mois_annee()
    import calendar
    debut = date(annee, mois, 1)
    fin   = date(annee, mois, calendar.monthrange(annee, mois)[1])
    pts = (Pointage.query
           .filter_by(tenant_id=t.id, journalier_id=id)
           .filter(Pointage.date_pointage >= debut, Pointage.date_pointage <= fin)
           .options(joinedload(Pointage.site))
           .order_by(Pointage.date_pointage).all())
    # Les journaliers ne relèvent pas de la ventilation conventionnelle BTP :
    # on cumule directement les colonnes pointées.
    ctx = _pointages_mois_contexte(t, pts, convention=None)
    return render_template("tenant/pointages_print.html",
        tenant=t, now=datetime.now(),
        personne={"nom_complet": j.nom_complet,
                  "reference": (j.profession or "Journalier"),
                  "type": "Journalier"},
        mois=mois, annee=annee, mois_libelle=_MOIS_FR[mois], **ctx)




@bp.route("/journaliers/imprimer")
@login_required
def journaliers_imprimer():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    return render_template("tenant/journaliers_print.html", journaliers=Journalier.query.filter_by(tenant_id=t.id).order_by(Journalier.nom).all(), tenant=t, now=datetime.now())

# ── Export & API ──────────────────────────────────────────────────────────────


@bp.route("/pointage/recap-semaine")
@login_required
def pointage_recap_semaine():
    """Récapitulatif de présence hebdomadaire par site."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))

    # Semaine sélectionnée
    date_str = request.args.get("date", datetime.now().strftime("%Y-%m-%d"))
    try:    date_ref = datetime.strptime(date_str, "%Y-%m-%d").date()
    except: date_ref = datetime.now().date()

    lundi  = date_ref - timedelta(days=date_ref.weekday())
    jours  = [lundi + timedelta(days=i) for i in range(6)]  # lundi→samedi
    samedi = jours[-1]

    # Sites actifs
    sites_list = Site.query.filter_by(tenant_id=t.id, actif=True).order_by(Site.nom).all()

    # Tous les pointages de la semaine
    pts_semaine = Pointage.query.filter_by(tenant_id=t.id)        .filter(Pointage.date_pointage >= lundi,
                Pointage.date_pointage <= samedi).all()

    # Affectations actives (site → workers)
    aff_sal  = {}  # salarie_id  → site_id
    aff_jour = {}  # journalier_id → site_id
    for a in AffectationSite.query.filter_by(tenant_id=t.id, actif=True).all():
        if a.salarie_id:    aff_sal[a.salarie_id]    = a.site_id
        if a.journalier_id: aff_jour[a.journalier_id] = a.site_id

    # ── Construire les données par site et par jour ───────────────────────────
    # Structure : { site_id: { date: { presents, absents, heures, h_sup, non_pointes } } }
    JOURS_FR = ["Lundi","Mardi","Mercredi","Jeudi","Vendredi","Samedi"]

    # Nb de travailleurs affectés à chaque site
    effectif_site = {}
    for s in sites_list:
        nb_sal  = sum(1 for v in aff_sal.values()  if v == s.id)
        nb_jour = sum(1 for v in aff_jour.values() if v == s.id)
        effectif_site[s.id] = nb_sal + nb_jour

    recap = {}
    for s in sites_list:
        recap[s.id] = {
            "site": s,
            "effectif": effectif_site.get(s.id, 0),
            "jours": {},
            "totaux": {"presents": 0, "absents": 0, "heures": 0.0, "h_sup": 0.0},
        }
        for j in jours:
            recap[s.id]["jours"][str(j)] = {
                "date":        j,
                "jour_fr":     JOURS_FR[j.weekday()],
                "presents":    0,
                "absents":     0,
                "non_pointes": effectif_site.get(s.id, 0),
                "heures":      0.0,
                "h_sup":       0.0,
            }

    # Sac sans site
    recap["sans_site"] = {
        "site": None,
        "effectif": 0,
        "jours": {},
        "totaux": {"presents": 0, "absents": 0, "heures": 0.0, "h_sup": 0.0},
    }
    for j in jours:
        recap["sans_site"]["jours"][str(j)] = {
            "date": j, "jour_fr": JOURS_FR[j.weekday()],
            "presents": 0, "absents": 0, "non_pointes": 0,
            "heures": 0.0, "h_sup": 0.0,
        }

    # Remplir avec les pointages réels
    for p in pts_semaine:
        d = str(p.date_pointage)
        site_id = aff_sal.get(p.salarie_id) or aff_jour.get(p.journalier_id)
        key = site_id if site_id and site_id in recap else "sans_site"

        if d not in recap[key]["jours"]:
            continue

        cell = recap[key]["jours"][d]
        if p.present:
            cell["presents"]    += 1
            h_norm = float(p.heures_normales or 8)
            h_sup  = (float(p.heures_sup_10 or 0) + float(p.heures_sup_30 or 0) +
                      float(p.heures_sup_40 or 0) + float(p.heures_sup_70 or 0) +
                      float(p.heures_sup or 0))
            cell["heures"] += h_norm
            cell["h_sup"]  += h_sup
        else:
            cell["absents"] += 1
        # Recalcul non_pointés
        cell["non_pointes"] = max(0,
            recap[key]["effectif"] - cell["presents"] - cell["absents"])

    # Calculer les totaux semaine par site
    for key in recap:
        tot = recap[key]["totaux"]
        for d, cell in recap[key]["jours"].items():
            tot["presents"] += cell["presents"]
            tot["absents"]  += cell["absents"]
            tot["heures"]   += cell["heures"]
            tot["h_sup"]    += cell["h_sup"]

    # Totaux globaux tous sites
    totaux_globaux = {"presents": 0, "absents": 0, "heures": 0.0, "h_sup": 0.0}
    for key in recap:
        t2 = recap[key]["totaux"]
        totaux_globaux["presents"] += t2["presents"]
        totaux_globaux["absents"]  += t2["absents"]
        totaux_globaux["heures"]   += t2["heures"]
        totaux_globaux["h_sup"]    += t2["h_sup"]

    # Filtrer "sans_site" si vide
    if recap["sans_site"]["totaux"]["presents"] == 0 and recap["sans_site"]["totaux"]["absents"] == 0:
        del recap["sans_site"]

    return render_template("tenant/pointage_recap_semaine.html",
        tenant=t,
        jours=jours,
        jours_fr=JOURS_FR,
        recap=recap,
        sites_list=sites_list,
        totaux_globaux=totaux_globaux,
        lundi=lundi,
        samedi=samedi,
        date_ref=date_ref,
        semaine_prec=(lundi - timedelta(days=7)).strftime("%Y-%m-%d"),
        semaine_suiv=(lundi + timedelta(days=7)).strftime("%Y-%m-%d"),
        now=datetime.now())

