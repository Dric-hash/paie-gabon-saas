# -*- coding: utf-8 -*-
"""Paramètres de l'entreprise, utilisateurs, rubriques, grille de salaires,
export de données, abonnement — extrait de tenant.py (même blueprint « tenant »)."""
from datetime import datetime, date
from flask import (render_template, request, redirect, url_for, flash, session,
                   current_app, abort, Response)
from flask_login import login_required, current_user
from blueprints.tenant import bp, ROLES_TENANT_AUTORISES, _config_rubriques_dict
from core import tenant_required, get_tenant, can_edit, admin_only
from audit import log_action
from models import (db, Plan, Utilisateur, Salarie, Contrat, PeriodePaie, BulletinPaie,
                    Conge, CategorieEmploi, ComposantPaie, ConfigRubrique, RubriquePaie)


@bp.route("/parametres/api/regenerer-token", methods=["POST"])
@tenant_required
@admin_only
def regenerer_token_api():
    """Régénère le token API du tenant. Le token en clair n'est affiché qu'ici,
    une seule fois (il est stocké haché). Réservé aux administrateurs du tenant."""
    t = get_tenant()
    raw = t.generate_token()
    db.session.commit()
    log_action("REGENERATE", "token_api", t.id, "Régénération du token API")
    db.session.commit()
    flash(f"Nouveau token API : {raw} — copiez-le maintenant, il ne sera plus affiché.",
          "success")
    return redirect(url_for("api_v1.api_clients_list"))




@bp.route("/parametres/export-donnees")
@tenant_required
def parametres_export_donnees():
    """Exporte toutes les données de l'entreprise dans un ZIP (classeur Excel multi-onglets)."""
    import openpyxl, zipfile
    from flask import Response
    from openpyxl.styles import Font, PatternFill
    t = get_tenant()

    # (Onglet, Modèle, requête filtrée sur le tenant)
    exports = [
        ("Salariés",     Salarie,        Salarie.query.filter_by(tenant_id=t.id)),
        ("Contrats",     Contrat,        Contrat.query.filter_by(tenant_id=t.id)),
        ("Bulletins",    BulletinPaie,   BulletinPaie.query.filter_by(tenant_id=t.id)),
        ("Périodes",     PeriodePaie,    PeriodePaie.query.filter_by(tenant_id=t.id)),
        ("Congés",       Conge,          Conge.query.filter_by(tenant_id=t.id)),
        ("Catégories",   CategorieEmploi, CategorieEmploi.query.filter_by(tenant_id=t.id)),
        ("Composants",   ComposantPaie,  ComposantPaie.query.filter_by(tenant_id=t.id)),
    ]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    HF = PatternFill("solid", fgColor="0f3d36"); HN = Font(bold=True, color="FFFFFF")

    def _val(v):
        if v is None: return ""
        if isinstance(v, (datetime, date)): return v.strftime("%d/%m/%Y %H:%M") if isinstance(v, datetime) else v.strftime("%d/%m/%Y")
        if isinstance(v, bool): return "Oui" if v else "Non"
        return v

    for titre, modele, query in exports:
        cols = [c.name for c in modele.__table__.columns]
        ws = wb.create_sheet(titre[:31])
        for ci, cn in enumerate(cols, 1):
            cell = ws.cell(1, ci, cn); cell.fill = HF; cell.font = HN
        for ri, obj in enumerate(query.all(), 2):
            for ci, cn in enumerate(cols, 1):
                ws.cell(ri, ci, _val(getattr(obj, cn, None)))
        ws.freeze_panes = "A2"

    xlsx_buf = io.BytesIO(); wb.save(xlsx_buf); xlsx_buf.seek(0)

    # ZIP : le classeur + un LISEZ-MOI
    nom_ent = (t.denomination or "entreprise").replace(" ", "_")[:30]
    horodatage = datetime.now().strftime("%Y-%m-%d_%H%M")
    readme = (
        f"EXPORT DES DONNÉES — {t.denomination}\n"
        f"Généré le {datetime.now().strftime('%d/%m/%Y à %H:%M')}\n\n"
        f"Ce fichier contient une copie de vos données PaieGabon :\n"
        f"  • Salariés, Contrats, Bulletins, Périodes, Congés, Catégories, Composants.\n\n"
        f"Ouvrez le fichier .xlsx avec Excel, LibreOffice ou Google Sheets.\n"
        f"Conservez cette archive en lieu sûr (disque, cloud personnel).\n\n"
        f"— Ameriack I.T. Solutions / PaieGabon\n"
    )
    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"donnees_{nom_ent}_{horodatage}.xlsx", xlsx_buf.getvalue())
        zf.writestr("LISEZ-MOI.txt", readme)
    zip_buf.seek(0)

    log_action("EXPORT", "tenant", t.id, "Export des données de l'entreprise",
               user_id=current_user.id, tenant_id=t.id)
    nom_zip = f"export_paiegabon_{nom_ent}_{horodatage}.zip"
    return Response(zip_buf.getvalue(), mimetype="application/zip",
                    headers={"Content-Disposition": f'attachment; filename="{nom_zip}"'})


@bp.route("/parametres")
@tenant_required
def parametres():
    t=get_tenant()
    # Passer tous les plans actifs pour l'onglet abonnement
    plans_dispo = Plan.query.filter_by(actif=True).order_by(Plan.prix_mensuel.asc()).all()
    return render_template("tenant/parametres.html", tenant=t,
        config_rubriques=_config_rubriques_dict(t.id),
        rubriques=RubriquePaie.query.filter_by(actif=True).all(),
        categories=CategorieEmploi.query.filter_by(tenant_id=t.id).all(),
        users=Utilisateur.query.filter_by(tenant_id=t.id).all(),
        plans_dispo=plans_dispo)

@bp.route("/parametres/grille-salaires", methods=["GET", "POST"])
@tenant_required
def grille_salaires():
    """Édition/vérification de la grille de salaires conventionnelle.
    L'auto-remplissage du salaire de base (fiche salarié) n'utilise QUE la grille
    enregistrée ici — jamais la graine brute — pour garantir des montants validés."""
    import json as _json
    from calculs_paie import GRILLE_CATEGORIES_AERIEN, grille_salaire_aerien_seed
    t = get_tenant()
    if not current_user.can_edit:
        abort(403)

    if request.method == "POST":
        grille = {}
        for key, val in request.form.items():
            if not key.startswith("montant_") or not (val or "").strip():
                continue
            try:
                _, code, ech = key.split("_", 2)
                montant = float(val.replace(" ", "").replace("\u202f", "").replace(",", "."))
            except (ValueError, IndexError):
                continue
            if montant > 0:
                grille.setdefault(code, {})[ech] = round(montant, 2)
        t.grille_salaires = _json.dumps(grille, ensure_ascii=False)
        db.session.commit()
        log_action("UPDATE", "tenant", t.id, "Grille de salaires mise à jour")
        flash("Grille de salaires enregistrée. Elle est désormais utilisée pour "
              "pré-remplir le salaire de base des salariés.", "success")
        return redirect(url_for("tenant.grille_salaires"))

    # GET : grille sauvegardée, sinon graine aérienne (à vérifier) si convention AERIEN.
    grille, est_seed = {}, False
    if t.grille_salaires:
        try:
            grille = _json.loads(t.grille_salaires)
        except (ValueError, TypeError):
            grille = {}
    if not grille and (t.convention or "").upper() == "AERIEN":
        grille = grille_salaire_aerien_seed()
        est_seed = True
    if not grille and (t.convention or "").upper() == "HOTELLERIE":
        from convention_hotellerie import grille_salaire_hotellerie_seed
        grille = grille_salaire_hotellerie_seed()
        est_seed = True
    if not grille and (t.convention or "").upper() == "BOIS":
        from convention_bois import grille_salaire_bois_seed
        grille = grille_salaire_bois_seed()
        est_seed = True
    if not grille and (t.convention or "").upper() == "MINIER":
        from convention_minier import grille_salaire_minier_seed
        grille = grille_salaire_minier_seed()
        est_seed = True
    if not grille and (t.convention or "").upper() == "FORET":
        from convention_foret import grille_salaire_foret_seed
        grille = grille_salaire_foret_seed()
        est_seed = True
    return render_template("tenant/grille_salaires.html", tenant=t,
        categories=GRILLE_CATEGORIES_AERIEN, grille=grille,
        est_seed=est_seed, nb_echelons=10)




@bp.route("/parametres/logo", methods=["POST"])
@login_required
def parametres_logo():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    logo_file = request.files.get("logo")
    if logo_file and logo_file.filename:
        import base64
        file_data = logo_file.read()
        if len(file_data) > 1_000_000:
            flash("Fichier trop volumineux. Maximum 1 Mo.", "error")
            return redirect(url_for("tenant.parametres"))

        # ── Validation de l'extension ─────────────────────────────────────────
        ext = logo_file.filename.rsplit(".", 1)[-1].lower() if "." in logo_file.filename else ""
        EXTENSIONS_AUTORISEES = {"png", "jpg", "jpeg", "gif", "webp"}
        if ext not in EXTENSIONS_AUTORISEES:
            flash("Format non autorisé. Utilisez PNG, JPG, JPEG, GIF ou WEBP (pas SVG).", "error")
            return redirect(url_for("tenant.parametres"))

        # ── Validation du MIME réel (magic bytes) — pas seulement l'extension ─
        MAGIC = {
            b"\x89PNG":   "image/png",
            b"\xff\xd8\xff": "image/jpeg",
            b"GIF8":      "image/gif",
            b"RIFF":      None,  # WebP — vérification complémentaire ci-dessous
        }
        detected_mime = None
        for magic, mime in MAGIC.items():
            if file_data[:len(magic)] == magic:
                if magic == b"RIFF" and file_data[8:12] == b"WEBP":
                    detected_mime = "image/webp"
                else:
                    detected_mime = mime
                break
        if not detected_mime:
            flash("Le contenu du fichier ne correspond pas à une image valide.", "error")
            return redirect(url_for("tenant.parametres"))

        b64 = base64.b64encode(file_data).decode("utf-8")
        logo_data = f"data:{detected_mime};base64,{b64}"
        try:
            db.session.execute(db.text("UPDATE tenants SET logo_url = :logo WHERE id = :id"),{"logo": logo_data, "id": t.id})
            db.session.commit(); db.session.expire(t)
            log_action("UPDATE", "parametres", t.id, "Mise à jour du logo de la société")
            db.session.commit()
            flash("Logo mis à jour avec succès.", "success")
        except Exception as e:
            db.session.rollback(); flash(f"Erreur: {str(e)}", "error")
    else:
        flash("Aucun fichier sélectionné.", "error")
    return redirect(url_for("tenant.parametres"))

@bp.route("/parametres/logo/supprimer", methods=["POST"])
@login_required
def parametres_logo_supprimer():
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    t.logo_url = None; db.session.commit()
    log_action("UPDATE", "parametres", t.id, "Suppression du logo de la société")
    db.session.commit()
    flash("Logo supprime.", "success")
    return redirect(url_for("tenant.parametres"))

@bp.route("/parametres/modele-bulletin", methods=["POST"])
@login_required
def parametres_modele_bulletin():
    """Changer le modèle d'impression des bulletins."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    modele = request.form.get("modele_bulletin", "classique")
    if modele not in ("classique", "moderne", "minimaliste", "grandlivre", "sgtg"):
        modele = "classique"
    t.modele_bulletin = modele
    db.session.commit()
    log_action("UPDATE", "parametres", t.id, f"Modèle d'impression bulletins : {modele}")
    db.session.commit()
    flash(f"Modèle d'impression « {modele.capitalize()} » appliqué.", "success")
    return redirect(url_for("tenant.parametres"))

@bp.route("/parametres/rubriques", methods=["POST"])
@login_required
def parametres_rubriques():
    """Enregistre le paramétrage des rubriques fixes souples (panier, transport,
    représentation, salisure) : entre dans le brut, position, soumis CNSS/CNAMGS/IRPP."""
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.can_edit:
        flash("Accès refusé.", "error"); return redirect(url_for("tenant.parametres"))
    from models import ConfigRubrique
    for cle in ("panier", "transport", "representation", "salisure"):
        c = ConfigRubrique.query.filter_by(tenant_id=t.id, cle=cle).first()
        if not c:
            c = ConfigRubrique(tenant_id=t.id, cle=cle)
            db.session.add(c)
        c.entre_dans_brut = request.form.get(f"{cle}_brut") == "on"
        c.position        = "HAUT" if request.form.get(f"{cle}_position") == "HAUT" else "BAS"
        c.soumis_cnss     = request.form.get(f"{cle}_cnss") == "on"
        c.soumis_cnamgs   = request.form.get(f"{cle}_cnamgs") == "on"
        c.soumis_irpp     = request.form.get(f"{cle}_irpp") == "on"
    db.session.commit()
    flash("Paramétrage des rubriques enregistré.", "success")
    return redirect(url_for("tenant.parametres"))


@bp.route("/parametres/societe", methods=["POST"])
@tenant_required
@can_edit
def parametres_societe():
    if not current_user.can_manage_parametres:
        flash("Accès refusé. Seul l'administrateur peut modifier les paramètres.", "error")
        return redirect(url_for("tenant.parametres"))
    t=get_tenant()
    for f in ["denomination","sigle","activite","secteur","nif","numero_cnss","numero_cnamgs","adresse","boite_postale","telephone","ville","region","representant_nom","representant_fonction"]:
        if f not in request.form:   # champ absent de ce formulaire → ne pas écraser
            continue
        try: setattr(t,f,request.form.get(f,"").strip() or None)
        except: pass
    # Convention collective applicable — mise à jour uniquement si le champ est
    # présent dans le formulaire soumis (évite d'écraser la valeur depuis un
    # formulaire qui ne l'inclut pas, ex. la carte « Informations générales »).
    from calculs_paie import CONVENTIONS_DISPONIBLES
    conv_raw = request.form.get("convention")
    if conv_raw is not None:
        conv = (conv_raw or "AUCUNE").upper()
        if conv in CONVENTIONS_DISPONIBLES:
            t.convention = conv
    # Seuil hebdomadaire de déclenchement des heures supplémentaires (dérogation).
    # La loi gabonaise fixe le seuil légal à 40h/semaine ; une dérogation peut le
    # porter jusqu'à 48h. On borne strictement la valeur saisie dans cet intervalle.
    seuil_raw = request.form.get("seuil_heures_sup_hebdo")
    if seuil_raw is not None and str(seuil_raw).strip() != "":
        try:
            seuil = float(str(seuil_raw).replace(",", "."))
            t.seuil_heures_sup_hebdo = max(40.0, min(seuil, 48.0))
        except (ValueError, TypeError):
            flash("Seuil d'heures supplémentaires invalide — valeur inchangée.", "error")
    # Langue de l'interface — idem, seulement si le champ est soumis.
    langue = request.form.get("langue")
    if langue is not None and langue in SUPPORTED_LANGUAGES:
        t.langue = langue
        set_language(langue)
    db.session.commit()
    log_action("UPDATE", "parametres", t.id, "Modification des informations de la société")
    db.session.commit()
    flash("Informations mises à jour." if (t.langue or "fr") == "fr" else "Settings updated.", "success")
    return redirect(url_for("tenant.parametres"))


@bp.route("/parametres/importer-grille-commerce", methods=["POST"])
@tenant_required
@can_edit
def importer_grille_commerce():
    """Crée/complète les catégories d'emploi à partir de la grille conventionnelle COMMERCE."""
    if not current_user.can_manage_parametres:
        flash("Accès refusé. Seul l'administrateur peut modifier les paramètres.", "error")
        return redirect(url_for("tenant.parametres"))
    t = get_tenant()
    from calculs_paie import GRILLE_COMMERCE
    existantes = {c.code for c in CategorieEmploi.query.filter_by(tenant_id=t.id).all()}
    ajout = 0
    maj = 0
    for code, libelle, mensuel, _horaire in GRILLE_COMMERCE:
        if code in existantes:
            cat = CategorieEmploi.query.filter_by(tenant_id=t.id, code=code).first()
            if cat and (cat.salaire_minimum is None or float(cat.salaire_minimum or 0) == 0):
                cat.salaire_minimum = mensuel
                maj += 1
        else:
            db.session.add(CategorieEmploi(
                tenant_id=t.id, code=code, libelle=libelle, salaire_minimum=mensuel,
                description="Grille Convention Collective du Commerce"))
            ajout += 1
    # Bascule la convention du tenant sur COMMERCE si pas déjà fait
    if t.convention != "COMMERCE":
        t.convention = "COMMERCE"
    db.session.commit()
    flash(f"Grille Commerce importée : {ajout} catégorie(s) ajoutée(s), {maj} mise(s) à jour.", "success")
    return redirect(url_for("tenant.parametres"))


@bp.route("/parametres/importer-grille-hydrocarbures", methods=["POST"])
@tenant_required
@can_edit
def importer_grille_hydrocarbures():
    """Crée/complète les catégories A→M à partir de la grille Hydrocarbures."""
    if not current_user.can_manage_parametres:
        flash("Accès refusé. Seul l'administrateur peut modifier les paramètres.", "error")
        return redirect(url_for("tenant.parametres"))
    t = get_tenant()
    from convention_hydrocarbures import GRILLE_HYDROCARBURES
    existantes = {c.code for c in CategorieEmploi.query.filter_by(tenant_id=t.id).all()}
    ajout = 0
    maj = 0
    for code, libelle, mensuel, _horaire in GRILLE_HYDROCARBURES:
        if code in existantes:
            cat = CategorieEmploi.query.filter_by(tenant_id=t.id, code=code).first()
            if cat and (cat.salaire_minimum is None or float(cat.salaire_minimum or 0) == 0):
                cat.salaire_minimum = mensuel
                maj += 1
        else:
            db.session.add(CategorieEmploi(
                tenant_id=t.id, code=code, libelle=libelle, salaire_minimum=mensuel,
                description="Grille Convention Hydrocarbures (Recherche & Exploitation)"))
            ajout += 1
    # Bascule la convention du tenant sur HYDROCARBURES
    if t.convention != "HYDROCARBURES":
        t.convention = "HYDROCARBURES"
    # Congés de base : 2,5 jours ouvrables / mois (Art. 42)
    if hasattr(t, "jours_conge_par_mois"):
        t.jours_conge_par_mois = 2.5
    db.session.commit()
    flash(f"Grille Hydrocarbures importée : {ajout} catégorie(s) ajoutée(s), {maj} mise(s) à jour.", "success")
    return redirect(url_for("tenant.parametres"))


@bp.route("/parametres/importer-grille-petrole", methods=["POST"])
@tenant_required
@can_edit
def importer_grille_petrole():
    """Crée/complète les catégories d'emploi à partir de la grille conventionnelle PÉTROLE.

    ⚠️ Les montants de l'Annexe n°2 (1983) sont obsolètes : les premières
    catégories sont sous le SMIG actuel. On importe donc la STRUCTURE des
    catégories en appliquant un plancher au SMIG légal, et on n'écrase jamais une
    catégorie existante. Les montants doivent être actualisés par l'entreprise.
    """
    if not current_user.can_manage_parametres:
        flash("Accès refusé. Seul l'administrateur peut modifier les paramètres.", "error")
        return redirect(url_for("tenant.parametres"))
    t = get_tenant()
    from calculs_paie import GRILLE_PETROLE, SMIG_GABON
    existantes = {c.code for c in CategorieEmploi.query.filter_by(tenant_id=t.id).all()}
    ajout = 0
    maj = 0
    plancher_applique = False
    for code, libelle, mensuel_1983 in GRILLE_PETROLE:
        # Plancher SMIG : aucune catégorie ne peut être créée sous le minimum légal.
        mensuel = max(int(mensuel_1983), int(SMIG_GABON))
        if mensuel != int(mensuel_1983):
            plancher_applique = True
        if code in existantes:
            cat = CategorieEmploi.query.filter_by(tenant_id=t.id, code=code).first()
            if cat and (cat.salaire_minimum is None or float(cat.salaire_minimum or 0) == 0):
                cat.salaire_minimum = mensuel
                maj += 1
        else:
            db.session.add(CategorieEmploi(
                tenant_id=t.id, code=code, libelle=libelle, salaire_minimum=mensuel,
                description="Grille Convention Pétrole (montants à actualiser)"))
            ajout += 1
    if t.convention != "PETROLE":
        t.convention = "PETROLE"
    db.session.commit()
    msg = f"Grille Pétrole importée : {ajout} catégorie(s) ajoutée(s), {maj} mise(s) à jour."
    if plancher_applique:
        msg += (f" ⚠️ Certains montants de 1983 étaient sous le SMIG "
                f"({int(SMIG_GABON):,} FCFA) et ont été relevés au plancher légal — "
                f"actualisez-les selon votre grille interne.").replace(",", " ")
    flash(msg, "success")
    return redirect(url_for("tenant.parametres"))




@bp.route("/parametres/demande-changement-plan", methods=["POST"])
@login_required
def demande_changement_plan():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.is_tenant_admin: abort(403)
    plan_souhaite_id = request.form.get("plan_id", type=int)
    motif = request.form.get("motif", "").strip()
    plan_souhaite = db.session.get(Plan, plan_souhaite_id) if plan_souhaite_id else None
    if not plan_souhaite:
        flash("Plan invalide.", "error")
        return redirect(url_for("tenant.parametres"))
    # Enregistrer la demande dans les notes + changer statut
    note_demande = (
        f"[DEMANDE CHANGEMENT PLAN — {datetime.now().strftime('%d/%m/%Y %H:%M')}] "
        f"Plan souhaité : {plan_souhaite.nom} ({int(plan_souhaite.prix_mensuel):,} FCFA/mois). "
        f"Motif : {motif or 'Non précisé'}. "
        f"Demandé par : {current_user.nom_complet} ({current_user.email})."
    )
    # Ajouter à la suite des notes existantes
    t.notes = (t.notes or "") + ("\n" if t.notes else "") + note_demande
    t.statut = "PAIEMENT_EN_ATTENTE"
    db.session.commit()
    flash(f"Demande de passage au plan « {plan_souhaite.nom} » enregistrée. L'équipe PaieGabon vous contactera sous 24h pour finaliser.", "success")
    return redirect(url_for("tenant.parametres"))

@bp.route("/parametres/annuler-abonnement", methods=["POST"])
@login_required
def annuler_abonnement():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.is_tenant_admin: abort(403)
    motif = request.form.get("motif", "").strip()
    t.statut = "ANNULATION_DEMANDEE"
    t.notes = f"Annulation demandée le {datetime.now().strftime('%d/%m/%Y')}. Motif: {motif}"
    db.session.commit()
    flash("Demande d annulation enregistrée. L equipe PaieGabon vous contactera sous 48h.", "success")
    return redirect(url_for("tenant.parametres"))

# ── Utilisateurs ──────────────────────────────────────────────────────────────
@bp.route("/utilisateurs")
@login_required
def utilisateurs():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    liste = Utilisateur.query.filter_by(tenant_id=t.id).order_by(Utilisateur.nom).all()
    return render_template("tenant/utilisateurs.html", tenant=t, utilisateurs=liste, users=liste)

@bp.route("/utilisateurs/nouveau", methods=["GET","POST"])
@login_required
def utilisateur_nouveau():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.is_tenant_admin:
        flash("Réservé à l administrateur.", "error")
        return redirect(url_for("tenant.utilisateurs"))
    # ── Vérifier la limite dès le GET (bloquer l'accès au formulaire) ───────
    if t.plan and t.plan.max_utilisateurs:
        nb_actuel = Utilisateur.query.filter_by(tenant_id=t.id, actif=True).count()
        if nb_actuel >= t.plan.max_utilisateurs:
            flash(
                f"Limite atteinte — Plan « {t.plan.nom} » : "
                f"{t.plan.max_utilisateurs} utilisateur(s) maximum "
                f"(vous en avez {nb_actuel}). "
                f"Passez au plan supérieur pour en ajouter d'autres.",
                "error"
            )
            return redirect(url_for("tenant.utilisateurs"))

    if request.method == "GET":
        nb_utilisateurs = Utilisateur.query.filter_by(tenant_id=t.id, actif=True).count()
        return render_template("tenant/utilisateur_form.html", tenant=t,
            nb_utilisateurs=nb_utilisateurs)
    email = request.form.get("email", "").strip().lower()
    nom = request.form.get("nom", "").strip().upper()
    prenom = request.form.get("prenom", "").strip()
    password = request.form.get("password", "")
    role = request.form.get("role", "GESTIONNAIRE").strip().upper()
    # Liste blanche : empêche l'attribution de SUPER_ADMIN ou d'un rôle inconnu.
    if role not in ROLES_TENANT_AUTORISES:
        flash("Rôle invalide.", "error")
        return redirect(url_for("tenant.utilisateurs"))
    if not email or not nom or not password:
        flash("Veuillez remplir tous les champs.", "error")
        return render_template("tenant/utilisateur_form.html", tenant=t)
    if Utilisateur.query.filter_by(email=email).first():
        flash("Email déjà utilisé.", "error")
        return render_template("tenant/utilisateur_form.html", tenant=t)
    u = Utilisateur(nom=nom, prenom=prenom, email=email, role=role, tenant_id=t.id, actif=True)
    u.set_password(password)
    db.session.add(u); db.session.commit()
    log_action("CREATE", "utilisateur", u.id,
               f"Création utilisateur {u.nom_complet} ({u.role})")
    db.session.commit()
    flash(f"Utilisateur {u.nom_complet} créé.", "success")
    return redirect(url_for("tenant.utilisateurs"))

@bp.route("/utilisateurs/<int:id>/toggle", methods=["POST"])
@login_required
def utilisateur_toggle(id):
    """Activer / désactiver un utilisateur."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.is_tenant_admin:
        flash("Réservé à l'administrateur.", "error")
        return redirect(url_for("tenant.utilisateurs"))
    u = Utilisateur.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if u.id == current_user.id:
        flash("Vous ne pouvez pas vous désactiver vous-même.", "error")
        return redirect(url_for("tenant.utilisateurs"))
    u.actif = not u.actif
    db.session.commit()
    etat = "activé" if u.actif else "désactivé"
    log_action("UPDATE", "utilisateur", u.id,
               f"Utilisateur {u.nom_complet} {etat} par {current_user.nom_complet}")
    db.session.commit()
    flash(f"Utilisateur {u.nom_complet} {etat}.", "success")
    return redirect(url_for("tenant.utilisateurs"))


@bp.route("/utilisateurs/<int:id>/modifier", methods=["GET","POST"])
@login_required
def utilisateur_modifier(id):
    """Modifier le rôle et les infos d'un utilisateur."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.is_tenant_admin:
        flash("Réservé à l'administrateur.", "error")
        return redirect(url_for("tenant.utilisateurs"))

    u = Utilisateur.query.filter_by(id=id, tenant_id=t.id).first_or_404()

    if request.method == "POST":
        ancien_role = u.role
        nouveau_role = request.form.get("role", u.role).strip().upper()
        # Liste blanche : empêche l'escalade vers SUPER_ADMIN ou un rôle inconnu.
        if nouveau_role not in ROLES_TENANT_AUTORISES:
            flash("Rôle invalide.", "error")
            return redirect(url_for("tenant.utilisateurs"))

        # Empêcher de retirer son propre rôle admin
        if u.id == current_user.id and nouveau_role != "TENANT_ADMIN":
            flash("Vous ne pouvez pas changer votre propre rôle.", "error")
            return redirect(url_for("tenant.utilisateurs"))

        u.nom    = request.form.get("nom", u.nom).strip().upper()
        u.prenom = request.form.get("prenom", u.prenom).strip()
        u.role   = nouveau_role

        # Changer le mot de passe si fourni
        nouveau_mdp = request.form.get("nouveau_mdp", "").strip()
        if nouveau_mdp:
            if len(nouveau_mdp) < 8:
                flash("Le mot de passe doit faire au moins 8 caractères.", "error")
                return render_template("tenant/utilisateur_form.html",
                                       utilisateur=u, tenant=t, mode="modifier")
            u.set_password(nouveau_mdp)

        db.session.commit()
        log_action("UPDATE", "utilisateur", u.id,
                   f"Modification {u.nom_complet} — rôle : {ancien_role} → {nouveau_role}")
        db.session.commit()
        flash(f"Utilisateur {u.nom_complet} mis à jour.", "success")
        return redirect(url_for("tenant.utilisateurs"))

    return render_template("tenant/utilisateur_form.html",
                           utilisateur=u, tenant=t, mode="modifier")


@bp.route("/utilisateurs/<int:id>/supprimer", methods=["POST"])
@login_required
def utilisateur_supprimer(id):
    """
    Supprime définitivement un utilisateur.
    Règles :
      - Seul l'admin du tenant peut supprimer
      - On ne peut pas se supprimer soi-même
      - On ne peut pas supprimer le dernier admin
    """
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.is_tenant_admin:
        flash("Réservé à l'administrateur.", "error")
        return redirect(url_for("tenant.utilisateurs"))

    u = Utilisateur.query.filter_by(id=id, tenant_id=t.id).first_or_404()

    # Règle 1 : pas de suicide
    if u.id == current_user.id:
        flash("Vous ne pouvez pas supprimer votre propre compte.", "error")
        return redirect(url_for("tenant.utilisateurs"))

    # Règle 2 : conserver au moins un admin actif
    if u.role == "TENANT_ADMIN":
        nb_admins = Utilisateur.query.filter_by(
            tenant_id=t.id, role="TENANT_ADMIN", actif=True
        ).filter(Utilisateur.id != u.id).count()
        if nb_admins == 0:
            flash("Impossible de supprimer le seul administrateur actif du compte. "
                  "Activez ou désignez d'abord un autre administrateur.", "error")
            return redirect(url_for("tenant.utilisateurs"))

    nom_sauvegarde = u.nom_complet
    log_action("DELETE", "utilisateur", u.id,
               f"Suppression définitive de {nom_sauvegarde} ({u.role_label})")
    db.session.delete(u)
    db.session.commit()
    flash(f"Utilisateur {nom_sauvegarde} supprimé définitivement.", "success")
    return redirect(url_for("tenant.utilisateurs"))

# ── Journaliers ───────────────────────────────────────────────────────────────
# ── Acomptes ──────────────────────────────────────────────────────────────────
