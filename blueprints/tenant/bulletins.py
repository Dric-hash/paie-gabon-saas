# -*- coding: utf-8 -*-
"""Bulletins de paie, périodes, composants, et API de calcul/simulation —
extrait de tenant.py (même blueprint « tenant »)."""
from datetime import datetime, date, timedelta
from flask import (render_template, request, redirect, url_for, flash, session,
                   current_app, abort, send_file, jsonify, Response)
from flask_login import login_required, current_user
from flask_mail import Message
from sqlalchemy.orm import joinedload
from sqlalchemy import desc, func
from blueprints.tenant import bp, logger, _config_rubriques_dict, _pdf_bulletin_bytes, _bulletin_imprimer_impl
from core import (tenant_required, get_tenant, can_edit, _parse_date, parse_date,
                  _cache_delete, _pd, send_email_async, calculer_parts_irpp)
from audit import log_action
from calculs_paie import (calculer_bulletin, calculer_masse_salariale,
                          calculer_taux_horaire, calculer_heures_sup_btp)
from models import (db, BulletinPaie, BulletinComposant, PeriodePaie, ComposantPaie,
                    Salarie, Contrat, Acompte, Pointage, Site, AffectationSite)


@bp.route("/api/simuler-paie/scenarios", methods=["POST"])
@login_required
def api_simuler_scenarios():
    """Compare jusqu'à 3 scénarios de paie côte à côte."""
    if current_user.is_super_admin: return jsonify({"error": "forbidden"}), 403
    t = get_tenant()
    if not t: return jsonify({"error": "non authentifié"}), 401

    data = request.get_json(force=True) or {}
    scenarios = data.get("scenarios", [])
    nb_parts  = float(data.get("nb_parts", 1.0))

    if not scenarios or len(scenarios) < 2:
        return jsonify({"error": "Au moins 2 scénarios requis"}), 400

    from simulation_paie import comparer_scenarios
    result = comparer_scenarios(scenarios, nb_parts=nb_parts)
    return jsonify(result)


@bp.route("/api/simuler-paie/net-vers-brut", methods=["POST"])
@login_required
def api_simuler_net_vers_brut():
    """Calcule le brut nécessaire pour atteindre un net cible."""
    if current_user.is_super_admin: return jsonify({"error": "forbidden"}), 403
    t = get_tenant()
    if not t: return jsonify({"error": "non authentifié"}), 401

    data = request.get_json(force=True) or {}
    net_cible = float(data.get("net_cible", 0))
    nb_parts  = float(data.get("nb_parts",  1.0))
    extras    = data.get("extras", {})

    if not net_cible or net_cible <= 0:
        return jsonify({"error": "Net cible invalide"}), 400

    from simulation_paie import simuler_depuis_net
    result = simuler_depuis_net(net_cible, nb_parts=nb_parts, donnees_extra=extras)
    return jsonify(result)


@bp.route("/api/simuler-paie/augmentation", methods=["POST"])
@login_required
def api_simuler_augmentation():
    """Simule l'impact d'une augmentation de salaire."""
    if current_user.is_super_admin: return jsonify({"error": "forbidden"}), 403
    t = get_tenant()
    if not t: return jsonify({"error": "non authentifié"}), 401

    data = request.get_json(force=True) or {}
    salaire_actuel      = float(data.get("salaire_actuel", 0))
    augmentation_pct    = data.get("augmentation_pct")
    augmentation_montant= data.get("augmentation_montant")
    nb_parts            = float(data.get("nb_parts", 1.0))
    extras              = data.get("extras", {})

    if not salaire_actuel or salaire_actuel <= 0:
        return jsonify({"error": "Salaire actuel invalide"}), 400

    from simulation_paie import simuler_augmentation
    result = simuler_augmentation(
        salaire_actuel       = salaire_actuel,
        augmentation_pct     = float(augmentation_pct) if augmentation_pct else None,
        augmentation_montant = float(augmentation_montant) if augmentation_montant else None,
        nb_parts             = nb_parts,
        donnees_extra        = extras,
    )
    return jsonify(result)


@bp.route("/api/simuler-paie", methods=["POST"])
@login_required
def api_simuler_paie():
    """API JSON : simuler un bulletin de paie complet sans le sauvegarder."""
    t = get_tenant()
    if not t: return jsonify({"erreur": "non connecté"})
    from calculs_paie import calculer_bulletin, calculer_heures_sup_btp
    try:
        d = request.get_json(force=True) or {}
        # Accepter aussi le form-data
        if not d:
            d = {k: request.form.get(k) for k in request.form}

        def flt(key, default=0):
            try: return float(str(d.get(key) or default).replace(",",".") or default)
            except: return float(default)

        # Conversion des heures supplémentaires (saisies en HEURES) → montants.
        # Effectuée quelle que soit la convention : les coefficients +10/+30/+40/+70
        # sont communs au BTP et au Commerce (la sélection de convention côté UI
        # ne change que les libellés et le préremplissage structurel BTP).
        sal_base = flt("salaire_base")
        if sal_base > 0:
            from calculs_paie import calculer_heures_sup_btp
            hs = calculer_heures_sup_btp(sal_base,
                h10=flt("h10"), h30=flt("h30"),
                h40=flt("h40"), h70=flt("h70"),
                h30b=flt("h30b"), convention=t.convention)
            d["heures_sup_10"] = hs["montant_10"]
            d["heures_sup_30"] = hs["montant_30"]
            d["heures_sup_30b"] = hs["montant_30b"]
            d["heures_sup_40"] = hs["montant_40"]
            d["heures_sup_70"] = hs["montant_70"]
        d["convention"] = t.convention

        # Nombre de parts IRPP
        sal_id = d.get("salarie_id")
        nb_parts = 1.0
        if sal_id:
            s = Salarie.query.filter_by(id=int(sal_id), tenant_id=t.id).first()
            if s: nb_parts = float(s.nombre_parts or 1)
        nb_parts = flt("nb_parts") or nb_parts

        # Composants personnalisés (aperçu temps réel) : lus depuis composant_<id>
        comps_live = []
        for comp in ComposantPaie.query.filter_by(tenant_id=t.id, actif=True).all():
            montant = flt(f"composant_{comp.id}")
            if montant:
                comps_live.append({
                    "libelle": comp.libelle, "sens": comp.sens, "montant": montant,
                    "soumis_cnss": comp.soumis_cnss, "soumis_cnamgs": comp.soumis_cnamgs,
                    "soumis_irpp": comp.soumis_irpp,
                    "entre_dans_brut": comp.entre_dans_brut, "position": comp.position})
        d["composants"] = comps_live
        d["config_rubriques"] = _config_rubriques_dict(t.id)

        result = calculer_bulletin(d, nb_parts=nb_parts)

        # Enrichir avec détails BTP
        from calculs_paie import calculer_taux_horaire, H_NORMALES_MENSUEL
        if sal_base > 0:
            th = calculer_taux_horaire(sal_base)
            result["taux_horaire"]     = round(th, 2)
            result["h_normales_mensuel"] = H_NORMALES_MENSUEL

        # Ajouter libellés pour l'affichage
        result["gains_detail"] = [
            {"label": "Salaire de base",          "montant": result["salaire_base"]},
            {"label": "H.sup +10%",               "montant": result["heures_sup_10"]},
            {"label": "H.sup +30%",               "montant": result["heures_sup_30"]},
            {"label": "H.sup +30% (repos/férié)", "montant": result.get("heures_sup_30b", 0)},
            {"label": "H.sup +40% (nuit/dim.)",   "montant": result["heures_sup_40"]},
            {"label": "H.sup +70% (fériés)",      "montant": result["heures_sup_70"]},
            {"label": "Sursalaire",               "montant": result["sursalaire"]},
            {"label": "Prime de transport",       "montant": result["prime_transport"]},
            {"label": "Prime de responsabilité",  "montant": result["prime_responsabilite"]},
            {"label": "Indemnité logement",       "montant": result["indem_logement"]},
            {"label": "Prime d'ancienneté",      "montant": result["prime_anciennete"]},
            {"label": "Autres primes",            "montant": sum([
                result["prime_caisse"], result["carburant"],
                result["prime_rendement"], result["prime_qualite"],
                result["prime_performance"], result["prime_assiduité"],
            ])},
        ]
        result["gains_detail"] = [g for g in result["gains_detail"] if g["montant"] > 0]

        result["retenues_detail"] = [
            {"label": f"CNSS salarié (5% / base {int(result['base_cnss']):,} FCFA)", "montant": result["cnss_salarie"]},
            {"label": f"CNAMGS salarié (2% / base {int(result['base_cnamgs']):,} FCFA)", "montant": result["cnamgs_salarie"]},
            {"label": f"TCS (5% / base {int(result['base_tcs']):,} FCFA)",             "montant": result["tcs"]},
            {"label": f"IRPP ({nb_parts} part(s))",                                    "montant": result["irpp"]},
            {"label": "Acompte",                                                        "montant": result["acompte"]},
            {"label": "Retenue absences",                                               "montant": result["absences"]},
        ]
        result["retenues_detail"] = [r for r in result["retenues_detail"] if r["montant"] > 0]

        result["charges_pat_detail"] = [
            {"label": "CNSS patronal (18%)",     "montant": result["cnss_patronale"]},
            {"label": "CNAMGS patronal (4.1%)",  "montant": result["cnamgs_patronale"]},
            {"label": "FNH (3%)",                "montant": result["fnh"]},
            {"label": "CFP (0.5%)",              "montant": result["cfp"]},
        ]

        return jsonify(result)
    except Exception as e:
        import traceback
        return jsonify({"erreur": str(e), "trace": traceback.format_exc()[-300:]})



@bp.route("/bulletins")
@login_required
def bulletins():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("auth.login"))
    pid          = request.args.get("periode_id", type=int)
    sf           = request.args.get("statut", "")
    site_filtre_id = request.args.get("site_id", type=int)
    periodes     = PeriodePaie.query.filter_by(tenant_id=t.id)                    .order_by(PeriodePaie.annee.desc(), PeriodePaie.mois.desc()).all()
    sites_list   = Site.query.filter_by(tenant_id=t.id, actif=True).order_by(Site.nom).all()
    site_filtre  = Site.query.filter_by(id=site_filtre_id, tenant_id=t.id).first() if site_filtre_id else None
    ps = None; buls = []; masse = {}; pagination = None

    if pid:
        ps = PeriodePaie.query.filter_by(id=pid, tenant_id=t.id).first_or_404()
        q = BulletinPaie.query.options(
            joinedload(BulletinPaie.salarie),
            joinedload(BulletinPaie.periode),
        ).filter_by(tenant_id=t.id, periode_id=pid)
        if sf:
            q = q.filter_by(statut=sf)

        # ── Filtre par site ──────────────────────────────────────────────────
        if site_filtre_id:
            # Récupérer les IDs des salariés affectés à ce site
            ids_sal = [a.salarie_id for a in AffectationSite.query.filter_by(
                tenant_id=t.id, site_id=site_filtre_id, actif=True
            ).filter(AffectationSite.salarie_id.isnot(None)).all()]
            q = q.filter(BulletinPaie.salarie_id.in_(ids_sal))

        page_bul   = request.args.get("page", 1, type=int)
        # Une seule query base avec le join, puis on pagine
        q_joined   = q.join(Salarie).order_by(Salarie.nom)
        buls_tous  = q_joined.all()
        masse      = calculer_masse_salariale(buls_tous)
        pagination = q_joined.paginate(page=page_bul, per_page=25, error_out=False)
        buls       = pagination.items

    # Affectation site de chaque salarié pour affichage dans le tableau
    aff_sal = {a.salarie_id: a.site for a in AffectationSite.query.filter_by(
        tenant_id=t.id, actif=True
    ).filter(AffectationSite.salarie_id.isnot(None)).all()}

    _args = {k: v for k, v in request.args.items() if k != 'page'}
    _base = request.path + '?' + '&'.join(f'{k}={v}' for k, v in _args.items())
    _sep  = '&' if _args else '?'
    return render_template("tenant/bulletins.html",
        periodes=periodes, periode_sel=ps,
        bulletins=buls, masse=masse, statut_filtre=sf,
        sites=sites_list, site_filtre=site_filtre, aff_sal=aff_sal,
        pagination=pagination if pid else None,
        pagination_base=_base + _sep,
        tenant=t)

@bp.route("/bulletins/bordereau")
@login_required
def bulletins_bordereau():
    """Bordereau de paie imprimable d'une période : liste des salariés, net à
    payer, mode de paiement, avec sous-totaux Espèces / Virement.
    """
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))

    pid = request.args.get("periode_id", type=int)
    if not pid:
        flash("Choisissez une période pour générer le bordereau.", "error")
        return redirect(url_for("tenant.bulletins"))
    ps = PeriodePaie.query.filter_by(id=pid, tenant_id=t.id).first_or_404()

    statut_f = request.args.get("statut", "")
    q = BulletinPaie.query.options(joinedload(BulletinPaie.salarie)) \
        .filter_by(tenant_id=t.id, periode_id=pid)
    if statut_f:
        q = q.filter_by(statut=statut_f)
    bulletins = q.join(Salarie).order_by(Salarie.nom).all()

    lignes = []
    recap_mode = {"ESPECES": {"total": 0.0, "nb": 0}, "VIREMENT": {"total": 0.0, "nb": 0}}
    total_general = 0.0
    for b in bulletins:
        net = float(b.net_a_payer or 0)
        mode = (b.mode_paiement if b.statut == "PAYÉ" and b.mode_paiement
                else (b.salarie.mode_paiement if b.salarie else "ESPECES")) or "ESPECES"
        if mode not in ("ESPECES", "VIREMENT"):
            mode = "ESPECES"
        recap_mode[mode]["total"] += net
        recap_mode[mode]["nb"]    += 1
        total_general += net
        lignes.append({
            "matricule": b.salarie.matricule if b.salarie else "—",
            "nom": b.salarie.nom_complet if b.salarie else "—",
            "emploi": (b.salarie.emploi if b.salarie else "") or "—",
            "net": net, "mode": mode, "statut": b.statut,
        })

    return render_template("tenant/bulletins_bordereau_print.html",
        tenant=t, periode=ps, lignes=lignes, recap_mode=recap_mode,
        total_general=total_general, nb_total=len(lignes),
        statut=statut_f, now=datetime.now())

@bp.route("/bulletins/generer-lot", methods=["POST"])
@login_required
def bulletins_generer_lot():
    """Génère un bulletin BROUILLON pour chaque salarié actif d'une période.

    Principes de sûreté :
      - ne crée QUE des brouillons (jamais de bulletin validé) ;
      - ne touche JAMAIS un bulletin existant (salarié ignoré, pas écrasé) ;
      - refuse une période clôturée ;
      - ignore les salariés sans contrat actif ou embauchés après la période.
    Le salaire de base vient du contrat actif ; la prime d'ancienneté et les
    acomptes en attente sont appliqués automatiquement. L'utilisateur ajuste
    ensuite les cas particuliers avant de valider.
    """
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))
    if not current_user.can_edit:
        abort(403)

    pid = request.form.get("periode_id", type=int)
    if not pid:
        flash("Choisissez une période avant de générer les bulletins.", "error")
        return redirect(url_for("tenant.bulletins"))
    periode = PeriodePaie.query.filter_by(id=pid, tenant_id=t.id).first_or_404()
    retour = f"/bulletins?periode_id={pid}"

    if periode.statut not in ("OUVERT", "OUVERTE"):
        flash(f"La période {periode.libelle_mois} {periode.annee} est "
              f"{periode.statut.lower()} : aucun bulletin ne peut y être généré.", "error")
        return redirect(retour)

    import calendar as _cal
    fin_periode = date(periode.annee, periode.mois,
                       _cal.monthrange(periode.annee, periode.mois)[1])

    # Salariés déjà dotés d'un bulletin sur cette période : intouchables.
    deja = {b.salarie_id for b in
            BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=pid).all()}

    salaries = (Salarie.query.filter_by(tenant_id=t.id, statut="ACTIF")
                .order_by(Salarie.nom, Salarie.prenom).all())

    crees = 0
    ignores_existants = ignores_sans_contrat = ignores_non_embauches = 0
    noms_sans_contrat = []

    for s in salaries:
        if s.id in deja:
            ignores_existants += 1
            continue
        if s.date_embauche and s.date_embauche > fin_periode:
            ignores_non_embauches += 1
            continue
        contrat = (Contrat.query
                   .filter_by(salarie_id=s.id, tenant_id=t.id, actif=True)
                   .order_by(Contrat.date_debut.desc()).first())
        if not contrat or not contrat.salaire_base:
            ignores_sans_contrat += 1
            if len(noms_sans_contrat) < 8:
                noms_sans_contrat.append(s.nom_complet)
            continue

        donnees = {"salaire_base": float(contrat.salaire_base)}
        if getattr(s, "travail_de_nuit", False):
            donnees["travail_de_nuit"] = True
        if s.date_embauche:
            donnees["anciennete_annees"] = max(
                0, (fin_periode - s.date_embauche).days // 365)

        # Acomptes en attente du mois : déduits comme dans la saisie unitaire.
        total_ac = float(db.session.query(db.func.sum(Acompte.montant))
                         .filter_by(tenant_id=t.id, salarie_id=s.id,
                                    mois=periode.mois, annee=periode.annee,
                                    statut="EN_ATTENTE").scalar() or 0)
        if total_ac > 0:
            donnees["acompte"] = total_ac

        res = calculer_bulletin(dict(donnees, convention=t.convention),
                                nb_parts=float(s.nombre_parts or 1))
        b = BulletinPaie(tenant_id=t.id, salarie_id=s.id, periode_id=pid)
        for k, v in res.items():
            if not k.startswith("_") and hasattr(b, k):
                setattr(b, k, v)
        b.statut = "BROUILLON"
        b.mode_paiement = s.mode_paiement or "ESPECES"
        db.session.add(b)
        crees += 1

    db.session.commit()
    log_action("GENERATE_BATCH", "bulletin", pid,
               f"Génération en lot {periode.libelle_mois} {periode.annee} : "
               f"{crees} brouillon(s) créé(s)")

    if crees:
        flash(f"{crees} bulletin(s) brouillon créé(s) pour "
              f"{periode.libelle_mois} {periode.annee}. Vérifiez-les avant validation.",
              "success")
    else:
        flash("Aucun bulletin créé : tous les salariés actifs en ont déjà un "
              "sur cette période, ou aucun n'a de contrat actif.", "warning")

    details = []
    if ignores_existants:
        details.append(f"{ignores_existants} salarié(s) avaient déjà un bulletin (inchangés)")
    if ignores_non_embauches:
        details.append(f"{ignores_non_embauches} embauché(s) après la période")
    if ignores_sans_contrat:
        noms = ", ".join(noms_sans_contrat)
        suite = "…" if ignores_sans_contrat > len(noms_sans_contrat) else ""
        details.append(f"{ignores_sans_contrat} sans contrat actif ({noms}{suite})")
    if details:
        flash(" · ".join(details), "info")

    return redirect(retour)


@bp.route("/bulletins/valider-lot", methods=["POST"])
@login_required
def bulletins_valider_lot():
    """Valider une sélection de bulletins ou tous les brouillons d'une période."""
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.can_edit: abort(403)

    pid      = request.form.get("periode_id", type=int)
    site_id  = request.form.get("site_id",    type=int)
    action   = request.form.get("action_lot", "valider")
    ids_str  = request.form.get("bulletin_ids", "")
    ids_sel  = [int(x) for x in ids_str.split(",") if x.strip().isdigit()]

    if not pid:
        flash("Période manquante.", "error")
        return redirect(url_for("tenant.bulletins"))

    # ── CORRECTION BUG : si aucun ID sélectionné + action valider → refuser ──
    # Avant : ids_sel vide → validait TOUS les bulletins de la période
    # Maintenant : ids_sel vide → uniquement pour les actions "tout valider" explicites
    if not ids_sel and action == "valider":
        # Vérifier que c'est bien une demande "tout valider" (bouton dédié)
        tout_valider = request.form.get("tout_valider", "0")
        if tout_valider != "1":
            flash("Aucun bulletin sélectionné. Cochez des bulletins ou utilisez 'Tout valider'.", "warning")
            return redirect(f"/bulletins?periode_id={pid}" + (f"&site_id={site_id}" if site_id else ""))

    # Charger les bulletins ciblés
    if ids_sel:
        buls = BulletinPaie.query.filter(
            BulletinPaie.id.in_(ids_sel),
            BulletinPaie.tenant_id == t.id
        ).all()
    else:
        q = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=pid, statut="BROUILLON")
        if site_id:
            ids_sal = [a.salarie_id for a in AffectationSite.query.filter_by(
                tenant_id=t.id, site_id=site_id, actif=True
            ).filter(AffectationSite.salarie_id.isnot(None)).all()]
            q = q.filter(BulletinPaie.salarie_id.in_(ids_sal))
        buls = q.all()

    nb = 0
    if action == "valider":
        for b in buls:
            if b.statut != "BROUILLON": continue
            b.statut          = "VALIDÉ"
            b.date_validation = utcnow()
            attribuer_numero_bulletin(b)
            for a in Acompte.query.filter_by(
                tenant_id=t.id, salarie_id=b.salarie_id,
                mois=b.periode.mois, annee=b.periode.annee, statut="EN_ATTENTE").all():
                a.statut = "DEDUIT"
            nb += 1
        msg = f"✅ {nb} bulletin(s) validé(s)."
        log_action("VALIDATE", "bulletin", pid,
                   f"Validation de {nb} bulletin(s) — période {pid}")

    elif action == "annuler_validation":
        # ── NOUVEAU : annuler la validation → repasser en BROUILLON ──────────
        for b in buls:
            if b.statut not in ("VALIDÉ", "VALIDE"): continue
            b.statut          = "BROUILLON"
            b.date_validation = None
            # Remettre les acomptes déduits en attente
            for a in Acompte.query.filter_by(
                tenant_id=t.id, salarie_id=b.salarie_id,
                mois=b.periode.mois, annee=b.periode.annee, statut="DEDUIT").all():
                a.statut = "EN_ATTENTE"
            nb += 1
        msg = f"↩️ {nb} bulletin(s) remis en brouillon."
        log_action("CANCEL", "bulletin", pid,
                   f"Annulation validation de {nb} bulletin(s) — période {pid}")

    elif action == "payer":
        for b in buls:
            if b.statut not in ("VALIDÉ", "VALIDE", "BROUILLON"): continue
            b.statut = "PAYÉ"
            nb += 1
        msg = f"💰 {nb} bulletin(s) marqué(s) comme payé(s)."
        log_action("PAY", "bulletin", pid,
                   f"{nb} bulletin(s) marqué(s) payé(s) — période {pid}")

    elif action == "supprimer_brouillons":
        for b in buls:
            if b.statut != "BROUILLON": continue
            db.session.delete(b); nb += 1
        msg = f"🗑️ {nb} brouillon(s) supprimé(s)."
        log_action("DELETE", "bulletin", pid,
                   f"Suppression de {nb} brouillon(s) — période {pid}")

    else:
        flash("Action inconnue.", "error")
        return redirect(f"/bulletins?periode_id={pid}")

    db.session.commit()
    _cache_delete(f"{t.id}:")
    flash(msg, "success")
    redir = f"/bulletins?periode_id={pid}"
    if site_id: redir += f"&site_id={site_id}"
    return redirect(redir)



@bp.route("/bulletins/saisie", methods=["GET","POST"])
@login_required
def bulletin_saisie():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("auth.login"))
    if not current_user.can_edit: abort(403)
    sals=Salarie.query.filter_by(tenant_id=t.id,statut="ACTIF").order_by(Salarie.nom).all()
    pers=PeriodePaie.query.filter_by(tenant_id=t.id,statut="OUVERT").order_by(PeriodePaie.annee.desc(),PeriodePaie.mois.desc()).all()
    if request.method=="POST":
        sid=int(request.form["salarie_id"]); pid=int(request.form["periode_id"])
        s=Salarie.query.filter_by(id=sid,tenant_id=t.id).first_or_404()
        periode = PeriodePaie.query.filter_by(id=pid, tenant_id=t.id).first_or_404()
        acomptes_en_attente = Acompte.query.filter_by(
            tenant_id=t.id, salarie_id=sid,
            mois=periode.mois, annee=periode.annee, statut="EN_ATTENTE").all()
        total_acomptes = sum(float(a.montant) for a in acomptes_en_attente)
        donnees={}
        for k,v in request.form.items():
            # Exclure les champs non-numériques et les champs base_/taux_ (traités séparément)
            if k in ("salarie_id","periode_id","csrf_token","action","nb_jours_travailles"):
                continue
            if k.startswith("base_") or k.startswith("taux_") or k.startswith("composant_"):
                continue
            try:
                donnees[k] = float(v) if v else 0
            except (ValueError, TypeError):
                donnees[k] = 0
        if total_acomptes > 0:
            donnees["acompte"] = max(donnees.get("acompte", 0), total_acomptes)

        # ── Composants personnalisés du tenant (gains/retenues) ────────────────
        # Lecture des montants saisis (champ composant_<id>) ; on construit la
        # liste passée au calcul et on en garde une copie pour la persistance.
        composants_actifs = ComposantPaie.query.filter_by(tenant_id=t.id, actif=True).all()
        composants_saisis = []
        for comp in composants_actifs:
            montant = request.form.get(f"composant_{comp.id}", type=float) or 0
            if montant:
                composants_saisis.append({
                    "composant_id": comp.id, "libelle": comp.libelle, "sens": comp.sens,
                    "montant": montant, "soumis_cnss": comp.soumis_cnss,
                    "soumis_cnamgs": comp.soumis_cnamgs, "soumis_irpp": comp.soumis_irpp,
                    "entre_dans_brut": comp.entre_dans_brut, "position": comp.position,
                    "base": request.form.get(f"base_composant_{comp.id}", type=float),
                    "taux": request.form.get(f"taux_composant_{comp.id}", type=float),
                })
        donnees["composants"] = composants_saisis

        # ── Prime d'ancienneté automatique (selon convention collective) ───────
        # L'ancienneté est calculée à la fin de la période de paie. Le calcul
        # effectif se fait dans calculer_bulletin() et ne s'applique QUE si
        # aucune prime n'a été saisie manuellement (l'override reste prioritaire).
        if s.date_embauche:
            import calendar as _cal2
            _fin_periode = date(periode.annee, periode.mois,
                                _cal2.monthrange(periode.annee, periode.mois)[1])
            _anc = max(0, (_fin_periode - s.date_embauche).days // 365)
            donnees["anciennete_annees"] = _anc
            if getattr(s, "travail_de_nuit", False) and not donnees.get("travail_de_nuit"):
                donnees["travail_de_nuit"] = True
            if _anc >= 2 and not donnees.get("prime_anciennete"):
                from calculs_paie import prime_anciennete as _calc_pa
                _pa = _calc_pa(t.convention, float(donnees.get("salaire_base") or 0), _anc)
                if _pa > 0:
                    flash(
                        f"Prime d'ancienneté calculée automatiquement : "
                        f"{int(_pa):,} FCFA ({_anc} ans d'ancienneté, "
                        f"convention {t.convention}). Vous pouvez la modifier "
                        f"manuellement.".replace(",", " "),
                        "info")

        # ── L6 : Allocation de congé (Code du travail 2021, Art. 225) ──────────
        # Calcul automatique si un congé ANNUEL débute dans le mois de la période
        # et que l'utilisateur n'a pas saisi de valeur manuelle (l'override
        # manuel reste prioritaire). Versée avant le départ en congé (Art. 225).
        if not donnees.get("allocations_conge"):
            import calendar as _cal
            _debut_mois = date(periode.annee, periode.mois, 1)
            _fin_mois   = date(periode.annee, periode.mois,
                               _cal.monthrange(periode.annee, periode.mois)[1])
            jours_conge = sum(
                float(c.jours_pris or 0)
                for c in s.conges
                if (c.type_conge in ("ANNUEL", None))
                   and c.statut in ("APPROUVÉ", "APPROUVE", "PRIS")
                   and c.date_depart and _debut_mois <= c.date_depart <= _fin_mois
            )
            if jours_conge > 0:
                from conges_avance import allocation_conge
                bulletins_12 = (BulletinPaie.query
                    .filter_by(tenant_id=t.id, salarie_id=sid)
                    .filter(BulletinPaie.statut.in_(["VALIDÉ", "VALIDE", "PAYÉ"]),
                            BulletinPaie.periode_id != pid)
                    .order_by(BulletinPaie.date_creation.desc()).limit(12).all())
                alloc = allocation_conge(bulletins_12, jours_conge)
                if alloc > 0:
                    donnees["allocations_conge"] = alloc
                    flash(
                        f"Allocation de congé calculée automatiquement pour "
                        f"{jours_conge:.0f} jour(s) pris : {int(alloc):,} FCFA "
                        f"(moyenne des 12 derniers mois, Art. 225). "
                        f"Vous pouvez la modifier manuellement."
                        .replace(",", " "),
                        "info")

        donnees["config_rubriques"] = _config_rubriques_dict(t.id)
        res=calculer_bulletin(dict(donnees, convention=t.convention),nb_parts=float(s.nombre_parts or 1))
        ex=BulletinPaie.query.filter_by(tenant_id=t.id,salarie_id=sid,periode_id=pid).first()
        # 🔒 Immuabilité : un bulletin validé est un document de paie officiel.
        # Il ne peut pas être réécrit en silence — il faut d'abord annuler sa
        # validation (action tracée), sinon l'historique de paie serait modifiable.
        if ex and ex.statut in ("VALIDÉ", "VALIDE"):
            log_action("BLOCK_EDIT", "bulletin", ex.id,
                       "Tentative de modification d'un bulletin validé (bloquée)")
            flash("Ce bulletin est validé : il ne peut pas être modifié. "
                  "Annulez d'abord sa validation depuis la liste des bulletins, "
                  "puis ressaisissez-le.", "error")
            return redirect(url_for("tenant.bulletin_detail", id=ex.id))
        b=ex or BulletinPaie(tenant_id=t.id,salarie_id=sid,periode_id=pid)
        if not ex: db.session.add(b)
        for k,v in res.items():
            if not k.startswith("_") and hasattr(b,k): setattr(b,k,v)
        b.nb_jours_travailles=int(request.form.get("nb_jours_travailles") or 0)
        # ✅ Sauvegarder base et taux saisis manuellement pour chaque rubrique
        RUBRIQUES_BT = ["salaire_base","heures_sup_10","heures_sup_30","heures_sup_30b","heures_sup_40","heures_sup_70",
            "heures_sup_fj","heures_sup_fn",
            "absences","sursalaire","prime_caisse","carburant","prime_anciennete",
            "indem_logement","indem_domesticite","indem_eau_electricite","indem_nourriture",
            "prime_transport","prime_responsabilite","prime_rendement","prime_assiduité",
            "prime_qualite","prime_performance","allocations_conge",
            "indem_compensatrice_conge","indem_services_rendus",
            "indem_compensatrice_preavis","indem_licenciement"]
        for r in RUBRIQUES_BT:
            base_val = request.form.get(f"base_{r}", "")
            taux_val = request.form.get(f"taux_{r}", "")
            if hasattr(b, f"base_{r}"):
                try: setattr(b, f"base_{r}", float(base_val) if base_val else None)
                except: pass
            if hasattr(b, f"taux_{r}"):
                setattr(b, f"taux_{r}", taux_val.strip()[:20] if taux_val else "")
        action=request.form.get("action","brouillon")
        if action=="valider":
            b.statut="VALIDÉ"; b.date_validation=utcnow()
            attribuer_numero_bulletin(b)
            for a in acomptes_en_attente: a.statut = "DEDUIT"
        else:
            b.statut="BROUILLON"
        # Persistance des composants personnalisés (instantané par bulletin)
        db.session.flush()  # garantit b.id pour un nouveau bulletin
        BulletinComposant.query.filter_by(bulletin_id=b.id).delete()
        for cs in composants_saisis:
            db.session.add(BulletinComposant(
                bulletin_id=b.id, composant_id=cs["composant_id"],
                libelle=cs["libelle"], sens=cs["sens"], montant=cs["montant"],
                position=cs.get("position") or "BAS",
                base=cs.get("base"), taux=cs.get("taux"),
                soumis_cnss=cs["soumis_cnss"], soumis_cnamgs=cs["soumis_cnamgs"],
                soumis_irpp=cs["soumis_irpp"]))
        db.session.commit()
        if total_acomptes > 0:
            flash(f"Bulletin sauvegardé. Acompte de {int(total_acomptes):,} FCFA déduit automatiquement.".replace(",", " "), "success")
        else:
            flash(f"Bulletin {'validé' if b.statut=='VALIDÉ' else 'sauvegardé'}.","success")
        return redirect(url_for("tenant.bulletin_detail",id=b.id))
    sid=request.args.get("salarie_id",type=int)
    # ── Mode édition : ?id=<bulletin> pré-remplit le formulaire ──────────────
    bulletin_edit=None; valeurs_edit={}
    edit_id=request.args.get("id",type=int) or request.args.get("bulletin_id",type=int)
    if edit_id:
        be=BulletinPaie.query.filter_by(id=edit_id,tenant_id=t.id).first_or_404()
        if be.statut in ("VALIDÉ","VALIDE","PAYÉ","PAYE"):
            flash("Ce bulletin est validé ou payé : annulez d'abord sa validation "
                  "depuis la liste des bulletins pour pouvoir le modifier.","error")
            return redirect(url_for("tenant.bulletin_detail",id=be.id))
        bulletin_edit=be; sid=be.salarie_id
        _exclus={"id","tenant_id","salarie_id","periode_id","numero","numero_seq",
                 "statut","date_creation","date_validation","date_paiement","mode_paiement"}
        for col in be.__table__.columns:
            if col.name in _exclus: continue
            val=getattr(be,col.name)
            if val is None: continue
            try: valeurs_edit[col.name]=float(val)
            except (TypeError,ValueError): valeurs_edit[col.name]=val
        _comps=BulletinComposant.query.filter_by(bulletin_id=be.id).all()
        valeurs_edit["__composants"]={f"composant_{c.composant_id}":float(c.montant or 0) for c in _comps}
    ss=Salarie.query.filter_by(id=sid,tenant_id=t.id).first() if sid else None
    c=Contrat.query.filter_by(salarie_id=sid,tenant_id=t.id,actif=True).first() if sid else None
    acomptes_attente = Acompte.query.filter_by(tenant_id=t.id, salarie_id=sid, statut="EN_ATTENTE").all() if sid else []
    total_acomptes = sum(float(a.montant) for a in acomptes_attente)
    composants_actifs = ComposantPaie.query.filter_by(tenant_id=t.id, actif=True).order_by(
        ComposantPaie.ordre, ComposantPaie.libelle).all()
    return render_template("tenant/bulletin_saisie.html", salaries=sals, periodes=pers, salarie_sel=ss, contrat=c, tenant=t,
        acomptes_attente=acomptes_attente, total_acomptes=total_acomptes,
        composants_actifs=composants_actifs,
        bulletin_edit=bulletin_edit, valeurs_edit=valeurs_edit,
        periode_sel_id=(bulletin_edit.periode_id if bulletin_edit else None))

@bp.route("/bulletins/<int:id>")
@login_required
def bulletin_detail(id):
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("auth.login"))
    bulletin = BulletinPaie.query.filter_by(id=id,tenant_id=t.id).first_or_404()
    composants = BulletinComposant.query.filter_by(bulletin_id=bulletin.id).all()
    from datetime import date as _date
    return render_template("tenant/bulletin_detail.html",
        bulletin=bulletin, tenant=t, composants=composants,
        today=_date.today().isoformat())

@bp.route("/bulletins/<int:id>/valider", methods=["POST"])
@login_required
def bulletin_valider(id):
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if b.statut == "VALIDÉ":
        flash("Ce bulletin est déjà validé.", "info")
        return redirect(url_for("tenant.bulletin_detail", id=id))
    acomptes = Acompte.query.filter_by(
        tenant_id=t.id, salarie_id=b.salarie_id,
        mois=b.periode.mois, annee=b.periode.annee, statut="EN_ATTENTE").all()
    for a in acomptes: a.statut = "DEDUIT"
    b.statut = "VALIDÉ"; b.date_validation = utcnow()
    attribuer_numero_bulletin(b)
    db.session.commit()
    log_action("VALIDATE", "bulletin", b.id,
               f"Validation bulletin {b.salarie.nom_complet if b.salarie else ''} "
               f"(net {float(b.net_a_payer or 0):,.0f} F)".replace(",", " "))
    db.session.commit()
    flash("Bulletin validé avec succès.", "success")
    return redirect(url_for("tenant.bulletin_detail", id=id))

@bp.route("/bulletins/<int:id>/arrondi", methods=["POST"])
@login_required
def bulletin_arrondi(id):
    """Le tenant admin saisit le net à payer arrondi souhaité. Le système
    calcule l'écart et l'enregistre comme ligne d'ajustement d'arrondi, en
    gardant le bulletin cohérent. Interdit sur un bulletin validé."""
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))
    if not current_user.is_tenant_admin:
        flash("Seul l'administrateur peut ajuster le net à payer.", "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))

    b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if b.statut == "VALIDÉ":
        flash("Impossible d'ajuster un bulletin validé. Il faut d'abord le rouvrir.", "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))

    # Net actuellement calculé, hors ajustement précédent
    ajust_actuel = float(b.ajustement_arrondi or 0)
    net_calcule = float(b.net_a_payer or 0) - ajust_actuel

    action = request.form.get("action", "definir")
    if action == "reinitialiser":
        b.net_a_payer = round(net_calcule, 2)
        b.ajustement_arrondi = 0
        db.session.commit()
        log_action("UPDATE", "bulletin", b.id, "Ajustement d'arrondi retiré")
        flash("Ajustement d'arrondi retiré. Net rétabli au montant calculé.", "success")
        return redirect(url_for("tenant.bulletin_detail", id=id))

    # Net rond souhaité
    net_souhaite_raw = request.form.get("net_souhaite", "").replace(" ", "").replace(",", ".")
    try:
        net_souhaite = round(float(net_souhaite_raw), 2)
    except (ValueError, TypeError):
        flash("Montant invalide.", "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))

    ajustement = round(net_souhaite - net_calcule, 2)
    # Garde-fou : limiter l'ajustement à un vrai arrondi (±1000 F max), pour
    # éviter qu'on détourne cette fonction pour fausser un net librement.
    if abs(ajustement) > 1000:
        flash("L'ajustement d'arrondi est limité à ±1 000 F. "
              f"Écart demandé : {ajustement:,.0f} F.".replace(",", " "), "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))

    b.ajustement_arrondi = ajustement
    b.net_a_payer = round(net_calcule + ajustement, 2)
    db.session.commit()
    log_action("UPDATE", "bulletin", b.id,
               f"Ajustement d'arrondi : {ajustement:+,.0f} F "
               f"(net {net_souhaite:,.0f} F)".replace(",", " "))
    flash(f"Net à payer ajusté à {net_souhaite:,.0f} F "
          f"(arrondi de {ajustement:+,.0f} F).".replace(",", " "), "success")
    return redirect(url_for("tenant.bulletin_detail", id=id))


@bp.route("/bulletins/<int:id>/payer", methods=["POST"])
@login_required
def bulletin_paye(id):
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    mode = (request.form.get("mode_paiement") or "").strip()
    if mode not in ("ESPECES", "VIREMENT"):
        mode = (b.salarie.mode_paiement if b.salarie else "ESPECES") or "ESPECES"
    b.mode_paiement = mode
    # Date de paiement : celle choisie dans le formulaire, sinon aujourd'hui.
    from datetime import date as _date, datetime as _dt
    date_pmt = _date.today()
    _df = (request.form.get("date_paiement") or "").strip()
    if _df:
        try:
            date_pmt = _dt.strptime(_df, "%Y-%m-%d").date()
        except ValueError:
            date_pmt = _date.today()
    b.date_paiement = date_pmt
    b.statut = "PAYÉ"; db.session.commit()
    log_action("PAY", "bulletin", b.id,
               f"Bulletin payé {b.salarie.nom_complet if b.salarie else ''}")
    db.session.commit()
    # Interconnexion Caisse : proposer la sortie correspondante (ne bloque jamais).
    try:
        from interco_caisse import proposer_ecriture
        nom = b.salarie.nom_complet if b.salarie else "salarié"
        periode = f"{b.periode.libelle_mois} {b.periode.annee}" if b.periode else ""
        proposer_ecriture(
            t, source_ref=f"bulletin-{b.id}",
            montant=float(b.net_a_payer or 0),
            motif=f"Salaire {nom} — {periode}".strip(" —"),
            compte_suggere="6611", date_operation=b.date_paiement)
    except Exception as _e:
        current_app.logger.warning(f"[COMPTA] Proposition d'écriture échouée (bulletin {b.id}) : {_e}")
    flash("Bulletin marqué comme payé.", "success")
    return redirect(url_for("tenant.bulletin_detail", id=id))

@bp.route("/bulletins/<int:id>/supprimer", methods=["POST"])
@login_required
def bulletin_supprimer(id):
    if current_user.is_super_admin:
        b = BulletinPaie.query.get_or_404(id)
        salarie_id = b.salarie_id
        db.session.delete(b); db.session.commit()
        flash("Bulletin supprimé (super admin).", "success")
        return redirect(url_for("tenant.salarie_detail", id=salarie_id))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if b.statut == "VALIDÉ":
        flash("Impossible de supprimer un bulletin validé.", "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))
    db.session.delete(b); db.session.commit()
    flash("Bulletin supprimé.", "success")
    return redirect(url_for("tenant.bulletins"))

@bp.route("/bulletins/<int:id>/pdf")
@login_required
def bulletin_pdf(id):
    """Génère et retourne le bulletin en PDF téléchargeable."""
    if current_user.is_super_admin:
        b = BulletinPaie.query.get_or_404(id)
        t = b.salarie.tenant
    else:
        t = get_tenant()
        if not t: return redirect(url_for("auth.login"))
        b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    try:
        from pdf_bulletin import generer_bulletin_pdf
        pdf_bytes = _pdf_bulletin_bytes(b, t)
        nom_fichier = (
            f"bulletin_{b.salarie.nom}_{b.salarie.prenom}_{b.periode.annee}_{b.periode.mois:02d}.pdf"
            .replace(" ", "_")
        )
        from flask import Response
        return Response(
            pdf_bytes,
            mimetype="application/pdf",
            headers={
                "Content-Disposition": f'attachment; filename="{nom_fichier}"',
                "Content-Length": str(len(pdf_bytes)),
            }
        )
    except Exception as e:
        flash(f"Erreur génération PDF : {e}", "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))


@bp.route("/bulletins/export-zip/<int:periode_id>")
@login_required
def bulletins_export_zip(periode_id):
    """Télécharge TOUS les bulletins d'une période dans UN SEUL PDF (un par page)."""
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))
    p = PeriodePaie.query.filter_by(id=periode_id, tenant_id=t.id).first_or_404()
    bulletins = (BulletinPaie.query.filter_by(periode_id=periode_id, tenant_id=t.id)
                 .join(Salarie).order_by(Salarie.nom).all())
    if not bulletins:
        flash("Aucun bulletin à exporter pour cette période.", "error")
        return redirect(url_for("tenant.bulletins"))

    from pdf_bulletin import generer_bulletins_pdf
    try:
        data = generer_bulletins_pdf(bulletins, t, modele=(t.modele_bulletin or "classique"))
    except Exception as e:
        logger.error(f"Erreur génération PDF groupé période {periode_id} : {e}")
        flash(f"Erreur lors de la génération du PDF : {e}", "error")
        return redirect(url_for("tenant.bulletins"))

    log_action("EXPORT", "bulletin", periode_id,
               f"Export PDF groupé {len(bulletins)} bulletins — {p.libelle_complet}",
               user_id=current_user.id, tenant_id=t.id)
    db.session.commit()

    nom_pdf = f"bulletins_{p.libelle_mois}_{p.annee}_{t.slug}.pdf".replace(" ", "_")
    from flask import Response
    return Response(data, mimetype="application/pdf",
                    headers={"Content-Disposition": f'attachment; filename="{nom_pdf}"',
                             "Content-Length": str(len(data))})


@bp.route("/bulletins/<int:id>/imprimer")
@login_required
def bulletin_imprimer(id):
    """Aperçu imprimable du bulletin (HTML), disponible dès le brouillon."""
    return _bulletin_imprimer_impl(id)




@bp.route("/bulletins/<int:id>/envoyer-email", methods=["POST"])
@login_required
def bulletin_envoyer_email(id):
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    b = BulletinPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    s = b.salarie
    dest_email = request.form.get("email_dest", "").strip()
    if not dest_email and s.email:
        dest_email = s.email
    if not dest_email:
        flash(f"{s.nom_complet} n'a pas d'adresse email. Renseignez-en une dans le formulaire.", "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))
    if not os.environ.get("MAIL_PASSWORD"):
        flash("Email non configuré sur le serveur : ajoutez la clé API Resend "
              "(variable MAIL_PASSWORD) dans les variables d'environnement.", "error")
        return redirect(url_for("tenant.bulletin_detail", id=id))
    try:
        corps = (f"Bonjour {s.prenom},\n\n"
                 f"Veuillez trouver votre bulletin de paie pour : {b.periode.libelle_complet}\n\n"
                 f"Salaire brut : {int(b.salaire_brut or 0)} FCFA\n"
                 f"Net a payer  : {int(b.net_a_payer or 0)} FCFA\n\n"
                 f"Cordialement,\n{t.denomination}")
        msg = Message(
            subject=f"Bulletin de paie {b.periode.libelle_complet} — {t.denomination}",
            recipients=[dest_email],
            body=corps,
            sender=current_app.config["MAIL_DEFAULT_SENDER"]
        )
        # Joindre le bulletin en PDF
        from pdf_bulletin import generer_bulletin_pdf
        pdf_bytes = _pdf_bulletin_bytes(b, t)
        nom_pdf = (f"bulletin_{s.nom}_{s.prenom}_"
                   f"{b.periode.annee}_{b.periode.mois:02d}.pdf")
        msg.attach(nom_pdf, "application/pdf", pdf_bytes)
        # ✅ Envoi dans un thread séparé → le serveur répond immédiatement
        send_email_async(current_app.extensions["mail"], msg)
        flash(f"Email en cours d'envoi à {dest_email}.", "success")
    except Exception as e:
        flash(f"Erreur préparation email: {str(e)}", "error")
    return redirect(url_for("tenant.bulletin_detail", id=id))

@bp.route("/bulletins/envoyer-tous", methods=["POST"])
@login_required
def bulletins_envoyer_tous():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t: return redirect(url_for("auth.login"))
    periode_id = request.form.get("periode_id", type=int)
    if not periode_id: flash("Période manquante.", "error"); return redirect(url_for("tenant.bulletins"))
    if not os.environ.get("MAIL_PASSWORD"):
        flash("Email non configuré sur le serveur : ajoutez la clé API Resend "
              "(variable MAIL_PASSWORD) dans les variables d'environnement.", "error")
        return redirect(url_for("tenant.bulletins"))
    buls = BulletinPaie.query.filter_by(tenant_id=t.id, periode_id=periode_id).all()
    nb_ok=0; nb_sans_email=0
    for b in buls:
        if not b.salarie.email: nb_sans_email+=1; continue
        try:
            corps = (f"Bonjour {b.salarie.prenom},\n\n"
                     f"Bulletin {b.periode.libelle_complet}\n"
                     f"Net a payer : {int(b.net_a_payer or 0)} FCFA\n\n"
                     f"Cordialement, {t.denomination}")
            msg = Message(subject=f"Bulletin {b.periode.libelle_complet}",
                recipients=[b.salarie.email], body=corps,
                sender=current_app.config["MAIL_DEFAULT_SENDER"])
            from pdf_bulletin import generer_bulletin_pdf
            pdf_bytes = _pdf_bulletin_bytes(b, t)
            nom_pdf = (f"bulletin_{b.salarie.nom}_{b.salarie.prenom}_"
                       f"{b.periode.annee}_{b.periode.mois:02d}.pdf")
            msg.attach(nom_pdf, "application/pdf", pdf_bytes)
            send_email_async(current_app.extensions["mail"], msg)
            nb_ok+=1
        except Exception as e:
            print(f"Erreur email {b.salarie.email}: {e}")
    flash(f"{nb_ok} email(s) en cours d'envoi. {nb_sans_email} salarié(s) sans email.", "success")
    return redirect(url_for("tenant.bulletins"))

# ── Périodes ──────────────────────────────────────────────────────────────────
@bp.route("/periodes")
@login_required
def periodes():
    if current_user.is_super_admin: return redirect(url_for("admin.admin_dashboard"))
    t=get_tenant()
    if not t: return redirect(url_for("auth.login"))
    periodes_liste = (PeriodePaie.query.filter_by(tenant_id=t.id)
                      .order_by(PeriodePaie.annee.desc(), PeriodePaie.mois.desc()).all())

    # Masse brute et nombre de bulletins par période (une seule requête groupée)
    stats = {}
    for pid, nb, brut, net in (db.session.query(
            BulletinPaie.periode_id,
            db.func.count(BulletinPaie.id),
            db.func.sum(BulletinPaie.salaire_brut),
            db.func.sum(BulletinPaie.net_a_payer))
            .filter_by(tenant_id=t.id)
            .group_by(BulletinPaie.periode_id).all()):
        stats[pid] = {"nb": nb or 0, "brut": float(brut or 0),
                      "net": float(net or 0), "brouillons": 0}

    # Brouillons restants : ce sont eux qui bloquent une clôture sereine
    for pid, nb in (db.session.query(
            BulletinPaie.periode_id, db.func.count(BulletinPaie.id))
            .filter_by(tenant_id=t.id, statut="BROUILLON")
            .group_by(BulletinPaie.periode_id).all()):
        if pid in stats:
            stats[pid]["brouillons"] = nb or 0

    return render_template("tenant/periodes.html", tenant=t,
        periodes=periodes_liste, stats_periodes=stats,
        now=datetime.now())

@bp.route("/periodes/<int:id>")
@login_required
def periode_detail(id):
    """Fiche récapitulative d'une période : masse salariale, retenues
    salariales, charges patronales et accès aux déclarations."""
    if current_user.is_super_admin:
        return redirect(url_for("admin.admin_dashboard"))
    t = get_tenant()
    if not t:
        return redirect(url_for("auth.login"))

    p = PeriodePaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    bulletins = (BulletinPaie.query.options(joinedload(BulletinPaie.salarie))
                 .filter_by(tenant_id=t.id, periode_id=p.id).all())

    def somme(champ, source=None):
        return round(sum(float(getattr(b, champ) or 0) for b in (source or bulletins)), 2)

    # Répartition par statut
    par_statut = {"BROUILLON": 0, "VALIDÉ": 0, "PAYÉ": 0}
    for b in bulletins:
        st = "VALIDÉ" if b.statut in ("VALIDÉ", "VALIDE") else b.statut
        par_statut[st] = par_statut.get(st, 0) + 1

    # Masse salariale
    brut = somme("salaire_brut")
    net = somme("salaire_net")
    net_a_payer = somme("net_a_payer")

    # Retenues salariales
    retenues = {
        "cnss": somme("cnss_salarie"),
        "cnamgs": somme("cnamgs_salarie"),
        "tcs": somme("tcs"),
        "irpp": somme("irpp"),
    }
    retenues["total"] = round(sum(retenues.values()), 2)

    # Charges patronales
    patronales = {
        "cnss": somme("cnss_patronale"),
        "cnamgs": somme("cnamgs_patronale"),
        "fnh": somme("fnh"),
        "cfp": somme("cfp"),
    }
    patronales["total"] = round(sum(patronales.values()), 2)

    # Reversements par organisme (part salariale + part patronale)
    organismes = [
        {"nom": "CNSS", "salarial": retenues["cnss"], "patronal": patronales["cnss"],
         "total": round(retenues["cnss"] + patronales["cnss"], 2),
         "note": "Déclaration trimestrielle"},
        {"nom": "CNAMGS", "salarial": retenues["cnamgs"], "patronal": patronales["cnamgs"],
         "total": round(retenues["cnamgs"] + patronales["cnamgs"], 2),
         "note": "Déclaration trimestrielle"},
        {"nom": "TCS", "salarial": retenues["tcs"], "patronal": 0,
         "total": retenues["tcs"], "note": "Taxe complémentaire sur les salaires"},
        {"nom": "IRPP", "salarial": retenues["irpp"], "patronal": 0,
         "total": retenues["irpp"], "note": "Retenue à la source, reversée à la DGI"},
        {"nom": "FNH", "salarial": 0, "patronal": patronales["fnh"],
         "total": patronales["fnh"], "note": "Fonds national de l'habitat"},
        {"nom": "CFP", "salarial": 0, "patronal": patronales["cfp"],
         "total": patronales["cfp"], "note": "Contribution à la formation professionnelle"},
    ]
    total_reversements = round(sum(o["total"] for o in organismes), 2)

    stats = {
        "nb_bulletins": len(bulletins),
        "par_statut": par_statut,
        "nb_brouillons": par_statut.get("BROUILLON", 0),
        "brut": brut, "net": net, "net_a_payer": net_a_payer,
        "base_cnss": somme("base_cnss"),
        "base_cnamgs": somme("base_cnamgs"),
        "base_tcs": somme("base_tcs"),
        "base_irpp": somme("base_irpp"),
        "cout_employeur": round(brut + patronales["total"], 2),
    }

    return render_template("tenant/periode_detail.html", tenant=t, p=p,
        bulletins=bulletins, stats=stats, retenues=retenues,
        patronales=patronales, organismes=organismes,
        total_reversements=total_reversements, now=datetime.now())


@bp.route("/periodes/nouvelle", methods=["POST"])
@tenant_required
@can_edit
def periode_nouvelle():
    t=get_tenant(); annee=int(request.form["annee"]); mois=int(request.form["mois"])
    noms=PeriodePaie.MOIS_NOMS
    if PeriodePaie.query.filter_by(tenant_id=t.id,annee=annee,mois=mois).first(): flash("Période existante.","warning")
    else:
        db.session.add(PeriodePaie(tenant_id=t.id,annee=annee,mois=mois,libelle_mois=noms[mois],
            trimestre=f"T{(mois-1)//3+1}",statut="OUVERT",date_ouverture=utcnow()))
        db.session.commit(); flash(f"Période {noms[mois]} {annee} créée.","success")
    return redirect(url_for("tenant.periodes"))

@bp.route("/periodes/<int:id>/cloturer", methods=["POST"])
@tenant_required
@can_edit
def periode_cloturer(id):
    """Clôture une période, en refusant de le faire à l'aveugle s'il reste des
    brouillons : ce sont des salariés potentiellement non payés. La clôture
    reste possible, mais elle doit alors être confirmée explicitement."""
    t = get_tenant()
    p = PeriodePaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()

    if p.statut != "OUVERT":
        flash("Cette période est déjà clôturée.", "warning")
        return redirect(url_for("tenant.periodes"))

    nb_brouillons = BulletinPaie.query.filter_by(
        tenant_id=t.id, periode_id=p.id, statut="BROUILLON").count()

    if nb_brouillons and request.form.get("confirmer") != "1":
        flash(f"Clôture annulée : {nb_brouillons} bulletin(s) encore en brouillon "
              f"sur {p.libelle_complet}. Validez-les d'abord, ou confirmez la "
              f"clôture depuis la fiche de la période si c'est volontaire.", "error")
        return redirect(url_for("tenant.periode_detail", id=p.id))

    p.statut = "CLÔTURÉ"
    p.date_cloture = utcnow()
    db.session.commit()
    log_action("CLOSE", "periode", p.id,
               f"Clôture de {p.libelle_complet}"
               + (f" avec {nb_brouillons} brouillon(s) restant(s)" if nb_brouillons else ""))

    if nb_brouillons:
        flash(f"Période {p.libelle_complet} clôturée avec {nb_brouillons} "
              f"brouillon(s) non validé(s). Rouvrez-la si vous devez les traiter.",
              "warning")
    else:
        flash(f"Période {p.libelle_complet} clôturée.", "success")
    return redirect(url_for("tenant.periodes"))


@bp.route("/periodes/<int:id>/rouvrir", methods=["POST"])
@tenant_required
@can_edit
def periode_rouvrir(id):
    """Rouvre une période clôturée. Une clôture par erreur reste ainsi
    rattrapable ; l'opération est tracée dans le journal d'audit."""
    t = get_tenant()
    p = PeriodePaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()

    if p.statut == "OUVERT":
        flash("Cette période est déjà ouverte.", "info")
        return redirect(url_for("tenant.periode_detail", id=p.id))

    ancienne_cloture = p.date_cloture
    p.statut = "OUVERT"
    p.date_cloture = None
    db.session.commit()
    log_action("REOPEN", "periode", p.id,
               f"Réouverture de {p.libelle_complet}"
               + (f" (clôturée le {ancienne_cloture:%d/%m/%Y})" if ancienne_cloture else ""))
    flash(f"Période {p.libelle_complet} rouverte. Vous pouvez de nouveau y "
          f"générer et valider des bulletins.", "success")
    return redirect(url_for("tenant.periode_detail", id=p.id))

# ── Paiement abonnement ───────────────────────────────────────────────────────


@bp.route("/composants")
@tenant_required
def composants():
    t = get_tenant()
    composants_list = ComposantPaie.query.filter_by(tenant_id=t.id).order_by(
        ComposantPaie.ordre, ComposantPaie.libelle).all()
    return render_template("tenant/composants.html", tenant=t, composants=composants_list)


@bp.route("/composants/nouveau", methods=["GET", "POST"])
@tenant_required
def composant_nouveau():
    t = get_tenant()
    if not current_user.can_edit: abort(403)
    if request.method == "POST":
        libelle = request.form.get("libelle", "").strip()
        if not libelle:
            flash("Le libellé est obligatoire.", "error")
            return render_template("tenant/composant_form.html", tenant=t, composant=None)
        c = ComposantPaie(
            tenant_id     = t.id,
            libelle       = libelle,
            sens          = "RETENUE" if request.form.get("sens") == "RETENUE" else "GAIN",
            soumis_cnss   = request.form.get("soumis_cnss")   == "on",
            soumis_cnamgs = request.form.get("soumis_cnamgs") == "on",
            soumis_irpp   = request.form.get("soumis_irpp")   == "on",
            entre_dans_brut = request.form.get("entre_dans_brut") == "on",
            position      = "HAUT" if request.form.get("position") == "HAUT" else "BAS",
            ordre         = request.form.get("ordre", type=int) or 0,
            actif         = True,
        )
        db.session.add(c)
        db.session.commit()
        flash(f"Composant « {c.libelle} » créé.", "success")
        return redirect(url_for("tenant.composants"))
    return render_template("tenant/composant_form.html", tenant=t, composant=None)


@bp.route("/composants/<int:id>/modifier", methods=["GET", "POST"])
@tenant_required
def composant_modifier(id):
    t = get_tenant()
    if not current_user.can_edit: abort(403)
    c = ComposantPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    if request.method == "POST":
        libelle = request.form.get("libelle", "").strip()
        if not libelle:
            flash("Le libellé est obligatoire.", "error")
            return render_template("tenant/composant_form.html", tenant=t, composant=c)
        c.libelle       = libelle
        c.sens          = "RETENUE" if request.form.get("sens") == "RETENUE" else "GAIN"
        c.soumis_cnss   = request.form.get("soumis_cnss")   == "on"
        c.soumis_cnamgs = request.form.get("soumis_cnamgs") == "on"
        c.soumis_irpp   = request.form.get("soumis_irpp")   == "on"
        c.entre_dans_brut = request.form.get("entre_dans_brut") == "on"
        c.position      = "HAUT" if request.form.get("position") == "HAUT" else "BAS"
        c.ordre         = request.form.get("ordre", type=int) or 0
        db.session.commit()
        flash(f"Composant « {c.libelle} » modifié.", "success")
        return redirect(url_for("tenant.composants"))
    return render_template("tenant/composant_form.html", tenant=t, composant=c)


@bp.route("/composants/<int:id>/toggle", methods=["POST"])
@tenant_required
def composant_toggle(id):
    t = get_tenant()
    if not current_user.can_edit: abort(403)
    c = ComposantPaie.query.filter_by(id=id, tenant_id=t.id).first_or_404()
    c.actif = not c.actif
    db.session.commit()
    flash(f"Composant « {c.libelle} » {'activé' if c.actif else 'désactivé'}.", "success")
    return redirect(url_for("tenant.composants"))




@bp.route("/bulletins/export/<int:periode_id>")
@tenant_required
def export_journal(periode_id):
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.styles import Font, PatternFill, Alignment
    t=get_tenant()
    p=PeriodePaie.query.filter_by(id=periode_id,tenant_id=t.id).first_or_404()
    buls=BulletinPaie.query.filter_by(periode_id=periode_id,tenant_id=t.id).join(Salarie).order_by(Salarie.nom).all()
    wb=Workbook(); ws=wb.active; ws.title=f"Journal {p.libelle_complet}"
    ws.merge_cells("A1:R1"); ws["A1"]=f"JOURNAL DE PAIE — {p.libelle_complet} — {t.denomination}"
    ws["A1"].font=Font(bold=True,size=13); ws["A1"].alignment=Alignment(horizontal="center")
    hdrs=["Matricule","Nom","Prénom","Emploi","Cat.","Base","Brut","CNSS Sal.","CNAMGS Sal.","TCS","IRPP","Net","Net à Payer","CNSS Pat.","CNAMGS Pat.","FNH","CFP","Statut"]
    for col,h in enumerate(hdrs,1):
        c=ws.cell(row=3,column=col,value=h); c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="1a2332"); c.alignment=Alignment(horizontal="center")
    for row,b in enumerate(buls,4):
        s=b.salarie
        vals=[s.matricule,s.nom,s.prenom,s.emploi,s.categorie.code if s.categorie else "",
              float(b.salaire_base or 0),float(b.salaire_brut or 0),float(b.cnss_salarie or 0),
              float(b.cnamgs_salarie or 0),float(b.tcs or 0),float(b.irpp or 0),
              float(b.salaire_net or 0),float(b.net_a_payer or 0),float(b.cnss_patronale or 0),
              float(b.cnamgs_patronale or 0),float(b.fnh or 0),float(b.cfp or 0),b.statut]
        for col,v in enumerate(vals,1):
            cell=ws.cell(row=row,column=col,value=csv_safe(v) if isinstance(v,str) else v)
            if isinstance(v,float): cell.number_format='#,##0'
            if row%2==0: cell.fill=PatternFill("solid",fgColor="F5F5F5")
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,download_name=f"Journal_{p.libelle_mois}_{p.annee}_{t.slug}.xlsx")

@bp.route("/api/calculer-bulletin", methods=["POST"])
@login_required
def api_calculer():
    try:
        t = get_tenant()
        data = request.get_json() or {}
        sid = data.pop("salarie_id", None)
        nb_parts = 1.0
        if sid and t:
            s = Salarie.query.filter_by(id=sid, tenant_id=t.id).first()
            if s: nb_parts = float(s.nombre_parts or 1)
        mois  = data.pop("mois_periode", None)
        annee = data.pop("annee_periode", None)
        # ── Ancienneté du salarié (pour la prime d'ancienneté automatique) ──
        # Référence : fin de la période de paie si connue, sinon aujourd'hui.
        if sid and t:
            s_anc = Salarie.query.filter_by(id=sid, tenant_id=t.id).first()
            if s_anc and s_anc.date_embauche:
                from datetime import date as _date
                try:
                    if mois and annee:
                        m, a = int(mois), int(annee)
                        ref = (_date(a + 1, 1, 1) if m == 12 else _date(a, m + 1, 1))
                        ref = ref - timedelta(days=1)      # dernier jour du mois
                    else:
                        ref = _date.today()
                except (TypeError, ValueError):
                    ref = _date.today()
                data["anciennete_annees"] = max(0, (ref - s_anc.date_embauche).days // 365)
        total_acomptes = 0.0
        if sid and t and mois and annee:
            total_acomptes = float(db.session.query(db.func.sum(Acompte.montant))
                .filter_by(tenant_id=t.id, salarie_id=int(sid), mois=int(mois),
                           annee=int(annee), statut="EN_ATTENTE").scalar() or 0)
        if total_acomptes > 0:
            data["acompte"] = max(float(data.get("acompte", 0)), total_acomptes)
        # Composants personnalisés (aperçu temps réel) : lus depuis composant_<id>
        if t:
            comps_live = []
            for comp in ComposantPaie.query.filter_by(tenant_id=t.id, actif=True).all():
                try:
                    montant = float(data.get(f"composant_{comp.id}") or 0)
                except (TypeError, ValueError):
                    montant = 0.0
                if montant:
                    comps_live.append({
                        "libelle": comp.libelle, "sens": comp.sens, "montant": montant,
                        "soumis_cnss": comp.soumis_cnss, "soumis_cnamgs": comp.soumis_cnamgs,
                        "soumis_irpp": comp.soumis_irpp,
                        "entre_dans_brut": comp.entre_dans_brut, "position": comp.position})
            data["composants"] = comps_live
        if t:
            data["convention"] = t.convention
            data["config_rubriques"] = _config_rubriques_dict(t.id)
        res = calculer_bulletin(data, nb_parts=nb_parts)
        res["acompte_auto"] = total_acomptes
        return jsonify(res)
    except Exception as e:
        return jsonify({"error": str(e)}), 500

@bp.route("/api/salarie/<int:id>/contrat")
@login_required
def api_contrat(id):
    t=get_tenant()
    s=Salarie.query.filter_by(id=id,tenant_id=t.id).first()
    if not s: return jsonify({})
    c=Contrat.query.filter_by(salarie_id=id,tenant_id=t.id,actif=True).first()
    base={"nom":s.nom_complet,"poste":s.emploi,"matricule":s.matricule,"nombre_parts":float(s.nombre_parts or 1)}
    if c:
        base["salaire_base"]=float(c.salaire_base); base["poste"]=c.poste or s.emploi
        try:
            import json as _json
            base["elements"] = _json.loads(c.elements_recurrents) if c.elements_recurrents else {}
        except Exception:
            base["elements"] = {}
    return jsonify(base)

@bp.route("/api/salarie/<int:id>/pointage-mois")
@login_required
def api_pointage_mois(id):
    """Retourne le cumul des heures du pointage pour un salarié sur un mois donné."""
    t = get_tenant()
    if not t: return jsonify({})
    mois  = request.args.get("mois",  type=int)
    annee = request.args.get("annee", type=int)
    if not mois or not annee:
        return jsonify({"erreur": "mois et annee requis"})
    import calendar
    dernier_jour = calendar.monthrange(annee, mois)[1]
    debut = date(annee, mois, 1)
    fin   = date(annee, mois, dernier_jour)
    pts = Pointage.query.filter_by(tenant_id=t.id, salarie_id=id)        .filter(Pointage.date_pointage >= debut, Pointage.date_pointage <= fin).all()
    pts_travailles = [p for p in pts if p.present and not p.absent]
    if not pts_travailles:
        return jsonify({"nb_jours": 0, "nb_absences": 0,
            "heures_sup_10": 0, "heures_sup_30": 0, "heures_sup_40": 0, "heures_sup_70": 0,
            "heures_normales_total": 0, "total_sup": 0,
            "message": "Aucun pointage pour cette période"})
    nb_jours = len(pts_travailles)

    if (t.convention or "").upper() in ("BTP", "PETROLE", "INDUSTRIE", "AERIEN", "MINIER"):
        # Ventilation réglementaire : semaine par semaine, ligne par ligne
        from calculs_paie import ventiler_heures_mois, pointage_vers_jours
        _feries = set()
        for _an in {p.date.year for p in pts if getattr(p, "date", None)}:
            _feries |= set(jours_feries_annee(_an).keys())
        v = ventiler_heures_mois(t.convention, pointage_vers_jours(pts), feries=_feries, seuil_normales=t.seuil_hs)
        heures_normales = v["heures_normales"]
        heures_sup_10   = v["heures_sup_10"]
        heures_sup_30   = v["heures_sup_30"]
        heures_sup_30b  = v.get("heures_sup_30b", 0.0)
        heures_sup_40   = v["heures_sup_40"]
        heures_sup_70   = v["heures_sup_70"]
        heures_sup_fj   = v.get("heures_sup_fj", 0.0)
        heures_sup_fn   = v.get("heures_sup_fn", 0.0)
        detail_semaines = v["detail_semaines"]
    else:
        # Autres conventions : cumul direct des colonnes déjà ventilées par jour
        heures_normales = sum(float(p.heures_normales or 8) for p in pts_travailles)
        heures_sup_10   = sum(float(p.heures_sup_10 or 0) for p in pts_travailles)
        heures_sup_30   = sum(float(p.heures_sup_30 or 0) for p in pts_travailles)
        heures_sup_30b  = sum(float(getattr(p, "heures_sup_30b", 0) or 0) for p in pts_travailles)
        heures_sup_40   = sum(float(p.heures_sup_40 or 0) for p in pts_travailles)
        heures_sup_70   = sum(float(p.heures_sup_70 or 0) for p in pts_travailles)
        heures_sup_fj   = 0.0
        heures_sup_fn   = 0.0
        detail_semaines = []
    pts_absents = Pointage.query.filter_by(tenant_id=t.id, salarie_id=id)        .filter(Pointage.date_pointage >= debut, Pointage.date_pointage <= fin,
                Pointage.absent == True).all()
    return jsonify({
        "nb_jours":              nb_jours,
        "nb_absences":           len(pts_absents),
        "heures_normales_total": round(heures_normales, 2),
        "heures_sup_10":         round(heures_sup_10, 2),
        "heures_sup_30":         round(heures_sup_30, 2),
        "heures_sup_30b":        round(heures_sup_30b, 2),
        "heures_sup_40":         round(heures_sup_40, 2),
        "heures_sup_70":         round(heures_sup_70, 2),
        "heures_sup_fj":         round(heures_sup_fj, 2),
        "heures_sup_fn":         round(heures_sup_fn, 2),
        "total_sup":             round(heures_sup_10+heures_sup_30+heures_sup_30b+heures_sup_40+heures_sup_70+heures_sup_fj+heures_sup_fn, 2),
        "detail_semaines":       detail_semaines,
        "message":               f"{nb_jours} jour(s) pointé(s) sur {dernier_jour}"
    })



@bp.route("/api/salarie/<int:id>/acomptes-mois")
@login_required
def api_acomptes_mois(id):
    t = get_tenant()
    mois=request.args.get("mois",type=int); annee=request.args.get("annee",type=int)
    if not t or not mois or not annee: return jsonify({"total":0})
    total = db.session.query(db.func.sum(Acompte.montant))\
            .filter_by(tenant_id=t.id,salarie_id=id,mois=mois,annee=annee,statut="EN_ATTENTE").scalar() or 0
    return jsonify({"total":float(total)})

